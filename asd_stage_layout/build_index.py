"""Stage-wise plot index workbook for the white rust ASD dataset.

Usage: python build_index.py deck_data.json Stage_wise_plot_index.xlsx
"""
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

D = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
DATES = ["110226", "240226", "030326", "130326"]
DLAB = {"110226": "11/02/2026", "240226": "24/02/2026", "030326": "03/03/2026", "130326": "13/03/2026"}
DAI = {"110226": 6, "240226": 19, "030326": 26, "130326": 36}
STAGES = ["Vegetative", "Flowering", "Pod formation"]
SHEET = {"Vegetative": "Vegetative", "Flowering": "Flowering", "Pod formation": "Pod_formation"}
DUP = {"130326": {13, 34, 38, 41, 42}}

F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
FW = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FT = Font(name="Arial", size=12, bold=True, color="1F4D2B")
FIN = Font(name="Arial", size=10, color="0000FF")
HDR = PatternFill("solid", fgColor="1F4D2B")
YEL = PatternFill("solid", fgColor="FFFF00")
SFILL = {"Vegetative": PatternFill("solid", fgColor="D5E8D4"), "Flowering": PatternFill("solid", fgColor="FBEFC4"),
         "Pod formation": PatternFill("solid", fgColor="EBD3C9")}
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

rows = []
for d in DATES:
    for p in sorted(D["plots"][d], key=int):
        x = D["plots"][d][p]
        note = "15 spectra were duplicated in the file; copies removed" if int(p) in DUP.get(d, set()) else ""
        rows.append([x["stage"], DLAB[d], DAI[d], x["das"], int(p), x["t"], x["sow"], x["var"],
                     "Control" if x["rep"] == "Healthy" else x["rep"], x["hd"], x["n"], x["first"], x["last"], note])
rows.sort(key=lambda r: (STAGES.index(r[0]), DATES.index([k for k, v in DLAB.items() if v == r[1]][0]), r[4]))

COLS = ["Growth stage", "Assessment date", "DAI", "DAS", "Plot", "Treatment", "Sowing", "Variety", "Replication",
        "Healthy/Diseased", "Unique spectra", "First spectrum file", "Last spectrum file", "Note",
        "Plot DI (%) - enter", "DI class (auto)"]
WID = [14, 14, 6, 6, 6, 10, 8, 8, 11, 15, 9, 22, 22, 44, 13, 14]
CLASS_F = ('=IF(O{r}="","",IF(O{r}=0,"Healthy",IF(O{r}<=5,"Very low",IF(O{r}<=10,"Low",'
           'IF(O{r}<=25,"Moderate",IF(O{r}<=50,"High","Very high"))))))')

wb = openpyxl.Workbook()
rd = wb.active
rd.title = "README"
txt = [
    ("STAGE-WISE PLOT INDEX - WHITE RUST ASD LEAF SPECTRA, IARI 2025-26", FT),
    ("Every plot-date (48 plots x 4 dates = 192) with its growth stage, treatment and the range of ASD files it holds.", F),
    ("Source: asd__data_complete_labelled_110226 / 240226 / 030326 / 130326.xlsx and treatment_detail.xlsx (user files).", F),
    ("Stage, variety and healthy/diseased labels in the ASD files were checked against the layout and stage table: all 192 agree.", F),
    ("13/03: plots 13, 34, 38, 41, 42 held each of their 15 spectra twice; spectrum counts here exclude the copies.", F),
    ("", F),
    ("TO FILL IN (yellow, blue text): 'Plot DI (%)' on each stage sheet. 'DI class' then follows the conventional PDI bands: "
     "0 Healthy, 0.1-5 Very low, 5.1-10 Low, 10.1-25 Moderate, 25.1-50 High, >50 Very high (Shrivastava et al. 2023).", FB),
    ("Sheets: All_plot_dates; one sheet per stage (the training set of each stage model); Summary (counts by formula).", F),
]
for i, (t, f) in enumerate(txt, start=1):
    c = rd.cell(row=i, column=1, value=t)
    c.font = f
rd.column_dimensions["A"].width = 150


def write(ws, data, editable):
    for j, h in enumerate(COLS if editable else COLS[:14], start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = FW, HDR, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 30
    for i, r in enumerate(data, start=2):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = F, BOX
            if j == 1:
                c.fill = SFILL[r[0]]
        if editable:
            c = ws.cell(row=i, column=15, value=0 if r[9] == "Healthy" else None)
            c.font, c.fill, c.border = FIN, YEL, BOX
            c.number_format = "0.0"
            c = ws.cell(row=i, column=16, value=CLASS_F.format(r=i))
            c.font, c.border = F, BOX
    n = len(COLS) if editable else 14
    for j in range(1, n + 1):
        ws.column_dimensions[CL(j)].width = WID[j - 1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{CL(n)}{len(data) + 1}"
    if editable:
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True,
                            error="DI must be between 0 and 100")
        ws.add_data_validation(dv)
        dv.add(f"O2:O{len(data) + 1}")


write(wb.create_sheet("All_plot_dates"), rows, False)
for st in STAGES:
    write(wb.create_sheet(SHEET[st]), [r for r in rows if r[0] == st], True)

sm = wb.create_sheet("Summary")
sm.cell(row=1, column=1, value="Counts by formula from All_plot_dates").font = FT
hdr = ["Growth stage", "Assessment date", "Plots healthy", "Plots diseased", "Spectra healthy", "Spectra diseased", "Spectra total"]
for j, h in enumerate(hdr, start=1):
    c = sm.cell(row=3, column=j, value=h)
    c.font, c.fill, c.border = FW, HDR, BOX
    c.alignment = Alignment(wrap_text=True, horizontal="center")
r = 4
A = "All_plot_dates!"
for st in STAGES:
    for d in DATES:
        if not any(x[0] == st and x[1] == DLAB[d] for x in rows):
            continue
        sm.cell(row=r, column=1, value=st).fill = SFILL[st]
        sm.cell(row=r, column=2, value=DLAB[d])
        sm.cell(row=r, column=3, value=f'=COUNTIFS({A}$A:$A,A{r},{A}$B:$B,B{r},{A}$J:$J,"Healthy")')
        sm.cell(row=r, column=4, value=f'=COUNTIFS({A}$A:$A,A{r},{A}$B:$B,B{r},{A}$J:$J,"Diseased")')
        sm.cell(row=r, column=5, value=f'=SUMIFS({A}$K:$K,{A}$A:$A,A{r},{A}$B:$B,B{r},{A}$J:$J,"Healthy")')
        sm.cell(row=r, column=6, value=f'=SUMIFS({A}$K:$K,{A}$A:$A,A{r},{A}$B:$B,B{r},{A}$J:$J,"Diseased")')
        sm.cell(row=r, column=7, value=f"=E{r}+F{r}")
        for j in range(1, 8):
            sm.cell(row=r, column=j).font = F
            sm.cell(row=r, column=j).border = BOX
        r += 1
sm.cell(row=r, column=1, value="Total").font = FB
for j, col in zip(range(3, 8), "CDEFG"):
    c = sm.cell(row=r, column=j, value=f"=SUM({col}4:{col}{r - 1})")
    c.font, c.border = FB, BOX
for j, w in enumerate([16, 16, 12, 12, 12, 12, 12], start=1):
    sm.column_dimensions[CL(j)].width = w
wb.save(OUT)
print("saved", OUT, len(rows))
