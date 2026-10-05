"""Batch 15 rules (updated 15): rules 108-112 with the author decisions of 2026-10-05."""
import sys
from copy import copy

import openpyxl

p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
ROWS = [
    ("Treatment codes",
     "A new companion paper of an old-master trial is coded like the trial's existing AUTHOR-CONFIRMED companions, with partial codes added for treatments the old entries "
     "excluded - even when a later rule (e.g. rule 103 for tilled DSR) would code it differently (flagged).",
     "59 (companion 3) Singh G. 2022: DSR-ZTW (+BM / +MBR) ZT, + rice residue CA (DSR dry-tilled in years 1-3), TPR-ZTW pZT, TPR-CTW CT.", "Author 2026-10-05", "IN FORCE (author)"),
    ("Treatment codes",
     "A paper that contradicts itself on residue retention (treatment table 'removed' vs results table listing retained tonnes) is coded from the TREATMENT-DESCRIPTION table, "
     "flagged.",
     "48 (companion) Mondal 2021 pCA1 (NT wheat after CT rice): Table 1 residue removed (Table 5 lists 2.0-2.4 Mg/ha) -> pZT.", "Author 2026-10-05", "IN FORCE (author)"),
    ("Parameter-specific",
     "Aggregate-associated C printed as AMOUNTS (per kg bulk soil = concentration x class mass fraction) is converted to a concentration in the aggregate class (macro = "
     "sum of >0.25 mm amounts / sum of their mass fractions; micro = 0.053-0.25 mm amount / its mass fraction) and entered on macro c / micro c - for SOC, KMnO4-labile C "
     "and DISSOLVED organic C alike (DOC in g/kg aggregate).",
     "177 Zhao 2021 Table 2 (SOC, LOC, DOC amounts) with the Fig. 1 class masses.", "Author 2026-10-05 (DOC on macro / micro c)", "IN FORCE (author)"),
    ("Site classification",
     "A reported 50-80 cm layer goes to >60 CM by maximum overlap (20 cm in >60 vs 10 cm in 45-60), not to the sequential 45-60 CM class.",
     "176 Devkota 2015 EJA ECe 50-80 cm.", "Author 2026-10-05", "IN FORCE (author)"),
    ("Parameter-specific",
     "Aggregate C STOCKS (Mg/ha per aggregate class and layer) are converted to aggregate C concentrations by inverting the paper's own stock equation "
     "(C_agg = stock / (BD x depth x 0.1 x aggregate % / 100), live links to the BD and aggregate rows) and summed (macro + micro) over layers to an aggregate-based total "
     "C stock per cumulative depth; both flagged as ALTERNATIVE estimates to the bulk-soil rows (do not pool both).",
     "48 (companion) Mondal 2021 Fig. 5 (TA 0-7.5 cm 7.39 / 6.25 g/kg vs Fig. 3 class means 7.52 / 6.35; 0-7.5 cm aggregate total 6.95 vs bulk 7.12 Mg/ha).",
     "Author 2026-10-05", "IN FORCE (author)"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("new rules", n0 + 1, "-", n0 + len(ROWS))
