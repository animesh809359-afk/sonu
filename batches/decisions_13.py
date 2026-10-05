"""Batch 13 rules (updated 13): rules 103-107 with the author decisions of 2026-10-05; rule 25 narrowed."""
import sys
from copy import copy

import openpyxl

p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 1).value == 25:
        ws.cell(r, 6).value = "IN FORCE - narrowed by rule %d (conventionally dry-tilled rice + ZT wheat = partial codes)" % (n0 + 1)
ROWS = [
    ("Treatment codes",
     "CONVENTIONALLY TILLED rice + ZT wheat = PARTIAL code, whether the rice is puddled or not: conventional DRY tillage (harrowing / tyne / plough passes) before direct-seeded or "
     "water-seeded unpuddled rice + zero-till or surface-seeded wheat -> pZT (no residue) / pCA (residue). Rule 25 (ZT / CA) is kept for minimally tilled / untilled unpuddled "
     "rice (e.g. UPTR, unpuddled mechanical transplanting).",
     "174 Gathala 2017: T1 / T3 pZT, T2 / T4 pCA (rows a/b = puddled TPR / dry-tilled DSR), T5 ZT, T6 CA. 176 / 176 (companion) Devkota 2015: WSRF-FI pZT (rows a-d), WSRF-AWD "
     "second pZT treatment (row e, 2009). Old-master entries under rule 25 not changed (rule 1).", "Author 2026-10-05", "IN FORCE (author)"),
    ("Treatment codes",
     "Residue retention is coded only when the paper itself states it: a companion paper that does not describe residue is coded WITHOUT residue (ZT / pZT / CT), even when "
     "the main paper of the same trial reports residue retention (rule 89 transfers tillage descriptions only).",
     "115 (companion) Pokharel 2018: CTTPR+CTW CT, CTTPR+ZTW pZT, ZTDSR+ZTW / UPTPR+ZTW ZT rows a/b (old-master 115 coded them CA).", "Author 2026-10-05", "IN FORCE (author)"),
    ("Parameter-specific",
     "B:C not defined in the paper: treated as gross / cost and converted (printed - 1, rule 57) when the printed numbers imply it (net/cost would give a total cost below the "
     "itemised costs); flagged on the row.",
     "175 Hossain 2022 (rice TA: gross 1133 US$/ha, BCR 1.07 -> cost 1059 as gross/cost vs 547 as net/cost < tillage + weeding 455 + fertiliser).", "Author 2026-10-05",
     "IN FORCE (author confirmed)"),
    ("Parameter-specific",
     "Crop-model papers: MEASURED values printed only as points of a measured-vs-simulated plot are entered as measured data (read from the measured axis and checked against the "
     "printed measured means); the paper's simulated long-term means are entered as separate MODEL OUTPUT rows (rule 55); simulation-only scenarios never tested in the field "
     "are not entered.",
     "176 (companion) Devkota 2015 AFM: Figs 3A / 4 measured yields 2008-2009 (means reproduce Tables 4-5 within 3 kg/ha); Tables 6-7 DSSAT 1971-2010 yields, N uptake, "
     "N leached; scenarios 9-11 excluded.", "Default 2026-10-05", "IN FORCE - awaiting author confirmation"),
    ("Site classification",
     "A printed site value that is clearly a typographical error is corrected on the author's instruction and the correction noted on the row.",
     "175 Hossain 2022: 'average annual rainfall 172 mm' -> 1720 mm (author 2026-10-05).", "Author 2026-10-05", "IN FORCE (author)"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("rule 25 narrowed; new rules", n0 + 1, "-", n0 + len(ROWS))
