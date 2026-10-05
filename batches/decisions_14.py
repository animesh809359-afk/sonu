"""Batch 14 defaults (updated 14): new rules 108-110 (awaiting author confirmation)."""
import sys
from copy import copy

import openpyxl

p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
ROWS = [
    ("Treatment codes",
     "A new companion paper of an old-master trial is coded like the trial's existing AUTHOR-CONFIRMED companions, with partial codes added for treatments the old entries "
     "excluded; when a later rule would code it differently, the row is flagged and the question raised.",
     "59 (companion 3) Singh G. 2022: DSR-ZTW (+BM / +MBR) ZT, + rice residue CA (as 59 companions), TPR-ZTW pZT; question - DSR was dry-tilled in years 1-3 (rule 103 -> pZT / pCA?).",
     "Default 2026-10-05", "IN FORCE - awaiting author confirmation"),
    ("Workflow & data management",
     "A paper that contradicts itself on a coding fact (e.g. residue 'removed' in the treatment table but tonnes retained in the results table) has that treatment left out "
     "and its values written into the Notes until the author decides (rules 3, 104).",
     "48 (companion) Mondal 2021 pCA1 (NT wheat after CT rice): Table 1 removed vs Table 5 2.0-2.4 Mg/ha retained -> pZT or pCA?", "Default 2026-10-05", "PENDING - author decision"),
    ("Parameter-specific",
     "Aggregate-associated C printed as AMOUNTS (g C per kg bulk soil = concentration x class mass fraction) is converted to a concentration in the aggregate class: "
     "macro C = (sum of >0.25 mm amounts) / (sum of >0.25 mm mass fractions); micro C = 0.053-0.25 mm amount / its mass fraction (DERIVED, flagged).",
     "177 Zhao 2021 Table 2 (SOC and KMnO4-LOC amounts) with the Fig. 1 class masses.", "Default 2026-10-05", "IN FORCE - awaiting author confirmation"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("new rules", n0 + 1, "-", n0 + len(ROWS))
