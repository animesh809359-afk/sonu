"""Loader for the transposed ASD export workbooks (one column per spectrum,
descriptor rows on top, one row per wavelength).

Rules (agreed with the data owner):
* columns whose header is a number are plot-average formulas (=AVERAGE) and
  are ignored;
* a second data block (a repeated "Wavelength" column) is accepted only if
  every spectrum in it is an exact copy of a spectrum of the same plot in the
  first block - then it is dropped as a duplicate (13-03-2026 file);
* spectra without a file name are kept and given an ID (flagged in the audit).
"""
from __future__ import annotations

import re

import numpy as np
import openpyxl
import pandas as pd


def read_file(path, date_label):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=False).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    k = 0
    while not isinstance(rows[k][0], (int, float)):
        k += 1
    desc_names = [str(rows[r][0]).strip() for r in range(1, k)]
    wl = np.array([rows[r][0] for r in range(k, len(rows))], float)
    hdr = rows[0]
    log = {"file": str(path), "date": date_label, "descriptor_rows": desc_names,
           "n_wavelengths": len(wl), "avg_columns": 0, "blocks": 1,
           "unnamed": [], "block2": []}
    block = 1
    recs, spectra = [], []
    for i, h in enumerate(hdr):
        if i == 0:
            continue
        if isinstance(h, str) and h.strip().lower() == "wavelength":
            block += 1
            log["blocks"] = block
            continue
        if h is not None and not isinstance(h, str):
            log["avg_columns"] += 1          # plot-average formula column
            continue
        vals = [rows[r][i] for r in range(k, len(rows))]
        if any(isinstance(v, str) for v in vals):
            raise ValueError(f"{path}: column {i} holds text/formulas but has a spectrum header")
        rec = {desc_names[d]: rows[1 + d][i] for d in range(len(desc_names))}
        rec.update({"Date": date_label, "Source column": openpyxl.utils.get_column_letter(i + 1),
                    "Block": block,
                    "File": h if h is not None else f"(unnamed, {date_label}, column {openpyxl.utils.get_column_letter(i + 1)})",
                    "Unnamed": h is None})
        if h is None:
            log["unnamed"].append(rec["File"])
        recs.append(rec)
        spectra.append(np.asarray(vals, float))
    meta = pd.DataFrame(recs)
    R = np.vstack(spectra)
    return meta, R, wl, log


def load_all(files):
    """files: list of (path, date_label). Returns meta, R, wl, log."""
    metas, Rs, logs, wl0 = [], [], [], None
    for path, date in files:
        meta, R, wl, log = read_file(path, date)
        if wl0 is None:
            wl0 = wl
        elif not np.array_equal(wl, wl0):
            raise ValueError(f"wavelength grid differs in {path}")
        # second block: drop only if every spectrum duplicates one of block 1 (same plot)
        if (meta["Block"] > 1).any():
            b1 = np.where(meta["Block"] == 1)[0]
            b2 = np.where(meta["Block"] > 1)[0]
            matched = []
            for j in b2:
                same = [i for i in b1 if meta.at[i, "Plot no."] == meta.at[j, "Plot no."]
                        and np.array_equal(R[i], R[j])]
                matched.append(same[0] if same else None)
            log["block2"] = [{"Block-2 column": meta.at[j, "Source column"], "File": meta.at[j, "File"],
                              "Plot": meta.at[j, "Plot no."],
                              "Identical to block-1 column": meta.at[m, "Source column"] if m is not None else None,
                              "Block-1 file": meta.at[m, "File"] if m is not None else None}
                             for j, m in zip(b2, matched)]
            if any(m is None for m in matched):
                raise ValueError(f"{path}: block 2 contains spectra not found in block 1 - check manually")
            keep = meta["Block"] == 1
            meta, R = meta[keep].reset_index(drop=True), R[keep.values]
        metas.append(meta)
        Rs.append(R)
        logs.append(log)
    meta = pd.concat(metas, ignore_index=True)
    R = np.vstack(Rs)
    meta = meta.rename(columns={"Plot no.": "Plot", "Healthy/Diseased": "Class",
                                "Growth stage": "Growth stage", "Variety": "Variety"})
    meta["Plot"] = meta["Plot"].astype(int)
    meta["Class"] = meta["Class"].astype(str).str.strip().str.capitalize()
    meta.insert(0, "Spectrum ID", np.arange(1, len(meta) + 1))
    drop = [c for c in meta.columns if c.lower() in ("leaf level", "disease_score")]
    meta = meta.drop(columns=drop)
    return meta, R, wl0, logs


def date_from_name(name):
    m = re.search(r"(\d{2})(\d{2})(\d{2})\.xlsx$", str(name))
    return f"{m.group(1)}-{m.group(2)}-20{m.group(3)}" if m else str(name)
