"""Author instructions 2026-10-05 (updated 11):

1. Water-stable aggregate size classes go to the aggregate sheets, not a separate sheet: classes >0.25 mm = MACRO, water-stable classes <0.25 mm = MICRO,
   MACRO + MICRO = WSA (g/g). Gathala 2011 (15 companion 2) Fig. 2 classes are moved accordingly and the AGG SIZE CLASS sheet is removed; MWD keeps its value
   (formula rewritten with the class values).
2. Obs recomputed for EVERY row of every study: Obs = Y + T + D (factors equal to 1 add nothing; minimum 1), where
   Y = years of data for that parameter / depth in the study (year-wise rows: number of years reported in the row's group; pooled means: years pooled),
   T = paper treatments under one code (rows a/b/c of rule 20, or treatments pooled into one value),
   D = paper depths falling in the row's depth class.
Usage: python batches/fix_11.py <in.xlsx> <out.xlsx>
"""
import re
import sys
from collections import defaultdict

import openpyxl

sys.path.insert(0, __import__("os").path.dirname(__file__))
from lib import Book  # noqa: E402

SRC, OUT = sys.argv[1], sys.argv[2]
B = Book(SRC)
wb = B.wb

# ------------------------------------------------------------------ 1. aggregate classes -> MICRO / WSA
agg = wb["AGG SIZE CLASS"]
ah = {c.value: c.column for c in agg[1] if c.value}
classes = defaultdict(dict)  # row label -> class -> {code: value}
for r in range(2, agg.max_row + 1):
    if agg.cell(r, 2).value is None:
        continue
    lab = re.search(r"ROW ([a-d])\b", str(agg.cell(r, ah["Notes/Doubts"]).value)).group(1)
    cl = agg.cell(r, ah["Aggregate size class (mm)"]).value
    classes[lab][cl] = {k[5:]: agg.cell(r, ah[k]).value for k in ah if k.startswith("AGGC_") and agg.cell(r, ah[k]).value is not None}
MID = {"8.00-4.75": 6.375, "4.75-2.00": 3.375, "2.00-1.00": 1.5, "1.00-0.50": 0.75, "0.50-0.25": 0.375, "0.25-0.11": 0.18, "0.11-0.00": 0.055}
mwd = wb["MWD 1"]
mh = B.headers("MWD 1")
for r in range(2, mwd.max_row + 1):
    if str(mwd.cell(r, 2).value) != "15 (companion 2)":
        continue
    lab = re.search(r"ROW ([a-d])\b", str(mwd.cell(r, mh["Notes/Doubts"]).value)).group(1)
    for k, c in mh.items():
        if k.startswith("MWD_") and mwd.cell(r, c).value is not None:
            code = k[4:]
            mwd.cell(r, c).value = "=ROUND(" + "+".join(f"{MID[cl]}*{classes[lab][cl][code]}/100" for cl in MID) + ",2)"
    n = mh["Notes/Doubts"]
    mwd.cell(r, n).value = str(mwd.cell(r, n).value).replace("(live links to AGG SIZE CLASS)", "(class values written into the formula; classes listed in the MICRO / WSA rows' Notes)")
# MACRO 2008-09 rows of Gathala are the base for MICRO and WSA
mac = wb["MACRO"]
xh = B.headers("MACRO")
macro_rows = {}
for r in range(2, mac.max_row + 1):
    if str(mac.cell(r, 2).value) == "15 (companion 2)" and mac.cell(r, xh["year of data collection/experiment"]).value == "2008-09":
        lab = re.search(r"ROW ([a-d])\b", str(mac.cell(r, xh["Notes/Doubts"]).value)).group(1)
        macro_rows[lab] = r
SKIP = {"SD", "Obs", "Data source", "Notes/Doubts", "UNIT", "Method used (from paper)"}
for lab, mr in sorted(macro_rows.items()):
    base = {k: mac.cell(mr, c).value for k, c in xh.items() if not k.startswith("MACRO_") and k not in SKIP}
    cls = classes[lab]
    clsnote = "; ".join(f"{cl} mm: " + ", ".join(f"{code} {v}" for code, v in cls[cl].items()) for cl in MID)
    note = (f"ROW {lab}. Fig. 2 water-stable aggregate size classes after 7 cycles (digitised, +/-0.5 %): {clsnote}. MICRO = the 0.25-0.11 mm class only - the finest class "
            "(0.11-0.00 mm) also holds silt + clay <0.053 mm and is not a water-stable aggregate fraction, so it is left out (flag: the 0.11-0.053 mm part of the micro-aggregates "
            "is therefore missing). MACRO for this sampling = Table 3 2008-09 row (the Fig. 2 classes >0.25 mm reproduce it within ~1 %). Author instruction 2026-10-05.")
    meth = ("Wet sieving (Yoder 1936) of 4.75-8 mm air-dried aggregates through 8.00, 4.75, 2.00, 1.00, 0.50, 0.25 and 0.11 mm sieves; mass per class")
    codes = list(cls["0.25-0.11"].keys())
    mi = B.add("MICRO", {**base, **{"MICRO_" + c: cls["0.25-0.11"][c] for c in codes}, "Obs": 1,
                         "UNIT": "% of soil (water-stable micro-aggregates 0.25-0.11 mm)", "Data source": "Fig. 2 (digitised, scanned figure)",
                         "Method used (from paper)": meth, "Notes/Doubts": note})
    B.add("WSA", {**base, **{"WSA_" + c: f"=ROUND(({B.ref('MACRO', 'MACRO_' + c, mr)}+{B.ref('MICRO', 'MICRO_' + c, mi)})/100,4)" for c in codes}, "Obs": 1,
                  "UNIT": "g/g soil (MACRO >0.25 mm + MICRO 0.25-0.11 mm, % / 100)", "Data source": "DERIVED (Table 3 MACRO + Fig. 2 MICRO)",
                  "Method used (from paper)": meth + "; WSA = macro + micro (live links)", "Notes/Doubts": "DERIVED (author 2026-10-05): WSA = MACRO + MICRO. " + note})
# remove the AGG SIZE CLASS sheet and its index row
del wb["AGG SIZE CLASS"]
sp = wb["STUDIES_BY_PARAMETER"]
for r in range(sp.max_row, 3, -1):
    if sp.cell(r, 2).value == "AGG SIZE CLASS":
        sp.delete_rows(r)
for i, r in enumerate(range(4, sp.max_row + 1), 1):
    if sp.cell(r, 2).value is not None:
        sp.cell(r, 1).value = i

# ------------------------------------------------------------------ 2. Obs = Y + T + D for every row
T_STUDY = {"15 (companion 2)": 2, "165 (companion)": 2, "40 (companion)": 2, "172": 3, "173": 2, "170": 2, "171": 2}
T_REASON = {"15 (companion 2)": "rows a/b (2 treatments per code)", "165 (companion)": "rows a/b (+/- bio-fertiliser)", "40 (companion)": "LR / HR residue levels (rows a/b or pooled)",
            "172": "rows a/b/c (3 residue treatments under CTR)", "173": "2 CA treatments per fertiliser level (rows a/c, b/d)", "170": "medium + intensive puddling pooled",
            "171": "N1 + N2 timing pooled"}
CUMUL = {"stock-SOC", "c sequestration rate"}
YR = re.compile(r"\b(1[89]|20)\d{2}(-\d{2,4})?\b")


def pooled(y):
    y = str(y)
    for pat in (r"\((\d+)-yr mean\)", r"(\d+)-season mean", r"(\d+)-yr trial"):
        m = re.search(pat, y)
        if m:
            return int(m.group(1))
    if "2009-2015 (mean)" in y:
        return 7
    if y in ("1993-96", "2008-11"):
        return 3
    return None


def norm(stage):
    s = str(stage or "")
    s = re.sub(r"\(LTE year \d+\)", "", s)
    s = re.sub(r"\(\d+(\.\d+)? years?\)", "", s)
    s = YR.sub("#", s)
    s = s.replace("one season # to # (unlabelled)", "#")
    return re.sub(r"\s+", " ", s).strip()


names = wb.sheetnames
log = defaultdict(int)
for n in names[names.index("BD"):]:
    if n == "EXCLUDED_rows":
        continue
    ws = wb[n]
    h = B.headers(n)
    if "Obs" not in h:
        continue
    g = lambda r, k: ws.cell(r, h[k]).value if k in h else None
    rows = []
    for r in range(2, ws.max_row + 1):
        s = ws.cell(r, 2).value
        if s is None or ws.cell(r, 3).value is None:
            continue
        note = str(g(r, "Notes/Doubts") or "")
        m = re.search(r"ROW ([a-d])\b", note)
        rows.append(dict(r=r, s=str(s), site=g(r, "Site/Location"), dc=g(r, "DEPTH"), dr=g(r, "DEPTH (as reported in paper)"), yr=g(r, "year of data collection/experiment"),
                         st=norm(g(r, "Crop/season of sampling")), tm=g(r, "Treatment mapping (paper's name -> code)"), un=g(r, "UNIT"), lab=m.group(1) if m else "-"))
    ygroups, dgroups = defaultdict(set), defaultdict(set)
    for x in rows:
        ygroups[(x["s"], x["site"], x["dc"], x["dr"], x["st"], x["tm"], x["un"], x["lab"])].add(str(x["yr"]))
        dgroups[(x["s"], x["site"], x["dc"], str(x["yr"]), x["st"], x["tm"], x["un"], x["lab"])].add(str(x["dr"]))
    for x in rows:
        p = pooled(x["yr"])
        Y = p if p else len(ygroups[(x["s"], x["site"], x["dc"], x["dr"], x["st"], x["tm"], x["un"], x["lab"])])
        if x["dc"] is None or n in CUMUL:
            D = 1
        else:
            mm = re.search(r"mean of (\d+) layers", str(x["dr"]))
            D = int(mm.group(1)) if mm else len(dgroups[(x["s"], x["site"], x["dc"], str(x["yr"]), x["st"], x["tm"], x["un"], x["lab"])])
        T = T_STUDY.get(x["s"], 1)
        if x["s"] == "15 (companion 2)" and x["lab"] == "-":
            T = 1  # soil temperature: only T1 / T3 / T5 plotted, no row b
        obs = sum(f for f in (Y, T, D) if f > 1) or 1
        r = x["r"]
        old = ws.cell(r, h["Obs"]).value
        ws.cell(r, h["Obs"]).value = obs
        why = (f" [OBS RECOMPUTED 2026-10-05 (author rule): Y={Y}" + (" (pooled mean)" if p else " (years of data in this row's series)") + f" + T={T}"
               + (f" ({T_REASON[x['s']]})" if T > 1 else "") + f" + D={D} -> Obs = {obs}; factors of 1 add nothing; supersedes any earlier Obs statement.]")
        c = ws.cell(r, h["Notes/Doubts"])
        c.value = (str(c.value) if c.value else "") + why
        log["changed" if old != obs else "same"] += 1
B.save(OUT)
print("MICRO/WSA rows added for", sorted(macro_rows), "| Obs:", dict(log))
