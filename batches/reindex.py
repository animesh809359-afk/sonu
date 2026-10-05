"""Regenerate PARAMETERS_BY_STUDY and STUDIES_BY_PARAMETER from the data sheets.

Usage: python batches/reindex.py <workbook.xlsx> [readme line]
"""
import sys
from collections import OrderedDict, defaultdict
from copy import copy

import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path)
names = wb.sheetnames
data_sheets = [n for n in names[names.index("BD"):] if n != "EXCLUDED_rows"]

rows = defaultdict(lambda: OrderedDict())   # serial -> {sheet: n}
sheet_rows = OrderedDict()
for n in data_sheets:
    ws = wb[n]
    cnt = OrderedDict()
    for r in range(2, ws.max_row + 1):
        s = ws.cell(r, 2).value
        if s is None or ws.cell(r, 3).value is None:
            continue
        cnt[s] = cnt.get(s, 0) + 1
        rows[s][n] = rows[s].get(n, 0) + 1
    sheet_rows[n] = cnt

si = wb["Study_Info"]
h = {c.value: c.column for c in si[1]}
studies = []
for r in range(2, si.max_row + 1):
    if si.cell(r, h["SERIAL NO"]).value is None:
        continue
    studies.append({k: si.cell(r, h[k]).value for k in ("SERIAL NO", "Authors", "Year", "Journal", "Parameters extracted", "Notes/Doubts")})

ws = wb["PARAMETERS_BY_STUDY"]
style = [copy(c._style) for c in ws[3]]
if ws.max_row >= 4:
    ws.delete_rows(4, ws.max_row - 3)
for i, s in enumerate(studies, 1):
    sp = rows.get(s["SERIAL NO"], {})
    status = "EXCLUDED" if str(s["Notes/Doubts"] or "").startswith("EXCLUDED") else "INCLUDED"
    vals = [i, s["SERIAL NO"], s["Authors"], s["Year"], s["Journal"], status, sum(sp.values()), len(sp),
            "; ".join(f"{k}: {v}" for k, v in sp.items()) or "-", s["Parameters extracted"]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(3 + i, j, v)
        c.font = openpyxl.styles.Font(name="Arial", size=10)

ws = wb["STUDIES_BY_PARAMETER"]
for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 2).value
    if name in sheet_rows:
        cnt = sheet_rows[name]
        ws.cell(r, 4, sum(cnt.values()))
        ws.cell(r, 5, len(cnt))
        ws.cell(r, 6, "; ".join(f"{k}: {v}" for k, v in cnt.items()) or "NO DATA")

if len(sys.argv) > 2:
    rd = wb["README"]
    r = rd.max_row + 1
    label, text = sys.argv[2].split("|", 1)
    rd.cell(r, 1, label).font = openpyxl.styles.Font(name="Arial", size=10, bold=True)
    c = rd.cell(r, 2, text)
    c.font = openpyxl.styles.Font(name="Arial", size=10)
    c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

wb.calculation.fullCalcOnLoad = True
wb.save(path)
print("reindexed", len(studies), "studies;", sum(1 for v in sheet_rows.values() if v), "sheets with data")
