"""Append the batch-04 decision log to RULES."""
import sys
from copy import copy
import openpyxl
p = sys.argv[1]
wb = openpyxl.load_workbook(p)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
rows = [
    ("Workflow & data management", "A paper EXCLUDED in the old master only because of a rotation mismatch that the partial-CA codes now cover is RE-ENTERED as 'NN (companion)' with the pCA/pZT/pMT/pMTR codes.",
     "Zhang 2023 = old-master 131 -> 131 (companion): PR (puddled rice + NT wheat + straw) = pCA, CN = CA.", "Default 2026-10-05 (rules 6, 68)", "PENDING CONFIRMATION"),
    ("Treatment codes", "Old-master trials whose partial-CA scenario was excluded contribute that scenario as pCA in new companion entries.",
     "CSSRI Karnal scenario trial (old-master 23): Sc2 (puddled TPR + ZT wheat + ZT mungbean, residue) -> pCA in 23 (companion 3) and 23 (companion 4).", "Default 2026-10-05 (rule 68)", "PENDING CONFIRMATION"),
    ("Parameter-specific", "Carbon pools printed only as stocks (Mg C/ha) with no treatment BD: concentration = pool stock / SOC stock x measured TOC concentration of the same sampling (DERIVED, flagged).",
     "23 (companion 4) VLC/LC/LLC/NLC; TOC from 23 (companion 3), Oct 2013.", "Default 2026-10-05 (rule 53)", "PENDING CONFIRMATION"),
    ("Parameter-specific", "Aggregate-associated C printed per size class without class masses: macro-C = unweighted mean of the >0.25 mm classes, micro-C = mean of the 0.053-0.25 mm classes (flagged).",
     "23 (companion 3) Figs 2-3.", "Default 2026-10-05", "PENDING CONFIRMATION"),
    ("Cropping system & design", "Organic-manure trials: manure only (CT) vs manure + crop residue (CTR); a residue + biofertiliser treatment is a second CTR row (row b, flagged). Tillage taken from a companion paper of the same project when the paper does not describe it (flagged).",
     "165 (companion) Meena 2020 (IARI organic project of Davari 2012): FYM / VC vs + CR (row a) and + CR + BF (row b).", "Default 2026-10-05 (rules 5, 20)", "PENDING CONFIRMATION"),
]
r = ws.max_row + 1
for k, (cat, rule, det, src, st) in enumerate(rows):
    for j, v in enumerate((n0 + 1 + k, cat, rule, det, src, st), 1):
        ws.cell(r + k, j, v)._style = copy(style[j - 1])
wb.calculation.fullCalcOnLoad = True
wb.save(p)
print("rules", n0 + 1, "-", n0 + len(rows))
