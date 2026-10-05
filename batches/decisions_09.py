"""Author 2026-10-05: rules 85, 86, 88, 89 confirmed; SOIL TEMP sheet confirmed (rule 96). No data values change."""
import sys
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
done = []
for r in range(1, ws.max_row + 1):
    n = ws.cell(r, 1).value
    if n in (85, 86, 88, 89):
        ws.cell(r, 6).value = "IN FORCE (author confirmed 2026-10-05)"
        done.append(n)
    if n == 96:
        ws.cell(r, 4).value = ("New sheets 2026-10-05: LLWR (author request; % v/v); SOIL TEMP (deg C, one row per week / reading date and time of day - author confirmed "
                               "2026-10-05); AGG SIZE CLASS (% per water-stable size class) created by default for Gathala 2011 Fig. 2 - PENDING CONFIRMATION of this sheet only.")
        done.append(n)
assert done == [85, 86, 88, 89, 96], done
rd = wb["README"]
r = rd.max_row + 1
rd.cell(r, 1, "Author confirmations (updated 09, 2026-10-05)").font = openpyxl.styles.Font(name="Arial", size=10, bold=True)
c = rd.cell(r, 2, "updated 09: rules 85 (re-entry of rotation-excluded old-master papers as companions), 86 (old-master partial-CA scenarios as pCA), "
                  "88 (unweighted macro/micro aggregate C) and 89 (organic-manure trials: manure CT vs manure + residue CTR, biofertiliser row b) confirmed by the author; "
                  "SOIL TEMP sheet confirmed (rule 96). No data values changed.")
c.font = openpyxl.styles.Font(name="Arial", size=10)
c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("updated", done)
