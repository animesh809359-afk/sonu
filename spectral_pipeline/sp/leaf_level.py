"""Leaf-level files: one sheet per date and class ('110226(Healthy)', '030326(diseased)', ...), first column
Wavelength, one column per leaf spectrum whose header is the leaf level (L3, L5, LL). The sheets carry no
plot numbers, so every spectrum is matched exactly to the labelled ASD export workbooks, which supply plot,
variety and growth stage (needed for plot-grouped validation)."""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DATE_RE = re.compile(r"(\d{2})(\d{2})(\d{2})")


def read_leaf_files(paths):
    rows, specs, wl0, inv = [], [], None, []
    for p in paths:
        wb = load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            vals = [r for r in ws.iter_rows(values_only=True) if r and r[0] is not None]
            hdr = vals[0]
            wl = np.array([float(r[0]) for r in vals[1:]])
            X = np.array([[float(v) for v in r[1:len(hdr)]] for r in vals[1:]], float).T
            if wl0 is None:
                wl0 = wl
            elif not np.array_equal(wl, wl0):
                raise ValueError(f"{p}:{ws.title} wavelength grid differs")
            m = DATE_RE.search(ws.title)
            date = f"{m.group(1)}-{m.group(2)}-20{m.group(3)}"
            cls = "Healthy" if "healthy" in ws.title.lower() else "Diseased" if "diseased" in ws.title.lower() else None
            levels = [str(h).strip() for h in hdr[1:]]
            for j in range(X.shape[0]):
                rows.append({"Leaf-level file": Path(p).name, "Sheet": ws.title, "Column": get_column_letter(j + 2),
                             "Leaf level": levels[j], "Date": date, "Class": cls})
                specs.append(X[j])
            inv.append({"File": Path(p).name, "Sheet": ws.title, "Date": date, "Class": cls, "Spectra": X.shape[0],
                        **{lv: levels.count(lv) for lv in sorted(set(levels))}})
    return pd.DataFrame(rows), np.array(specs), wl0, pd.DataFrame(inv).fillna(0)


def attach(meta_all, R_all, wl, paths):
    """Adds 'Leaf level' to meta_all by exact spectral match. Returns (meta_all, R_all) restricted to spectra present
    in the leaf-level files, plus a report dict (inventory, match table, unmatched, duplicates, checks)."""
    lm, LR, lwl, inv = read_leaf_files(paths)
    if not np.array_equal(lwl, wl):
        raise ValueError("leaf-level files and labelled files have different wavelength grids")
    index = {}
    for i, v in enumerate(R_all):
        index.setdefault(v.tobytes(), []).append(i)
    matches = [index.get(v.tobytes(), []) for v in LR]
    lm["Matches in labelled files"] = [len(m) for m in matches]
    lm["Matched file"] = [";".join(str(meta_all.at[i, "File"]) for i in m) for m in matches]
    lm["Plot"] = [int(meta_all.at[m[0], "Plot"]) if m else None for m in matches]
    lm["Variety"] = [meta_all.at[m[0], "Variety"] if m else None for m in matches]
    lm["Growth stage"] = [meta_all.at[m[0], "Growth stage"] if m else None for m in matches]
    lm["Class agrees"] = [bool(m) and all(meta_all.at[i, "Class"] == c for i in m) for m, c in zip(matches, lm["Class"])]
    lm["Date agrees"] = [bool(m) and all(meta_all.at[i, "Date"] == d for i in m) for m, d in zip(matches, lm["Date"])]
    level = pd.Series([None] * len(meta_all), dtype=object)
    conflict = 0
    for m, lv in zip(matches, lm["Leaf level"]):
        for i in m:
            if level[i] is not None and level[i] != lv:
                conflict += 1
            level[i] = lv
    meta_all = meta_all.copy()
    meta_all["Leaf level"] = level.values
    in_files = level.notna().to_numpy()
    unmatched_old = meta_all.loc[~in_files, ["Date", "Plot", "Variety", "Growth stage", "Class", "File", "Source column"]].copy()
    unmatched_old["Unnamed"] = meta_all.loc[~in_files, "Unnamed"].values
    # duplicate spectra inside the leaf-level files
    seen, dups = {}, []
    for k, v in enumerate(LR):
        key = v.tobytes()
        if key in seen:
            a = seen[key]
            dups.append({"Sheet": lm.at[k, "Sheet"], "Column": lm.at[k, "Column"], "Leaf level": lm.at[k, "Leaf level"],
                         "Identical to column": lm.at[a, "Column"], "Its leaf level": lm.at[a, "Leaf level"],
                         "Plot": lm.at[k, "Plot"], "Labelled file": lm.at[k, "Matched file"].split(";")[0]})
        else:
            seen[key] = k
    n_unm = int((lm["Matches in labelled files"] == 0).sum())
    checks = [
        ("Leaf-level files: every spectrum found exactly (|Δ| = 0 at all 2151 bands) in the labelled ASD files",
         f"{len(lm) - n_unm} of {len(lm)} spectra matched", "PASS" if n_unm == 0 else "FAIL"),
        ("Healthy/diseased label of the leaf-level sheet equals the label in the labelled files",
         f"{int(lm['Class agrees'].sum())} of {len(lm)}", "PASS" if lm["Class agrees"].all() else "FAIL"),
        ("Date of the leaf-level sheet equals the date of the labelled file",
         f"{int(lm['Date agrees'].sum())} of {len(lm)}", "PASS" if lm["Date agrees"].all() else "FAIL"),
        ("One leaf level per spectrum (no spectrum labelled with two levels)", f"{conflict} conflicts",
         "PASS" if conflict == 0 else "FAIL"),
        ("Identical spectra inside the leaf-level files", f"{len(dups)} pairs (the plot-27 copies of 11-02-2026) → one copy kept",
         "FLAG" if dups else "PASS"),
        ("Spectra of the labelled files that are not in the leaf-level files",
         "; ".join(f"{r.File} (plot {r.Plot}, {r.Date}, {r.Class})" for r in unmatched_old.itertuples() if not r.Unnamed) or "none",
         "FLAG" if (~unmatched_old["Unnamed"]).any() else "PASS"),
    ]
    rep = {"inventory": inv, "match": lm, "unmatched": unmatched_old, "dups": pd.DataFrame(dups), "checks": checks,
           "n_files": len(lm)}
    keep = in_files
    return meta_all[keep].reset_index(drop=True), R_all[keep], rep
