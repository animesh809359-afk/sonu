"""Step 1 - assemble the plot x date table from the four PlotAverage_Features workbooks and the
treatment layout. Usage: python -I p01_load.py <treatment.xlsx> <workbook_1.xlsx> ... <workbook_4.xlsx>"""
import os, sys, re, json
import numpy as np
import pandas as pd
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pcommon import *

SHEETS = {'raw': {'REF': '02', 'CR': '04', 'VI': '06', 'FD1': '07', 'FD2': '08', 'FD3': '09', 'RE': '10', 'OPT': '12'},
          'smooth': {'REF': '14', 'CR': '15', 'VI': '17', 'FD1': '18', 'FD2': '19', 'FD3': '20', 'RE': '21', 'OPT': '23'}}
WAVE = {'REF', 'FD1', 'FD2', 'FD3'}


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def is_int(v):
    try:
        return float(v) == int(float(v))
    except (TypeError, ValueError):
        return False


def sheet(wb, prefix):
    return wb[[s for s in wb.sheetnames if s.startswith(prefix + '_')][0]]


def read_wave(ws):
    plots, wl, rows, started = None, [], [], False
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if plots is None and a == 'Plot':
            plots = [int(float(v)) for v in r[1:] if is_int(v)]
            continue
        if plots is not None and isinstance(a, str) and a.startswith('Wavelength'):
            started = True
            continue
        if started and is_int(a):
            wl.append(int(float(a)))
            rows.append([num(v) for v in r[1:1 + len(plots)]])
    return np.array(plots), np.array(wl), np.array(rows).T  # plots x wavelengths


def read_table(ws):
    head, out = None, []
    for r in ws.iter_rows(values_only=True):
        if head is None:
            if r[0] == 'Plot' and len(r) > 1 and r[1] == 'Variety':
                head = [h for h in r]
            continue
        if not is_int(r[0]):
            if out:
                break
            continue
        out.append(r[:len(head)])
    ncol = max(i for i, h in enumerate(head) if h is not None) + 1
    df = pd.DataFrame([list(x[:ncol]) for x in out], columns=[str(h) for h in head[:ncol]])
    return df


def read_info(ws):
    rows, head = [], None
    for r in ws.iter_rows(values_only=True):
        if head is None:
            if r[0] == 'Plot' and r[1] == 'Variety':
                head = r
            continue
        if not is_int(r[0]):
            break
        rows.append(dict(plot=int(float(r[0])), variety=r[1], stage=r[2], label=r[3], n_leaves=int(float(r[4])),
                         files=r[5]))
    return pd.DataFrame(rows)


def main(treat_path, books):
    T = pd.read_excel(treat_path)
    T.columns = ['plot', 'inoc', 'variety_t', 'block']
    T['block'] = T['block'].replace({'Control/healthy': 'C'})
    metas, store = [], {b: {} for b in BRANCHES}
    spectra = {b: [] for b in BRANCHES}
    defs = None
    for path in books:
        dd = re.search(r'(\d\d)-(\d\d)-(\d{4})', os.path.basename(path)).groups()
        date = f'{dd[2]}-{dd[1]}-{dd[0]}'
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        info = read_info(sheet(wb, '01'))
        info['date'] = date
        metas.append(info)
        for b in BRANCHES:
            for fam, pre in SHEETS[b].items():
                ws = sheet(wb, pre)
                if fam in WAVE:
                    plots, wl, X = read_wave(ws)
                    assert list(plots) == list(info['plot']), (path, fam)
                    df = pd.DataFrame(X, columns=wl)
                else:
                    df = read_table(ws)
                    assert list(df['Plot'].astype(float).astype(int)) == list(info['plot']), (path, fam)
                    df = df.iloc[:, 5:].apply(pd.to_numeric, errors='coerce')
                df.insert(0, 'date', date); df.insert(1, 'plot', info['plot'].to_numpy())
                store[b].setdefault(fam, []).append(df)
        if defs is None:
            ws = sheet(wb, '26')
            rows = [r for r in ws.iter_rows(min_row=4, values_only=True)]
            defs = pd.DataFrame([r[:6] for r in rows[1:] if r[0]], columns=[str(c) for c in rows[0][:6]])
        wb.close()
        print('read', date, len(info), 'plots')
    M = pd.concat(metas, ignore_index=True)
    M = M.merge(T, on='plot', how='left')
    assert (M['variety'] == M['variety_t']).all(), 'variety mismatch with treatment file'
    assert ((M['label'] == 'Healthy') == (M['block'] == 'C')).all(), 'class mismatch with treatment file'
    M = M.drop(columns='variety_t')
    M['y'] = (M['label'] == 'Diseased').astype(int)
    M['cell'] = M['variety'] + '-' + M['inoc']
    M['date'] = pd.Categorical(M['date'], DATE_ORDER, ordered=True)
    order = M.sort_values(['date', 'plot']).index
    M = M.loc[order].reset_index(drop=True)
    M['date'] = M['date'].astype(str)
    M['pd'] = M['plot'].astype(str) + '_' + M['date']
    M.drop(columns='files').to_csv(os.path.join(CACHE, 'meta.csv'), index=False)
    key = M[['date', 'plot']]
    for b in BRANCHES:
        for fam, parts in store[b].items():
            df = pd.concat(parts, ignore_index=True)
            df = key.merge(df, on=['date', 'plot'], how='left')
            vals = df.drop(columns=['date', 'plot'])
            if fam in WAVE:
                wl = vals.columns.to_numpy(int)
                X = vals.to_numpy(float)
                if fam == 'REF':
                    np.savez_compressed(os.path.join(CACHE, f'{b}_spectra.npz'), R1=X, wl=wl)
                keep = (wl >= 400) & (wl <= 2400)
                B = bin10(X[:, keep], wl[keep])
                pref = {'REF': 'R', 'FD1': 'D1_', 'FD2': 'D2_', 'FD3': 'D3_'}[fam]
                out = pd.DataFrame(B, columns=[f'{pref}{w}' for w in WL10])
            else:
                out = vals.reset_index(drop=True)
                if fam == 'CR':
                    out.columns = [c.replace(' absorption ', ' ').replace(': ', ' ') for c in out.columns]
            out.to_pickle(os.path.join(CACHE, f'{b}_{fam}.pkl'))
            print(b, fam, out.shape, 'NaN:', int(out.isna().sum().sum()))
    defs.to_csv(os.path.join(CACHE, 'vi_definitions.csv'), index=False)
    T.to_csv(os.path.join(CACHE, 'treatment.csv'), index=False)
    json.dump(dict(treatment=os.path.basename(treat_path), books=[os.path.basename(p) for p in books]),
              open(os.path.join(CACHE, 'inputs.json'), 'w'))
    print(M.groupby(['date', 'label']).size().unstack())


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2:])
