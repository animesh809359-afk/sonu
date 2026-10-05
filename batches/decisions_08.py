"""Rules for batch 07/08 (2026-10-05) with the author's answers: method column, puddling coding (Kukal), bed/ZT coding (Gathala), factorial two-way tables,
duplicate means, all data points + new sheets. Pending rules 85, 86, 88, 89 are untouched."""
import sys
from copy import copy
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
ROWS = [
    ("Workflow & data management",
     "METHODOLOGY: every data row records the paper's method for that parameter in 'Method used (from paper)' (instrument / extractant / procedure and its citation, "
     "sampling depth and timing where given); 'Not stated in paper' when the paper gives none; derived values state the formula.",
     "Checked 2026-10-05: all 495 rows of updated 06 already carried a method; every batch 07/08 row filled the same way.", "Author 2026-10-05", "IN FORCE"),
    ("Treatment codes",
     "PUDDLING DEPTH / INTENSITY trials: unpuddled rice -> ZT; shallow (reduced) puddling -> MT; normal-depth puddling (any intensity) -> CT - also when the wheat phase is "
     "conventionally tilled in every treatment (the contrast is the rice-phase tillage; flagged).",
     "170 Kukal & Aggarwal 2003 (unpuddled plots outside the layout - flag). Old-master 72/73 ReP / DSWP-CTW exclusions are not changed (rule 1) unless the author asks.",
     "Author 2026-10-05", "IN FORCE"),
    ("Treatment codes",
     "Permanent raised beds RESHAPED at every wheat sowing -> MT; zero-till transplanting into unpuddled soil (plots flooded 1 d before) + ZT wheat -> ZT; puddled AWD rice + ZT wheat "
     "with only incidental stubble -> pZT. Several treatments under one code -> rows a/b paired by rice establishment (a = DSR, b = TPR).",
     "15 (companion 2) Gathala 2011 SSSAJ: T1 CT, T2 pZT, T3/T4 MT (rows a/b), T5/T6 ZT (rows a/b). Refines rule 15 for reshaped permanent beds.", "Author 2026-10-05", "IN FORCE"),
    ("Pooled means",
     "Tillage x residue x N factorial printed only as two-way tables: soil properties from the tillage x residue cells (rows a/b, pooled over N, flagged); yields and NUE from the tillage x N cells "
     "(one row per N rate, pooled over residue, flagged); economics from the tillage main effect. The other two-way table goes to Notes - never both.",
     "40 (companion) Kader 2022: CTR vs MTR.", "Author 2026-10-05", "IN FORCE (author confirmed)"),
    ("Workflow & data management",
     "A multi-year MEAN that summarises year-wise values already entered for an old-master study is KEPT, with a 'DUPLICATE FLAG' in Notes.",
     "15 (companion 2) Table 6 7-yr mean yields (T1/T5 year-wise yields are in old-master 15).", "Author 2026-10-05", "IN FORCE (author: keep)"),
    ("Workflow & data management",
     "ALL DATA POINTS: every value printed in a table or plotted in a figure is extracted - each depth / layer, each crop stage or week, each size class - as its own row; figures are "
     "digitised (vector exactly, scans with calibrated grids). Where no sheet exists a new parameter sheet is created with the standard layout (code groups incl. pCA/pZT/pMT/pMTR).",
     "New sheets 2026-10-05: LLWR (author request; % v/v); SOIL TEMP (deg C, one row per week/time) and AGG SIZE CLASS (% per water-stable size class) created by default for "
     "Gathala 2011 Figs 3 and 2 - PENDING CONFIRMATION of these two sheets.", "Author 2026-10-05", "IN FORCE"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("rules", n0 + 1, "-", n0 + len(ROWS))
