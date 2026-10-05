"""Batch 17 rules (updated 17): rules 113-121 confirmed by the author 2026-10-05; new rules 122-123 (author decisions on Mann et al. 2008)."""
import sys
from copy import copy

import openpyxl

p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
D = "Author 2026-10-05"
ROWS = [
    ("Workflow & data management",
     "A re-submitted paper that is already in the NEW workbook is entered again under rule 6 by COPYING every row of the original study (values, notes, styles) to "
     "'NN (companion)', with live formulas re-pointed to the copied rows and 'DUPLICATE RE-SUBMISSION' at the start of Notes; the original rows are not changed.",
     "173 (companion) = Biswakarma 2023 (127_real.pdf = 121_now.pdf), 212 rows. Delete the companion rows if the upload was a mistake.", D, "IN FORCE (author confirmed)"),
    ("Parameter-specific",
     "Laboratory C-mineralisation incubations: cumulative CO2 at the last day -> RESP as the mean daily rate over the incubation (mg CO2/kg/day; mg C x 44/12; "
     "per 100 g x 10). Incubations of aggregate fractions are separate rows flagged by fraction; temperatures are separate rows; the time series and Q10 / Ea / "
     "decay constants (no sheet) go to Notes.",
     "159 (companion) Dutta 2023 Figs 1-6 (59 d, 27 / 37 C); 179 Ahmad Figs S4-S5 (60 d, >2 / <2 mm).", D, "IN FORCE (author confirmed)"),
    ("Parameter-specific",
     "Glomalin measured within aggregate fractions goes on the GLOMALIN sheet as separate rows per fraction (flagged 'WITHIN MACRO- / MICRO-AGGREGATES').",
     "159 (companion) Dutta 2023 Fig. 7 (EEG, mg/g aggregate).", D, "IN FORCE (author confirmed)"),
    ("Treatment codes",
     "Residue BURNING = no residue retained: coded like residue removal (CT / ZT / pZT / pMT ...) as an extra row (row b) next to the residue-removed row.",
     "159 (companion) Dutta 2023 CT-RB / ZT-RB / ST-RB (old-master 60, 63, 88, 94, 97 coded burnt = CT / ZT).", D, "IN FORCE (author confirmed)"),
    ("Statistics (Obs, Rep, SD)",
     "On-farm paired-field trials (one treatment field per farm): farms are the replicates (Rep = number of farms); the row value is the mean over farms (the paper's "
     "stated mean, or a live AVERAGE of the farm values when only farm values are printed); farm values in Notes.",
     "180 Mann et al. 2008 (5 farms, ZT vs CT fields).", D, "IN FORCE (author confirmed)"),
    ("Workflow & data management",
     "A supplement uploaded WITHOUT its main paper: the supplement's eligible data are entered under a new serial (study-level fields the supplement does not give are "
     "left blank and flagged) and the main paper is added later: when the author supplies it, the study-level fields (year, journal, soil, duration, methods) are completed and any further data entered.",
     "179 Ahmad N. et al. (127.docx: Tables S1-S3, Figs S2-S9) - author 2026-10-05: extract now, main paper to follow.", D, "IN FORCE (author confirmed)"),
    ("Pooled means",
     "Box plots or summaries pooled over ALL treatments (e.g. rice vs wheat season only) carry no treatment contrast and are not entered (EXCLUDED_rows note).",
     "179 Ahmad Figs S6-S7 (MWD, GMD).", D, "IN FORCE (author confirmed)"),
    ("Parameter-specific",
     "C input printed by soil layer is summed over the layers for the C input sheet (no depth column), layer values in Notes; rice- and wheat-season inputs are "
     "separate rows and are NOT red (crop input, not a soil sampling).",
     "179 Ahmad Figs S8-S9 (straw + root + rhizodeposition C, 0-45 cm).", D, "IN FORCE (author confirmed)"),
    ("Treatment codes",
     "Re-submission of an INCLUDED old-master paper (rule 6): row a reproduces the old-master code set (DUPLICATE FLAG); treatments the old entry excluded only under "
     "rules since superseded (fertiliser levels, second residue level, mungbean / third crop) are added as rows b, c, d under rules 20 / 33 / 71 (flagged).",
     "24 (companion 2) Choudhary 2018 ASE: a T1/T4/T5/T9; b T2/T4/T6/T8; c T1/T3/T5/T7; d T1/T4/T5/T10.", D, "IN FORCE (author confirmed)"),
    ("Treatment codes",
     "Combine-harvested rice stubble left on the surface as mulch in zero-till wheat counts as residue retention -> CA (pCA when the rice phase is puddled); conventional "
     "tillage that destroys / incorporates the same stubble -> CTR - even when residue is not a treatment factor in the paper.",
     "180 Mann et al. 2008: zero tillage -> CA, conventional (rototiller + 4 cultivator + 2 planking) -> CTR.", "Author 2026-10-05", "IN FORCE (author)"),
    ("Treatment codes",
     "Raised beds whose formation (permanent vs re-made, tillage before bed making) is not described -> MT; several bed variants (two-row / three-row beds) = rows a / b.",
     "180 Mann et al. 2008 experiment 2 (Table VIII): beds with two rows (row a) and three rows (row b) -> MT.", "Author 2026-10-05", "IN FORCE (author)"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("new rules", n0 + 1, "-", n0 + len(ROWS))
