"""Batch 10 defaults (2026-10-05) recorded as RULES; pending author answers flagged."""
import sys
from copy import copy
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
ROWS = [
    ("Treatment codes",
     "Residue-type trials under conventional puddled rice-wheat (straw incorporated vs none): CT = no-residue plot at the SAME fertiliser level, CTR = straw plot; one row per straw type. "
     "Unfertilised tiers excluded (rule 78); composted straw held pending the author.",
     "172 Tirol-Padre 2005 (+N tier: urea vs rice straw row a / wheat straw row b; -N tier and rice-straw compost not entered).", "Default 2026-10-05 (rules 20, 78)", "PENDING CONFIRMATION"),
    ("Treatment codes",
     "Integrated-crop-management modules: rows pair the same fertiliser module (100 % RF / 75 % RF + bio-fertiliser) across tillage codes; '+ mungbean' CA modules as further rows (c/d). "
     "Annual deep chisel ploughing (30 cm) before puddling in the authors' conventional module kept as CT (flag); tilled DSR + freshly made raised-bed wheat held pending the author.",
     "173 Biswakarma 2023: ICM1/2 CT; ICM5-8 CA (rows a-d); ICM3/4 not entered.", "Default 2026-10-05 (rules 15, 20, 28, 33)", "PENDING CONFIRMATION"),
    ("Workflow & data management",
     "Companion re-entry of an INCLUDED old-master paper (rules 85/86) repeats the old CT / CA-family values in the same row as the new partial code, with a DUPLICATE FLAG.",
     "124 (companion) Hoque 2023: AT = pMTR; CT and MTR equal old-master 124 values.", "Default 2026-10-05 (rules 85, 86, 95)", "IN FORCE"),
    ("Parameter-specific",
     "Unit conversions used: urease ug NH4-N -> ug urea x 60/28; respiration mg C -> mg CO2 x 44/12; carbon budget printed in kg CO2-eq -> kg C x 12/44; 'CE (%)' printed as CO/CI -> CER ratio; "
     "mg/100 g x 10 = mg/kg; PMN and Olsen P to kg/ha with the paper's BD and an assumed 15-cm plough layer when depth is not stated (flagged).",
     "172 Tirol-Padre 2005; 173 Biswakarma 2023.", "Default 2026-10-05", "IN FORCE"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("rules", n0 + 1, "-", n0 + len(ROWS))
