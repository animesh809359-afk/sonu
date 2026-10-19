"""Plot-level feature extraction for the date-wise workbooks (contact-probe ASD leaf spectra).

Everything works on plot-average spectra R (plots × bands, 1-nm grid 350–2500 nm). Water-absorption bands are
kept (contact probe: no atmospheric path). The same functions are applied twice: to the raw plot averages and to
the smoothed plot averages (detector-join correction + Savitzky–Golay), so the two branches differ only in the
input spectrum.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from scipy.ndimage import gaussian_filter1d, uniform_filter1d
from scipy.signal import savgol_filter
from scipy import stats as st

import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from sp.load_asd import load_all  # noqa: E402
from sp.preprocess import splice_correct  # noqa: E402

JOINS = [1000, 1800]                 # last band of VNIR and SWIR-1 detectors
SG = (11, 2)                         # Savitzky–Golay window (bands = nm) and polynomial order
CR_WINDOWS = [("Blue absorption (carotenoids/chlorophyll, ~495 nm)", "CR1", 400, 548),
              ("Red absorption (chlorophyll, ~675 nm)", "CR2", 548, 765),
              ("Water absorption ~970 nm", "CR3", 914, 1069),
              ("Water absorption ~1200 nm", "CR4", 1069, 1258),
              ("Water absorption ~1450 nm", "CR5", 1258, 1721),
              ("Water absorption ~1940 nm", "CR6", 1721, 2304)]
OSI_RANGE = (400, 2500)


# ============================================================ data
def load_date(path, date):
    """Leaf spectra of one date: second data blocks that only repeat block-1 spectra are dropped by load_all;
    exact duplicate spectra of the same plot (unnamed copies) are dropped here."""
    meta, R, wl, logs = load_all([(path, date)])
    keep = np.ones(len(meta), bool)
    dups = []
    for p, g in meta.groupby("Plot"):
        idx = g.index.to_numpy()
        for a_i, a in enumerate(idx):
            if not keep[a]:
                continue
            for b in idx[a_i + 1:]:
                if keep[b] and np.array_equal(R[a], R[b]):
                    drop = b if not meta.at[a, "Unnamed"] or meta.at[b, "Unnamed"] else a
                    other = a if drop == b else b
                    keep[drop] = False
                    dups.append({"Plot": int(p), "Dropped column": meta.at[drop, "Source column"],
                                 "Dropped file": meta.at[drop, "File"], "Identical to column": meta.at[other, "Source column"],
                                 "Kept file": meta.at[other, "File"]})
    return meta[keep].reset_index(drop=True), R[keep], wl, logs[0], pd.DataFrame(dups)


def plot_average(meta, R):
    rows, avg = [], []
    for p, g in meta.groupby("Plot", sort=True):
        idx = g.index.to_numpy()
        f = g.iloc[0]
        rows.append({"Plot": int(p), "Variety": f["Variety"], "Growth stage": f["Growth stage"], "Class": f["Class"],
                     "Leaf spectra averaged": len(idx),
                     "Spectra (file names)": "; ".join(str(x) for x in g["File"]),
                     "Source columns": ", ".join(g["Source column"])})
        avg.append(R[idx].mean(0))
    return pd.DataFrame(rows), np.vstack(avg)


def workbook_averages(path):
    """Cached values of the plot-average (=AVERAGE) columns of the labelled workbook: {plot: (formula, values)}."""
    import openpyxl
    wf = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rf, rv = list(wf.iter_rows(values_only=True)), list(wv.iter_rows(values_only=True))
    k = 0
    while not isinstance(rf[k][0], (int, float)):
        k += 1
    out = {}
    for i, h in enumerate(rf[0]):
        if i and h is not None and not isinstance(h, str):
            vals = np.array([rv[r][i] if isinstance(rv[r][i], (int, float)) else np.nan for r in range(k, len(rv))], float)
            out.setdefault(int(h), []).append((str(rf[k][i]), vals))
    return out


# ============================================================ smoothing
def smooth(R, wl):
    Rsc, off = splice_correct(R, wl, JOINS, 10)
    return savgol_filter(Rsc, SG[0], SG[1], axis=1, mode="interp"), Rsc, off


def smoothing_alternatives(R, wl):
    Rsc, _ = splice_correct(R, wl, JOINS, 10)
    return {"Moving average, 11 nm": uniform_filter1d(Rsc, 11, axis=1, mode="nearest"),
            "Gaussian filter, σ = 3 nm": gaussian_filter1d(Rsc, 3, axis=1, mode="nearest"),
            "Savitzky–Golay, 11 nm, order 2 (selected)": savgol_filter(Rsc, 11, 2, axis=1, mode="interp"),
            "Savitzky–Golay, 21 nm, order 3": savgol_filter(Rsc, 21, 3, axis=1, mode="interp"),
            "Savitzky–Golay, 31 nm, order 2": savgol_filter(Rsc, 31, 2, axis=1, mode="interp")}


# ============================================================ derivatives (central finite differences, Δλ = 1 nm)
def derivatives(R):
    n = R.shape[1]
    D1 = np.full_like(R, np.nan)
    D2 = np.full_like(R, np.nan)
    D3 = np.full_like(R, np.nan)
    D1[:, 1:-1] = (R[:, 2:] - R[:, :-2]) / 2
    D2[:, 1:-1] = R[:, 2:] - 2 * R[:, 1:-1] + R[:, :-2]
    D3[:, 2:n - 2] = (R[:, 4:] - 2 * R[:, 3:-1] + 2 * R[:, 1:-3] - R[:, :-4]) / 2
    return D1, D2, D3


# ============================================================ continuum removal
def upper_hull(x, y):
    pts = []
    for i in range(len(x)):
        while len(pts) >= 2:
            (x1, y1), (x2, y2) = pts[-2], pts[-1]
            if (x2 - x1) * (y[i] - y1) - (y2 - y1) * (x[i] - x1) >= 0:
                pts.pop()
            else:
                break
        pts.append((x[i], y[i]))
    hx, hy = zip(*pts)
    return np.interp(x, hx, hy)


CR_PARAMS = [("depth", "maximum band depth, 1 − CR at the absorption minimum"),
             ("position (nm)", "wavelength of the maximum band depth"),
             ("width (nm)", "full width at half of the maximum band depth"),
             ("area (nm)", "area of the band-depth curve, Σ(1 − CR)·Δλ"),
             ("left area (nm)", "area from the window start to the absorption minimum"),
             ("right area (nm)", "area from the absorption minimum to the window end"),
             ("asymmetry", "left area / total area (0.5 = symmetric)"),
             ("depth / area (1/nm)", "maximum band depth normalised by area (BNA at the minimum)")]


def continuum(R, wl):
    """Returns parameter table (plots × windows·params) and the CR curves per window."""
    cols, names, curves = {}, [], {}
    for long, code, lo, hi in CR_WINDOWS:
        m = (wl >= lo) & (wl <= hi)
        x = wl[m]
        CR = np.vstack([r[m] / upper_hull(x, r[m]) for r in R])
        curves[code] = (x, CR)
        BD = 1 - CR
        imin = BD.argmax(1)
        depth = BD.max(1)
        area = BD.sum(1)
        left = np.array([BD[i, :imin[i] + 1].sum() for i in range(len(R))])
        width = np.empty(len(R))
        for i in range(len(R)):
            half = BD[i] >= depth[i] / 2
            a = b = imin[i]
            while a > 0 and half[a - 1]:
                a -= 1
            while b < len(x) - 1 and half[b + 1]:
                b += 1
            width[i] = x[b] - x[a] + 1
        vals = [depth, x[imin].astype(float), width, area, left, area - left,
                np.where(area > 0, left / np.where(area > 0, area, 1), np.nan),
                np.where(area > 0, depth / np.where(area > 0, area, 1), np.nan)]
        for (p, _), v in zip(CR_PARAMS, vals):
            nm = f"{code} {long.split(' (')[0].split(' ~')[0].lower()} {lo}–{hi} nm: {p}"
            cols[nm] = v
            names.append(nm)
    return pd.DataFrame(cols), curves


# ============================================================ vegetation indices
def _bands(R, wl):
    def f(lam):
        return R[:, int(np.argmin(np.abs(wl - lam)))]
    return f


def _mean(R, wl, a, b):
    m = (wl >= a) & (wl <= b)
    return R[:, m].mean(1)


# (abbreviation, full name, equation, reference as given in the activity table or original source, group,
#  in the activity table?, function(R, Rm) where R(λ) is the band and Rm(a, b) a band mean)
VI = [
    # ---------------- the activity table (31 indices)
    ("WI", "Water Index", "R900 / R970", "Peñuelas et al., 1997", "Water", True, lambda R, M: R(900) / R(970)),
    ("mCARI1", "Modified Chlorophyll Absorption in Reflectance Index 1", "1.2[2.5(R761 − R651) − 1.3(R761 − R581)]",
     "Haboudane et al., 2004", "Chlorophyll / greenness", True, lambda R, M: 1.2 * (2.5 * (R(761) - R(651)) - 1.3 * (R(761) - R(581)))),
    ("MTVI1", "Modified Triangular Vegetation Index 1", "1.2[1.2(R761 − R581) − 2.5(R651 − R581)]", "Haboudane et al., 2004",
     "Chlorophyll / greenness", True, lambda R, M: 1.2 * (1.2 * (R(761) - R(581)) - 2.5 * (R(651) - R(581)))),
    ("MTVI2", "Modified Triangular Vegetation Index 2",
     "1.5[1.2(R761 − R581) − 2.5(R651 − R581)] / √[(2R761 + 1)² − (6R761 − 5√R651) − 0.5]", "Haboudane et al., 2004",
     "Chlorophyll / greenness", True,
     lambda R, M: 1.5 * (1.2 * (R(761) - R(581)) - 2.5 * (R(651) - R(581))) / np.sqrt((2 * R(761) + 1) ** 2 - (6 * R(761) - 5 * np.sqrt(R(651))) - 0.5)),
    ("SR900", "Simple Ratio Index 900", "R900 / R680", "Jordan, 1969", "Chlorophyll / greenness", True, lambda R, M: R(900) / R(680)),
    ("SR850", "Simple Ratio Index 850", "R850 / R680", "Jordan, 1969", "Chlorophyll / greenness", True, lambda R, M: R(850) / R(680)),
    ("GNDVI", "Green Normalized Difference Vegetation Index", "(R850 − R580) / (R850 + R580)", "Gitelson and Merzlyak, 1996",
     "Chlorophyll / greenness", True, lambda R, M: (R(850) - R(580)) / (R(850) + R(580))),
    ("PRI", "Photochemical Reflectance Index", "(R531 − R570) / (R531 + R570)", "Gamon et al., 1992", "Photosynthetic efficiency / xanthophyll",
     True, lambda R, M: (R(531) - R(570)) / (R(531) + R(570))),
    ("RARSa", "Ratio Analysis of Reflectance Spectra, chlorophyll a", "R675 / R700", "Chappelle et al., 1992", "Pigments", True,
     lambda R, M: R(675) / R(700)),
    ("RARSb", "Ratio Analysis of Reflectance Spectra, chlorophyll b", "R675 / (R700 × R650)", "Chappelle et al., 1992", "Pigments", True,
     lambda R, M: R(675) / (R(700) * R(650))),
    ("RARSc", "Ratio Analysis of Reflectance Spectra, carotenoids", "R750 / R700", "Chappelle et al., 1992", "Pigments", True,
     lambda R, M: R(750) / R(700)),
    ("PSSRa", "Pigment Specific Simple Ratio, chlorophyll a", "R800 / R680", "Blackburn, 1998", "Pigments", True, lambda R, M: R(800) / R(680)),
    ("NDVI780", "Normalized Difference Vegetation Index 780", "(R780 − R670) / (R780 + R670)", "Raun et al., 2001", "Chlorophyll / greenness",
     True, lambda R, M: (R(780) - R(670)) / (R(780) + R(670))),
    ("SIPI", "Structure Insensitive Pigment Index", "(R800 − R445) / (R800 − R680)", "Peñuelas et al., 1995", "Pigments", True,
     lambda R, M: (R(800) - R(445)) / (R(800) - R(680))),
    ("NCPI", "Normalized Chlorophyll Pigment Ratio Index", "(R670 − R450) / (R670 + R450)", "Peñuelas et al., 1994", "Pigments", True,
     lambda R, M: (R(670) - R(450)) / (R(670) + R(450))),
    ("NPQI", "Normalized Phaeophytinization Index", "(R415 − R435) / (R415 + R435)", "Barnes et al., 1992", "Senescence / chlorophyll degradation",
     True, lambda R, M: (R(415) - R(435)) / (R(415) + R(435))),
    ("PSRI", "Plant Senescence Reflectance Index", "(R680 − R500) / R750", "Merzlyak et al., 1999 (activity table: Peñuelas et al., 1995)",
     "Senescence / chlorophyll degradation", True, lambda R, M: (R(680) - R(500)) / R(750)),
    ("WI/ND", "Ratio of Water Index and NDVI", "(R900 / R970) / [(R900 − R680) / (R900 + R680)]",
     "Peñuelas et al., 1997 (activity table: Hunt et al., 1989)", "Water", True,
     lambda R, M: (R(900) / R(970)) / ((R(900) - R(680)) / (R(900) + R(680)))),
    ("TCARI", "Transformed Chlorophyll Absorption in Reflectance Index", "3[(R700 − R670) − 0.2(R700 − R550)(R700 / R670)]",
     "Haboudane et al., 2002", "Chlorophyll / greenness", True,
     lambda R, M: 3 * ((R(700) - R(670)) - 0.2 * (R(700) - R(550)) * (R(700) / R(670)))),
    ("VIGreen", "Green Vegetation Index", "(R760 − R550) / (R760 + R550)", "Gitelson et al., 2002a", "Chlorophyll / greenness", True,
     lambda R, M: (R(760) - R(550)) / (R(760) + R(550))),
    ("RVS1", "Red-Edge Vegetation Stress Index 1", "R651 + R750", "Merton, 1998", "Red edge / stress", True, lambda R, M: R(651) + R(750)),
    ("RVS2", "Red-Edge Vegetation Stress Index 2", "(R651 + R750) / (R651 − R750)", "Merton, 1998", "Red edge / stress", True,
     lambda R, M: (R(651) + R(750)) / (R(651) - R(750))),
    ("TVI", "Triangle Vegetation Index", "0.5[120(R761 − R581) − 200(R651 − R581)]", "Broge and Leblanc, 2001", "Chlorophyll / greenness", True,
     lambda R, M: 0.5 * (120 * (R(761) - R(581)) - 200 * (R(651) - R(581)))),
    ("RDVI", "Renormalized Difference Vegetation Index", "(R761 − R651) / √(R761 + R651)", "Roujean and Breon, 1995", "Chlorophyll / greenness",
     True, lambda R, M: (R(761) - R(651)) / np.sqrt(R(761) + R(651))),
    ("NDVI850", "Normalized Difference Vegetation Index 850", "(R850 − R651) / (R850 + R651)", "Raun et al., 2001", "Chlorophyll / greenness",
     True, lambda R, M: (R(850) - R(651)) / (R(850) + R(651))),
    ("SR761", "Simple Ratio Index 761", "R761 / R651", "Jordan, 1969", "Chlorophyll / greenness", True, lambda R, M: R(761) / R(651)),
    ("NDVI761", "Normalized Difference Vegetation Index 761", "(R761 − R651) / (R761 + R651)", "Raun et al., 2001", "Chlorophyll / greenness",
     True, lambda R, M: (R(761) - R(651)) / (R(761) + R(651))),
    ("PPR", "Plant Pigment Ratio", "R550 / R450", "Metternicht, 2003", "Pigments", True, lambda R, M: R(550) / R(450)),
    ("NWI2", "Normalized Water Index 2", "(R970 − R850) / (R970 + R850)", "Babar et al., 2006", "Water", True,
     lambda R, M: (R(970) - R(850)) / (R(970) + R(850))),
    ("NRI", "Nitrogen Reflectance Index", "(R570 − R670) / (R570 + R670)", "Bausch and Duke, 1996", "Nitrogen", True,
     lambda R, M: (R(570) - R(670)) / (R(570) + R(670))),
    ("ARI", "Anthocyanin Reflectance Index", "1/R550 − 1/R700", "Gitelson et al., 2001", "Pigments", True, lambda R, M: 1 / R(550) - 1 / R(700)),
    # ---------------- additional indices related to disease, stress, pigments and water
    ("NDVI", "Normalized Difference Vegetation Index", "(R800 − R670) / (R800 + R670)", "Rouse et al., 1974", "Chlorophyll / greenness", False,
     lambda R, M: (R(800) - R(670)) / (R(800) + R(670))),
    ("DVI", "Difference Vegetation Index", "R800 − R670", "Tucker, 1979", "Chlorophyll / greenness", False, lambda R, M: R(800) - R(670)),
    ("SAVI", "Soil-Adjusted Vegetation Index", "1.5(R800 − R670) / (R800 + R670 + 0.5)", "Huete, 1988", "Chlorophyll / greenness", False,
     lambda R, M: 1.5 * (R(800) - R(670)) / (R(800) + R(670) + 0.5)),
    ("OSAVI", "Optimized Soil-Adjusted Vegetation Index", "1.16(R800 − R670) / (R800 + R670 + 0.16)", "Rondeaux et al., 1996",
     "Chlorophyll / greenness", False, lambda R, M: 1.16 * (R(800) - R(670)) / (R(800) + R(670) + 0.16)),
    ("MSAVI", "Modified Soil-Adjusted Vegetation Index", "0.5[2R800 + 1 − √((2R800 + 1)² − 8(R800 − R670))]", "Qi et al., 1994",
     "Chlorophyll / greenness", False, lambda R, M: 0.5 * (2 * R(800) + 1 - np.sqrt((2 * R(800) + 1) ** 2 - 8 * (R(800) - R(670))))),
    ("EVI", "Enhanced Vegetation Index", "2.5(R800 − R670) / (R800 + 6R670 − 7.5R475 + 1)", "Huete et al., 2002", "Chlorophyll / greenness",
     False, lambda R, M: 2.5 * (R(800) - R(670)) / (R(800) + 6 * R(670) - 7.5 * R(475) + 1)),
    ("MSR", "Modified Simple Ratio", "(R800/R670 − 1) / √(R800/R670 + 1)", "Chen, 1996", "Chlorophyll / greenness", False,
     lambda R, M: (R(800) / R(670) - 1) / np.sqrt(R(800) / R(670) + 1)),
    ("MCARI", "Modified Chlorophyll Absorption in Reflectance Index", "[(R700 − R670) − 0.2(R700 − R550)](R700 / R670)",
     "Daughtry et al., 2000", "Chlorophyll / greenness", False, lambda R, M: ((R(700) - R(670)) - 0.2 * (R(700) - R(550))) * (R(700) / R(670))),
    ("MCARI2", "Modified Chlorophyll Absorption in Reflectance Index 2",
     "1.5[2.5(R800 − R670) − 1.3(R800 − R550)] / √[(2R800 + 1)² − (6R800 − 5√R670) − 0.5]", "Haboudane et al., 2004",
     "Chlorophyll / greenness", False,
     lambda R, M: 1.5 * (2.5 * (R(800) - R(670)) - 1.3 * (R(800) - R(550))) / np.sqrt((2 * R(800) + 1) ** 2 - (6 * R(800) - 5 * np.sqrt(R(670))) - 0.5)),
    ("TCARI/OSAVI", "TCARI / OSAVI ratio", "TCARI / OSAVI", "Haboudane et al., 2002", "Chlorophyll / greenness", False,
     lambda R, M: 3 * ((R(700) - R(670)) - 0.2 * (R(700) - R(550)) * (R(700) / R(670))) / (1.16 * (R(800) - R(670)) / (R(800) + R(670) + 0.16))),
    ("TGI", "Triangular Greenness Index", "−0.5[190(R670 − R550) − 120(R670 − R480)]", "Hunt et al., 2013", "Chlorophyll / greenness", False,
     lambda R, M: -0.5 * (190 * (R(670) - R(550)) - 120 * (R(670) - R(480)))),
    ("GI", "Greenness Index", "R554 / R677", "Zarco-Tejada et al., 2005", "Chlorophyll / greenness", False, lambda R, M: R(554) / R(677)),
    ("CIgreen", "Chlorophyll Index green", "R780 / R550 − 1", "Gitelson et al., 2003", "Chlorophyll / greenness", False,
     lambda R, M: R(780) / R(550) - 1),
    ("CIrededge", "Chlorophyll Index red edge", "R780 / R710 − 1", "Gitelson et al., 2003", "Red edge / stress", False,
     lambda R, M: R(780) / R(710) - 1),
    ("ND705", "Red-edge Normalized Difference 705", "(R750 − R705) / (R750 + R705)", "Gitelson and Merzlyak, 1994", "Red edge / stress", False,
     lambda R, M: (R(750) - R(705)) / (R(750) + R(705))),
    ("mND705", "Modified Red-edge Normalized Difference 705", "(R750 − R705) / (R750 + R705 − 2R445)", "Sims and Gamon, 2002",
     "Red edge / stress", False, lambda R, M: (R(750) - R(705)) / (R(750) + R(705) - 2 * R(445))),
    ("mSR705", "Modified Red-edge Simple Ratio 705", "(R750 − R445) / (R705 − R445)", "Sims and Gamon, 2002", "Red edge / stress", False,
     lambda R, M: (R(750) - R(445)) / (R(705) - R(445))),
    ("NDRE", "Normalized Difference Red Edge", "(R790 − R720) / (R790 + R720)", "Barnes et al., 2000", "Red edge / stress", False,
     lambda R, M: (R(790) - R(720)) / (R(790) + R(720))),
    ("MTCI", "MERIS Terrestrial Chlorophyll Index", "(R754 − R709) / (R709 − R681)", "Dash and Curran, 2004", "Red edge / stress", False,
     lambda R, M: (R(754) - R(709)) / (R(709) - R(681))),
    ("VOG1", "Vogelmann Red Edge Index 1", "R740 / R720", "Vogelmann et al., 1993", "Red edge / stress", False, lambda R, M: R(740) / R(720)),
    ("VOG2", "Vogelmann Red Edge Index 2", "(R734 − R747) / (R715 + R726)", "Zarco-Tejada et al., 2001", "Red edge / stress", False,
     lambda R, M: (R(734) - R(747)) / (R(715) + R(726))),
    ("ZTM", "Zarco-Tejada and Miller index", "R750 / R710", "Zarco-Tejada et al., 2001", "Red edge / stress", False, lambda R, M: R(750) / R(710)),
    ("Datt", "Datt chlorophyll index", "(R850 − R710) / (R850 − R680)", "Datt, 1999", "Red edge / stress", False,
     lambda R, M: (R(850) - R(710)) / (R(850) - R(680))),
    ("RVSI", "Red-Edge Vegetation Stress Index (original definition)", "(R714 + R752) / 2 − R733", "Merton and Huntington, 1999",
     "Red edge / stress", False, lambda R, M: (R(714) + R(752)) / 2 - R(733)),
    ("GM1", "Gitelson–Merzlyak index 1", "R750 / R550", "Gitelson and Merzlyak, 1997", "Pigments", False, lambda R, M: R(750) / R(550)),
    ("GM2", "Gitelson–Merzlyak index 2", "R750 / R700", "Gitelson and Merzlyak, 1997", "Pigments", False, lambda R, M: R(750) / R(700)),
    ("PSSRb", "Pigment Specific Simple Ratio, chlorophyll b", "R800 / R635", "Blackburn, 1998", "Pigments", False, lambda R, M: R(800) / R(635)),
    ("PSSRc", "Pigment Specific Simple Ratio, carotenoids", "R800 / R470", "Blackburn, 1998", "Pigments", False, lambda R, M: R(800) / R(470)),
    ("PSNDa", "Pigment Specific Normalized Difference, chlorophyll a", "(R800 − R680) / (R800 + R680)", "Blackburn, 1998", "Pigments", False,
     lambda R, M: (R(800) - R(680)) / (R(800) + R(680))),
    ("PSNDb", "Pigment Specific Normalized Difference, chlorophyll b", "(R800 − R635) / (R800 + R635)", "Blackburn, 1998", "Pigments", False,
     lambda R, M: (R(800) - R(635)) / (R(800) + R(635))),
    ("CRI1", "Carotenoid Reflectance Index 1", "1/R510 − 1/R550", "Gitelson et al., 2002b", "Pigments", False, lambda R, M: 1 / R(510) - 1 / R(550)),
    ("CRI2", "Carotenoid Reflectance Index 2", "1/R510 − 1/R700", "Gitelson et al., 2002b", "Pigments", False, lambda R, M: 1 / R(510) - 1 / R(700)),
    ("ARI2", "Anthocyanin Reflectance Index 2", "R800(1/R550 − 1/R700)", "Gitelson et al., 2001", "Pigments", False,
     lambda R, M: R(800) * (1 / R(550) - 1 / R(700))),
    ("RGRI", "Red/Green Ratio Index", "mean(R600–699) / mean(R500–599)", "Gamon and Surfus, 1999", "Pigments", False,
     lambda R, M: M(600, 699) / M(500, 599)),
    ("PRI515", "Photochemical Reflectance Index 515", "(R515 − R531) / (R515 + R531)", "Hernández-Clemente et al., 2011",
     "Photosynthetic efficiency / xanthophyll", False, lambda R, M: (R(515) - R(531)) / (R(515) + R(531))),
    ("Ctr1", "Carter stress index 1", "R695 / R420", "Carter, 1994", "Red edge / stress", False, lambda R, M: R(695) / R(420)),
    ("Ctr2", "Carter stress index 2", "R695 / R760", "Carter, 1994", "Red edge / stress", False, lambda R, M: R(695) / R(760)),
    ("Lic", "Lichtenthaler blue/red ratio", "R440 / R690", "Lichtenthaler et al., 1996", "Red edge / stress", False, lambda R, M: R(440) / R(690)),
    ("HI", "Healthy Index", "(R534 − R698) / (R534 + R698) − 0.5R704", "Mahlein et al., 2013", "Disease-specific", False,
     lambda R, M: (R(534) - R(698)) / (R(534) + R(698)) - 0.5 * R(704)),
    ("CLSI", "Cercospora Leaf Spot Index", "(R698 − R570) / (R698 + R570) − R734", "Mahlein et al., 2013", "Disease-specific", False,
     lambda R, M: (R(698) - R(570)) / (R(698) + R(570)) - R(734)),
    ("SBRI", "Sugar Beet Rust Index", "(R570 − R513) / (R570 + R513) + 0.5R704", "Mahlein et al., 2013", "Disease-specific", False,
     lambda R, M: (R(570) - R(513)) / (R(570) + R(513)) + 0.5 * R(704)),
    ("PMI", "Powdery Mildew Index", "(R520 − R584) / (R520 + R584) + R724", "Mahlein et al., 2013", "Disease-specific", False,
     lambda R, M: (R(520) - R(584)) / (R(520) + R(584)) + R(724)),
    ("LRDSI1", "Leaf Rust Disease Severity Index 1", "6.9(R605 / R455) − 1.2", "Ashourloo et al., 2014", "Disease-specific", False,
     lambda R, M: 6.9 * R(605) / R(455) - 1.2),
    ("LRDSI2", "Leaf Rust Disease Severity Index 2", "4.2(R695 / R455) − 0.38", "Ashourloo et al., 2014", "Disease-specific", False,
     lambda R, M: 4.2 * R(695) / R(455) - 0.38),
    ("DSWI", "Disease–Water Stress Index", "(R802 + R547) / (R1657 + R682)", "Apan et al., 2004", "Disease-specific", False,
     lambda R, M: (R(802) + R(547)) / (R(1657) + R(682))),
    ("NWI1", "Normalized Water Index 1", "(R970 − R900) / (R970 + R900)", "Babar et al., 2006", "Water", False,
     lambda R, M: (R(970) - R(900)) / (R(970) + R(900))),
    ("NWI3", "Normalized Water Index 3", "(R970 − R920) / (R970 + R920)", "Babar et al., 2006", "Water", False,
     lambda R, M: (R(970) - R(920)) / (R(970) + R(920))),
    ("NWI4", "Normalized Water Index 4", "(R970 − R880) / (R970 + R880)", "Babar et al., 2006", "Water", False,
     lambda R, M: (R(970) - R(880)) / (R(970) + R(880))),
    ("NDWI", "Normalized Difference Water Index", "(R860 − R1240) / (R860 + R1240)", "Gao, 1996", "Water", False,
     lambda R, M: (R(860) - R(1240)) / (R(860) + R(1240))),
    ("MSI", "Moisture Stress Index", "R1600 / R820", "Hunt and Rock, 1989", "Water", False, lambda R, M: R(1600) / R(820)),
    ("NDII", "Normalized Difference Infrared Index", "(R819 − R1649) / (R819 + R1649)", "Hardisky et al., 1983", "Water", False,
     lambda R, M: (R(819) - R(1649)) / (R(819) + R(1649))),
    ("LWVI1", "Leaf Water Vegetation Index 1", "(R1094 − R983) / (R1094 + R983)", "Galvão et al., 2005", "Water", False,
     lambda R, M: (R(1094) - R(983)) / (R(1094) + R(983))),
    ("LWVI2", "Leaf Water Vegetation Index 2", "(R1094 − R1205) / (R1094 + R1205)", "Galvão et al., 2005", "Water", False,
     lambda R, M: (R(1094) - R(1205)) / (R(1094) + R(1205))),
    ("NDNI", "Normalized Difference Nitrogen Index", "[log(1/R1510) − log(1/R1680)] / [log(1/R1510) + log(1/R1680)]", "Serrano et al., 2002",
     "Dry matter / biochemistry", False,
     lambda R, M: (np.log10(1 / R(1510)) - np.log10(1 / R(1680))) / (np.log10(1 / R(1510)) + np.log10(1 / R(1680)))),
    ("NDLI", "Normalized Difference Lignin Index", "[log(1/R1754) − log(1/R1680)] / [log(1/R1754) + log(1/R1680)]", "Serrano et al., 2002",
     "Dry matter / biochemistry", False,
     lambda R, M: (np.log10(1 / R(1754)) - np.log10(1 / R(1680))) / (np.log10(1 / R(1754)) + np.log10(1 / R(1680)))),
    ("CAI", "Cellulose Absorption Index", "0.5(R2000 + R2200) − R2100", "Nagler et al., 2003", "Dry matter / biochemistry", False,
     lambda R, M: 0.5 * (R(2000) + R(2200)) - R(2100)),
]


def indices(R, wl):
    Rf = _bands(R, wl)
    M = lambda a, b: _mean(R, wl, a, b)  # noqa: E731
    with np.errstate(divide="ignore", invalid="ignore"):
        return pd.DataFrame({v[0]: v[6](Rf, M) for v in VI})


def index_definitions():
    return pd.DataFrame([{"Index": v[0], "Full name": v[1], "Equation (R = reflectance at the wavelength in nm)": v[2],
                          "Reference": v[3], "Group": v[4],
                          "Source": "Activity 1.4 table" if v[5] else "Additional (disease / stress related)"} for v in VI])


# ============================================================ red-edge parameters
RE_DEF = [
    ("REIP, maximum first derivative (nm)", "wavelength of the maximum first derivative in 680–760 nm", "Horler et al., 1983"),
    ("REIP, linear four-point interpolation (nm)", "700 + 40[(R670 + R780)/2 − R700] / (R740 − R700)", "Guyot and Baret, 1988"),
    ("REIP, Lagrangian interpolation (nm)", "vertex of the parabola through the first derivative at the maximum and its two neighbours",
     "Dawson and Curran, 1998"),
    ("REIP, linear extrapolation (nm)", "intersection of the first-derivative lines 680–694 nm and 724–760 nm", "Cho and Skidmore, 2006"),
    ("Red-edge amplitude Dr (1/nm)", "maximum first derivative in 680–760 nm", "Gong et al., 2002"),
    ("Red-edge area SDr", "sum of the first derivative in 680–760 nm", "Gong et al., 2002"),
    ("Blue-edge amplitude Db (1/nm)", "maximum first derivative in 490–530 nm", "Gong et al., 2002"),
    ("Blue-edge position λb (nm)", "wavelength of Db", "Gong et al., 2002"),
    ("Blue-edge area SDb", "sum of the first derivative in 490–530 nm", "Gong et al., 2002"),
    ("Yellow-edge amplitude Dy (1/nm)", "minimum first derivative in 550–582 nm", "Gong et al., 2002"),
    ("Yellow-edge position λy (nm)", "wavelength of Dy", "Gong et al., 2002"),
    ("Yellow-edge area SDy", "sum of the first derivative in 550–582 nm", "Gong et al., 2002"),
    ("Green peak reflectance Rg", "maximum reflectance in 510–560 nm", "Gong et al., 2002"),
    ("Green peak position λg (nm)", "wavelength of Rg", "Gong et al., 2002"),
    ("Red well reflectance Ro", "minimum reflectance in 640–680 nm", "Gong et al., 2002"),
    ("Red well position λo (nm)", "wavelength of Ro", "Gong et al., 2002"),
    ("Rg / Ro", "green peak / red well", "Gong et al., 2002"),
    ("(Rg − Ro) / (Rg + Ro)", "normalised green peak – red well", "Gong et al., 2002"),
    ("SDr / SDb", "red-edge area / blue-edge area", "Gong et al., 2002"),
    ("SDr / SDy", "red-edge area / yellow-edge area", "Gong et al., 2002"),
    ("(SDr − SDb) / (SDr + SDb)", "normalised red-edge – blue-edge area", "Gong et al., 2002"),
    ("(SDr − SDy) / (SDr + SDy)", "normalised red-edge – yellow-edge area", "Gong et al., 2002"),
    ("Red-edge slope (R750 − R700)/50 (1/nm)", "mean slope of the red edge", "Filella and Peñuelas, 1994"),
]


def red_edge(R, D1, wl):
    n = len(R)
    ar = np.arange(n)

    def win(a, b):
        return (wl >= a) & (wl <= b)
    out = {}
    m = win(680, 760)
    lam = wl[m]
    d = D1[:, m]
    i = d.argmax(1)
    out[RE_DEF[0][0]] = lam[i].astype(float)
    Rf = _bands(R, wl)
    out[RE_DEF[1][0]] = 700 + 40 * ((Rf(670) + Rf(780)) / 2 - Rf(700)) / (Rf(740) - Rf(700))
    im = np.clip(i, 1, m.sum() - 2)
    d_m1, d_0, d_p1 = (d[ar, im + k] for k in (-1, 0, 1))
    l_m1, l_0, l_p1 = lam[im - 1], lam[im], lam[im + 1]
    A = d_m1 / ((l_m1 - l_0) * (l_m1 - l_p1))
    B = d_0 / ((l_0 - l_m1) * (l_0 - l_p1))
    C = d_p1 / ((l_p1 - l_m1) * (l_p1 - l_0))
    out[RE_DEF[2][0]] = (A * (l_0 + l_p1) + B * (l_m1 + l_p1) + C * (l_m1 + l_0)) / (2 * (A + B + C))
    Df = _bands(D1, wl)
    m1 = (Df(694) - Df(680)) / 14
    c1 = Df(680) - m1 * 680
    m2 = (Df(760) - Df(724)) / 36
    c2 = Df(724) - m2 * 724
    out[RE_DEF[3][0]] = -(c1 - c2) / (m1 - m2)
    out[RE_DEF[4][0]] = d[ar, i]
    out[RE_DEF[5][0]] = d.sum(1)
    for tag, (a, b), fn, names in (("b", (490, 530), np.argmax, RE_DEF[6:9]), ("y", (550, 582), np.argmin, RE_DEF[9:12])):
        mm = win(a, b)
        j = fn(D1[:, mm], axis=1)
        out[names[0][0]] = D1[:, mm][ar, j]
        out[names[1][0]] = wl[mm][j].astype(float)
        out[names[2][0]] = D1[:, mm].sum(1)
    mm = win(510, 560)
    j = R[:, mm].argmax(1)
    Rg, out[RE_DEF[13][0]] = R[:, mm][ar, j], wl[mm][j].astype(float)
    out[RE_DEF[12][0]] = Rg
    mm = win(640, 680)
    j = R[:, mm].argmin(1)
    Ro, out[RE_DEF[15][0]] = R[:, mm][ar, j], wl[mm][j].astype(float)
    out[RE_DEF[14][0]] = Ro
    SDr, SDb, SDy = out[RE_DEF[5][0]], out[RE_DEF[8][0]], out[RE_DEF[11][0]]
    out[RE_DEF[16][0]] = Rg / Ro
    out[RE_DEF[17][0]] = (Rg - Ro) / (Rg + Ro)
    out[RE_DEF[18][0]] = SDr / SDb
    out[RE_DEF[19][0]] = SDr / SDy
    out[RE_DEF[20][0]] = (SDr - SDb) / (SDr + SDb)
    out[RE_DEF[21][0]] = (SDr - SDy) / (SDr + SDy)
    out[RE_DEF[22][0]] = (Rf(750) - Rf(700)) / 50
    return pd.DataFrame({k: out[k] for k, _, _ in RE_DEF})


# ============================================================ optimised two-band indices
def _r2_rows(F, yc, ysd):
    Fc = F - F.mean(0)
    sd = Fc.std(0)
    with np.errstate(divide="ignore", invalid="ignore"):
        r = (Fc * yc[:, None]).mean(0) / (sd * ysd)
    return np.where(sd > 0, r, 0.0)


def optimise(R, wl, y, lo=OSI_RANGE[0], hi=OSI_RANGE[1]):
    """Correlation r of every two-band index with disease status (0 healthy, 1 diseased; point-biserial r).
    Returns {kind: (wl_used, r matrix)} with r[i, j] for index(λi, λj)."""
    m = (wl >= lo) & (wl <= hi)
    x, X = wl[m], R[:, m]
    yc = y - y.mean()
    ysd = y.std()
    out = {k: np.zeros((len(x), len(x)), np.float32) for k in ("NDSI", "SR", "DI")}
    for i in range(len(x)):
        a = X[:, [i]]
        with np.errstate(divide="ignore", invalid="ignore"):
            out["NDSI"][i] = _r2_rows((a - X) / (a + X), yc, ysd)
            out["SR"][i] = _r2_rows(a / X, yc, ysd)
            out["DI"][i] = _r2_rows(a - X, yc, ysd)
    for k in out:
        np.fill_diagonal(out[k], 0)
    return x, out


def straddles(a, b):
    lo, hi = min(a, b), max(a, b)
    return any(lo <= j < hi for j in JOINS)


def top_pairs(x, r, n=10, min_sep=10, kind="NDSI"):
    """Top pairs by r² (non-redundant: each new pair ≥ min_sep nm from every earlier pick in both bands)."""
    r2 = r.astype(float) ** 2
    if kind in ("NDSI", "DI"):          # antisymmetric indices: |r| symmetric, keep i < j
        r2 = np.triu(r2, 1)
    order = np.argsort(r2, axis=None)[::-1]
    picks = []
    for flat in order:
        i, j = np.unravel_index(flat, r2.shape)
        if r2[i, j] <= 0:
            break
        a, b = x[i], x[j]
        if all(abs(a - p[0]) >= min_sep or abs(b - p[1]) >= min_sep for p in picks):
            picks.append((a, b, float(r[i, j]), float(r2[i, j])))
        if len(picks) == n:
            break
    return picks


def index_value(R, wl, kind, a, b):
    Rf = _bands(R, wl)
    A, B = Rf(a), Rf(b)
    return {"NDSI": (A - B) / (A + B), "SR": A / B, "DI": A - B}[kind]


# ============================================================ statistics (plot level)
def auc(x, y):
    ok = ~np.isnan(x)
    x, y = x[ok], y[ok]
    if len(np.unique(y)) < 2:
        return np.nan
    rk = st.rankdata(x)
    n1 = (y == 1).sum()
    n0 = (y == 0).sum()
    return (rk[y == 1].sum() - n1 * (n1 + 1) / 2) / (n0 * n1)


def class_stats(F, y):
    rows = []
    for c in F.columns:
        v = F[c].to_numpy(float)
        h, d = v[(y == 0) & ~np.isnan(v)], v[(y == 1) & ~np.isnan(v)]
        if len(h) < 2 or len(d) < 2 or (np.std(h) == 0 and np.std(d) == 0):
            rows.append({"Feature": c, "Mean healthy": np.nanmean(h) if len(h) else np.nan, "Mean diseased": np.nanmean(d) if len(d) else np.nan})
            continue
        t, p = st.ttest_ind(d, h, equal_var=False)
        sp = np.sqrt(((len(h) - 1) * h.var(ddof=1) + (len(d) - 1) * d.var(ddof=1)) / (len(h) + len(d) - 2))
        g = (d.mean() - h.mean()) / sp * (1 - 3 / (4 * (len(h) + len(d)) - 9)) if sp > 0 else np.nan
        a = auc(v, y)
        rows.append({"Feature": c, "Mean healthy": h.mean(), "SD healthy": h.std(ddof=1), "Mean diseased": d.mean(),
                     "SD diseased": d.std(ddof=1), "Difference (D − H)": d.mean() - h.mean(), "Welch t": t, "p": p,
                     "Hedges g": g, "AUC (direction-free)": max(a, 1 - a), "Direction": "D > H" if d.mean() > h.mean() else "D < H"})
    df = pd.DataFrame(rows)
    if "p" in df:
        ok = df["p"].notna()
        q = np.full(len(df), np.nan)
        if ok.any():
            from statsmodels.stats.multitest import multipletests
            q[ok.to_numpy()] = multipletests(df.loc[ok, "p"], method="fdr_bh")[1]
        df.insert(df.columns.get_loc("p") + 1, "q (BH-FDR)", q)
    return df
