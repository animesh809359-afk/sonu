"""Author decisions 2026-10-05 (updated 11): compost = CTR; Biswakarma ICM1/2 DT, ICM3/4 CT; aggregate classes -> MACRO / MICRO / WSA; Obs restated and recomputed."""
import sys
from copy import copy
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
done = []
for r in range(1, ws.max_row + 1):
    n = ws.cell(r, 1).value
    if n == 97:
        ws.cell(r, 3).value = ("Residue-type trials under conventional puddled rice-wheat (straw incorporated vs none): CT = no-residue plot at the SAME fertiliser level, CTR = straw plot; "
                               "one row per straw type. Composted crop residue (rice-straw compost) counts as CTR (author). Unfertilised tiers excluded (rule 78).")
        ws.cell(r, 4).value = "172 Tirol-Padre 2005 (+N tier: urea vs rice straw row a / wheat straw row b / rice-straw compost row c; -N tier not entered)."
        ws.cell(r, 5).value, ws.cell(r, 6).value = "Author 2026-10-05", "IN FORCE (author confirmed)"
        done.append(n)
    if n == 98:
        ws.cell(r, 3).value = ("Annual deep (chisel) ploughing to ~30 cm before puddled rice = DT, even when the authors call it conventional; chisel-ploughed direct-seeded rice + raised-bed "
                               "wheat formed after disc + cultivator tillage = CT. ICM modules: rows pair the same fertiliser module across tillage codes; '+ mungbean' CA modules as further rows.")
        ws.cell(r, 4).value = "173 Biswakarma 2023: ICM1/2 DT, ICM3/4 CT, ICM5-8 CA (rows a-d). Flag: the paper states chisel ploughing for ICM1-4 alike."
        ws.cell(r, 5).value, ws.cell(r, 6).value = "Author 2026-10-05", "IN FORCE (author confirmed)"
        done.append(n)
    if n == 96:
        ws.cell(r, 4).value = ("New sheets 2026-10-05: LLWR (author request); SOIL TEMP (author confirmed). Aggregate size classes are NOT kept on a separate sheet (AGG SIZE CLASS removed "
                               "in updated 11): classes >0.25 mm -> MACRO, water-stable classes <0.25 mm -> MICRO, MACRO + MICRO -> WSA (g/g).")
        ws.cell(r, 6).value = "IN FORCE (author confirmed)"
        done.append(n)
    if n == 45:
        ws.cell(r, 6).value = "IN FORCE - restated by rule %d (recomputed for all studies in updated 11)"
        done.append(n)
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 1).value == 45:
        ws.cell(r, 6).value = ws.cell(r, 6).value % (n0 + 1)
ROWS = [
    ("Statistics (Obs, Rep, SD)",
     "OBS (author restatement 2026-10-05): Obs = Y + T + D for each row, a factor equal to 1 adding nothing (minimum 1). Y = years of data for that parameter at that depth in the study "
     "(year-wise rows: the number of years in the row's series; pooled means: years pooled); T = paper treatments under one code (rule-20 rows a/b/c, or treatments pooled into one "
     "value); D = paper depths falling in the depth class. Example: 3 years + 2 treatments + 3 depths -> Obs 8 for that depth class.",
     "Recomputed for every row of every study in updated 11 (360 of 1078 rows changed; e.g. year-wise rows now carry Y = years in the series). Each row's Notes end with the Y/T/D breakdown.",
     "Author 2026-10-05", "IN FORCE"),
    ("Parameter-specific",
     "WATER-STABLE AGGREGATE SIZE CLASSES: sum of classes >0.25 mm -> MACRO; water-stable classes <0.25 mm (down to 0.053 mm) -> MICRO; WSA = MACRO + MICRO in g/g. A finest class that "
     "includes <0.053 mm silt + clay is left out of MICRO (flag).",
     "15 (companion 2) Gathala 2011 Fig. 2: MICRO = 0.25-0.11 mm class; MACRO = Table 3 2008-09; WSA = (MACRO + MICRO)/100.", "Author 2026-10-05", "IN FORCE"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("updated", done, "new rules", n0 + 1, "-", n0 + len(ROWS))
