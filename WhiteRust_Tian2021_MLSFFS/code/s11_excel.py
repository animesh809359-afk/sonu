"""Final Excel workbook: every step of the Tian et al. (2021) ML-SFFS workflow applied to the white-rust data."""
import sys, pickle, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
sys.path.insert(0, ".")
from mlsffs import *
from content import REFS, TIAN_REF_ERRORS, CORRECTIONS, PAPER_VS_STUDY

S = sys.argv[1]; OUT = sys.argv[2]; O = f"{S}/out"; FIG = f"{S}/figs"
T = pickle.load(open(f"{O}/tables.pkl", "rb"))
CK = pickle.load(open(f"{O}/crosschecks.pkl", "rb"))
feats, y, _ = pickle.load(open(f"{O}/feats.pkl", "rb"))
F = pd.DataFrame(pickle.load(open(f"{O}/res_faithful.pkl", "rb")))
SR = pickle.load(open(f"{O}/res_stage.pkl", "rb"))
TR = pd.read_pickle(f"{O}/treatment.pkl")
CODE = T["code"]; POOL = T["pool"]
F["OFC codes"] = F.ofc.apply(lambda t: ", ".join(sorted((CODE[n] for n in t), key=lambda c_: int(c_[1:]))))

AR = "Arial"
F_TITLE = Font(name=AR, size=14, bold=True, color="0D366B")
F_H = Font(name=AR, size=10, bold=True, color="FFFFFF")
F_B = Font(name=AR, size=10)
F_BB = Font(name=AR, size=10, bold=True)
F_NOTE = Font(name=AR, size=9, italic=True, color="52514E")
F_SEC = Font(name=AR, size=11, bold=True, color="1C5CAB")
FILL_H = PatternFill("solid", fgColor="1C5CAB")
FILL_F = PatternFill("solid", fgColor="E8F6EF")     # formula cells
FILL_FLAG = PatternFill("solid", fgColor="FDF0EA")
THIN = Side(style="thin", color="D9D8D3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wb.remove(wb.active)


def sheet(name, title, lines):
    ws = wb.create_sheet(("S" + name)[:31])
    ws["A1"] = title; ws["A1"].font = F_TITLE
    r = 2
    for ln in lines:
        ws.cell(r, 1, ln).font = F_B if not ln.startswith("Note") else F_NOTE
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        ws.row_dimensions[r].height = max(15, 14 * (1 + len(ln) // 150))
        r += 1
    ws.sheet_view.showGridLines = False
    return ws, r + 1


def section(ws, r, text):
    ws.cell(r, 1, text).font = F_SEC
    return r + 1


def table(ws, r, df, c0=1, fmt=None, widths=None, wrap_cols=(), flag_col=None):
    fmt = fmt or {}
    for j, col in enumerate(df.columns):
        cell = ws.cell(r, c0 + j, str(col)); cell.font = F_H; cell.fill = FILL_H; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 30
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, (np.floating, float)) and np.isnan(v):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            elif isinstance(v, (bool, np.bool_)):
                v = "Yes" if v else "No"
            elif isinstance(v, (tuple, list)):
                v = ", ".join(map(str, v))
            cell = ws.cell(r + i, c0 + j, v); cell.font = F_B; cell.border = BOX
            col = df.columns[j]
            if col in fmt:
                cell.number_format = fmt[col]
            elif isinstance(v, float):
                cell.number_format = "0.000"
            if col in wrap_cols:
                cell.alignment = WRAP
            if flag_col is not None and df.columns[j] == flag_col and v in ("FLAG", "FAIL"):
                cell.fill = FILL_FLAG
    if widths:
        for j, w in enumerate(widths):
            if w:
                ws.column_dimensions[get_column_letter(c0 + j)].width = w
    return r + len(df) + 2


def image(ws, r, path, width_px=900, col="A"):
    img = XLImage(path)
    sc = width_px / img.width
    img.width, img.height = int(img.width * sc), int(img.height * sc)
    ws.add_image(img, f"{col}{r}")
    return r + int(img.height / 20) + 2


def setw(ws, widths):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


pct = "0.0"
# ================================================================ 00 Contents
ws, r = sheet("00_Contents", "Spectroscopic detection of white rust in mustard with ML-SFFS – workbook following Tian et al. (2021)",
              ["Methodology: Tian, L., Xue, B., Wang, Z., Li, D., Yao, X., Cao, Q., Zhu, Y., Cao, W., Cheng, T., 2021. Spectroscopic detection of rice leaf blast infection from asymptomatic to mild stages with integrated machine learning and feature selection. Remote Sensing of Environment 257, 112350.",
               "Data: plot-average raw reflectance of 48 plots (12 uninoculated controls, 36 inoculated: 3 varieties × 4 inoculation timings D1–D4 × 3 replications) measured on 4 dates T1 = 11 Feb, T2 = 24 Feb, T3 = 3 Mar, T4 = 13 Mar 2026, with the features derived from it (continuum-removal parameters, 87 vegetation indices, 1st–3rd derivatives, red-edge parameters). Wavelet coefficients (CWT) were added to follow the paper.",
               "Every step of the paper is applied in order. Where the paper's procedure needed correction for these data, the paper-faithful result AND the corrected result are both reported (corrections C1–C9, sheet 03).",
               "Green cells contain live Excel formulas that recompute results from the data in this workbook (cross-checks). All other numbers were computed in Python 3 (scikit-learn 1.9, PyWavelets 1.10, SciPy 1.18, statsmodels 0.15)."])
contents = [
    ("S01_Workflow", "Overview", "Step-by-step workflow, paper vs this study, and where each output is"),
    ("S02_Paper_vs_Study", "Overview", "Side-by-side comparison of the paper's design and this study"),
    ("S03_Corrections", "Overview", "Corrections C1–C9 to the paper's procedure, with reasons and references"),
    ("S04_Step1_Data_QC", "Step 1", "Data, treatment design, quality control, confounds, candidate feature counts"),
    ("S05_Step1_Treatment", "Step 1", "Treatment file (plot, inoculation timing, variety, replication) and growth stages"),
    ("S06_Step2_Physiology", "Step 2", "Physiological interpretation: proxy indices for chlorophyll, carotenoids, anthocyanins, water"),
    ("S07_Step2_All_Scalar_Tests", "Step 2", "Mann–Whitney U + BH-FDR for every index / CR / red-edge parameter (supplement)"),
    ("S08_Step3_Sep_Scalars", "Step 3", "Threshold separability of the 158 scalar features on each date"),
    ("S09_Step3_Sep_Bands_Derivs", "Step 3", "Threshold separability of every reflectance band and 1st–3rd derivative"),
    ("S10_Step3_Sep_CWT", "Step 3", "Threshold separability of every wavelet coefficient (scale × wavelength)"),
    ("S11_Step3_Scalograms", "Step 3", "Classification scalograms (figures)"),
    ("S12_Step4_DSSF_Selection", "Step 4", "Top-5 % / top-10 % selection, consistency across dates, regions and representatives"),
    ("S13_Step5_VIF", "Step 5", "Iterative VIF reduction log"),
    ("S14_Step5_Final_DSSF", "Step 5", "Final disease-specific spectral features F1–F16 (analogue of Tian et al. Table 2)"),
    ("S15_Step6A_Faithful", "Step 6A", "Paper-faithful ML-SFFS: OFC and OA per date and classifier (analogue of Table 3, Fig. 8)"),
    ("S16_Step6A_OFC_Detail", "Step 6A", "Most frequent OFCs and selection frequency of every DSSF"),
    ("S17_Step6A_Repetitions", "Step 6A", "All 1200 repetitions (date × classifier × 100) – source of the formulas in sheet 15"),
    ("S18_Step6B_Nested", "Step 6B", "Corrected estimate: nested cross-validation; optimism of the paper-faithful estimate"),
    ("S19_Step6C_Permutation", "Step 6C", "Permutation null of the paper-faithful procedure"),
    ("S20_Step6D_Stage", "Step 6D", "Detection by infection stage: date × inoculation timing; variety-wise accuracy"),
    ("S21_Step7_DSSF_Values", "Step 7", "Values of F1–F16 for every plot and date (data for the formula checks)"),
    ("S22_Step7_Formula_Check", "Step 7", "Excel formulas recompute class means and threshold balanced accuracy of every DSSF"),
    ("S23_Step7_Crosschecks", "Step 7", "Independent cross-checks of every step (PASS / FLAG)"),
    ("S24_References", "References", "All references with verification status; errors found in the paper's reference list"),
]
r = section(ws, r, "Sheets")
r = table(ws, r, pd.DataFrame(contents, columns=["Sheet", "Step", "Content"]), widths=[26, 12, 110])

# ================================================================ 01 Workflow
ws, r = sheet("01_Workflow", "Workflow – Tian et al. (2021) applied step by step",
              ["The paper's three modules (data preparation, feature selection, classification; their Fig. 2) are applied as Steps 1–6 below; Step 7 collects the cross-checks."])
wf = [
    ("1", "Data preparation", "Spectra + biochemistry of paired healthy/infected leaves; SIs; CWT", "Plot-average spectra of 4 dates; QC; guard zones (C1); CWT computed; treatment design checked (C9)", "04, 05"),
    ("2", "Physiological basis", "Biochemical parameters compared with t / Wilcoxon / Mann–Whitney tests", "Index proxies of the same constituents (C8), same test rule + BH-FDR (C7)", "06, 07"),
    ("3", "Spectral separability", "Thresholding (Mahlein et al., 2013), OA, classification scalograms", "Same thresholding, balanced accuracy (C2), scalograms of R, D1–D3, CWT", "08–11"),
    ("4", "DSSF extraction", "Top 5 % per DAI → intersection → 1 per region; top-10 SIs per year → union", "Same with explicit consistency level k* (C3); scalar families top 10 % union", "12"),
    ("5", "Collinearity", "VIF < 10", "Iterative relevance-aware VIF < 10 (C5)", "13, 14"),
    ("6A", "ML-SFFS (paper-faithful)", "LDA / k-NN / SVM-SFFS, 4-fold CV × 100, OA, most frequent OFC, SVM-SFFS vs SVM-all", "Identical (C6 scaling details)", "15–17"),
    ("6B", "ML-SFFS (corrected)", "—", "Nested CV: pipeline of Steps 3–6 re-run inside every outer training fold, balanced accuracy (C4)", "18"),
    ("6C", "Permutation null", "—", "Procedure A repeated with shuffled plot labels (Ojala and Garriga, 2010)", "19"),
    ("6D", "Infection stages", "OA per DAI (asymptomatic, early, mild)", "Controls vs each inoculation timing D1–D4 on each date T1–T4", "20"),
    ("7", "Cross-checks", "—", "Excel formula checks, independent recomputation, reference verification", "21–24"),
]
r = table(ws, r, pd.DataFrame(wf, columns=["Step", "Module", "Tian et al. (2021)", "This study", "Sheets"]), widths=[7, 24, 60, 75, 10], wrap_cols=("Tian et al. (2021)", "This study"))
r = image(ws, r, f"{FIG}/fig01_workflow.png", 620)

# ================================================================ 02 / 03
ws, r = sheet("02_Paper_vs_Study", "Paper vs this study", ["Design and method of Tian et al. (2021) beside the corresponding choice in this study."])
r = table(ws, r, pd.DataFrame(PAPER_VS_STUDY, columns=["Aspect", "Tian et al. (2021)", "This study"]), widths=[22, 60, 80], wrap_cols=("Tian et al. (2021)", "This study"))
ws, r = sheet("03_Corrections", "Corrections to the paper's procedure (C1–C9)",
              ["Each correction keeps the paper's logic but fixes a problem that the white-rust data (or the paper's own reporting) raised. The paper-faithful result is always reported as well."])
corr = sorted(CORRECTIONS, key=lambda x: int(x[0][1:]))
r = table(ws, r, pd.DataFrame(corr, columns=["Code", "Issue", "Tian et al. (2021)", "Correction applied", "Reason"]),
          widths=[7, 26, 45, 70, 55], wrap_cols=("Issue", "Tian et al. (2021)", "Correction applied", "Reason"))

# ================================================================ 04 QC
ws, r = sheet("04_Step1_Data_QC", "Step 1 · Data preparation and quality control",
              ["Concept: the experimental unit is the plot; each plot spectrum is the mean of its leaf spectra (3–15 leaves), which reduces leaf-to-leaf and instrument noise and avoids pseudo-replication.",
               "Method: labels, varieties and treatment were cross-checked between the four workbooks and the treatment file; spectral range 400–2400 nm without the ASD detector-join guard zones 990–1010 and 1795–1835 nm (C1); continuous wavelet transform (Mexican hat, dyadic scales 2¹–2⁸, scales 3–8 retained; Cheng et al., 2010) computed from the raw plot-average reflectance.",
               "Note: the 12 control plots form their own block (C) and were measured first on every date (spectra 1–100); disease is therefore partly confounded with field position and measurement time (C9)."])
r = section(ws, r, "Plots, leaves and growth stages per date")
r = table(ws, r, T["qc"], widths=[14, 10, 10, 10, 12, 12, 30, 30, 34, 16, 12, 12])
r = section(ws, r, "Number of candidate features per family (after range restriction)")
r = table(ws, r, T["feature_counts"])
r = section(ws, r, "Treatment design: number of plots per block × variety × inoculation timing")
dz = T["design"].copy(); dz.columns = [f"{a} {b}" for a, b in dz.columns]; dz = dz.reset_index().rename(columns={"Replication": "Block"})
r = table(ws, r, dz)
r = image(ws, r, f"{FIG}/fig02_mean_spectra.png", 820)

ws, r = sheet("05_Step1_Treatment", "Step 1 · Treatment file and growth stage of every inoculation group",
              ["Inoculation timing D1–D4 = inoculation on different days after sowing (D1 assumed earliest). Block C = uninoculated controls; R1–R3 = replications."])
r = table(ws, r, T["treatment"], widths=[8, 26, 10, 18])
r = section(ws, r, "Growth stage of the plots of each inoculation timing on each date")
r = table(ws, r, T["stage_by_inoc"])

# ================================================================ 06 / 07 physiology
ph = T["physio"].copy()
ph = ph[["date_label", "proxy", "constituent", "reference", "n_H", "n_D", "outliers_removed", "normal", "equal_var", "test", "statistic", "p", "q_BH",
         "significant (q<0.05)", "mean_H", "sd_H", "mean_D", "sd_D", "hedges_g", "direction"]]
ph.columns = ["Date", "Proxy", "Constituent", "Reference", "n H", "n D", "Outliers removed", "Normal (SW)", "Equal var (Levene)", "Test", "Statistic",
              "p", "q (BH)", "q < 0.05", "Mean H", "SD H", "Mean D", "SD D", "Hedges g", "Direction"]
ws, r = sheet("06_Step2_Physiology", "Step 2 · Physiological interpretation (analogue of Tian et al. Fig. 4)",
              ["Concept: Tian et al. showed that infected leaves lose chlorophyll, carotenoids and water and gain anthocyanins, which explains why pigment- and water-sensitive features detect the disease. Leaf biochemistry was not measured here, so published index proxies of the same constituents are compared (C8).",
               "Method (as in the paper): outliers outside median ± 3 MAD removed (Leys et al., 2013); Student t-test if both groups are normal (Shapiro–Wilk) with equal variances (Levene), otherwise Mann–Whitney U; Benjamini–Hochberg FDR over the 32 tests (C7). Hedges g > 0: higher in diseased plots."])
r = table(ws, r, ph, fmt={"p": "0.0E+00", "q (BH)": "0.0E+00"}, widths=[13, 30, 22, 22] + [9] * 16)
r = image(ws, r, f"{FIG}/fig06_physiology_boxplots.png", 950)
st = T["scalar_tests"].copy(); st["date"] = st.date.map(DLAB); st["family"] = st.family.map(FAM_NAME)
st.columns = ["Date", "Family", "Feature", "Mean H", "Mean D", "U", "p", "AUC", "q (BH within family & date)"]
ws, r = sheet("07_Step2_All_Scalar_Tests", "Step 2 (supplement) · Mann–Whitney U test of every scalar feature",
              ["Healthy vs diseased plots, two-sided Mann–Whitney U, AUC direction-free, BH-FDR within each family and date."])
r = table(ws, r, st, fmt={"p": "0.0E+00", "q (BH within family & date)": "0.0E+00"}, widths=[13, 26, 50, 12, 12, 8, 10, 8, 12])

# ================================================================ 08-11 separability
ws, r = sheet("08_Step3_Sep_Scalars", "Step 3 · Threshold separability of every vegetation index, CR parameter and red-edge parameter",
              ["Concept (Tian et al. Fig. 3; Mahlein et al., 2013): the interval between the class means Mh (healthy) and Mi (infected) is divided into 100 equal steps; every step is used as a threshold and the best one is kept. Separability = balanced accuracy (BA) at that threshold (C2); OA is shown for comparison.",
               "Top10% = in the top 10 % of its family on that date (ties kept). The union of the top-10 % sets over dates enters the DSSF pool (Step 4)."])
r = table(ws, r, T["sep_scalar"], widths=[24, 46] + [10] * 30)
ws, r = sheet("09_Step3_Sep_Bands_Derivs", "Step 3 · Threshold separability of every reflectance band and derivative (400–2400 nm, guard zones excluded)",
              ["One row per wavelength; Top5% = among the top 5 % of that family on that date; 'Dates in top 5%' counts the dates; 'Common' = selected at the consistency level k* (C3)."])
long = []
for f in ["R", "D1", "D2", "D3"]:
    d_ = T["sep_spectral"][f].copy(); d_.insert(0, "Family", FAM_NAME[f]); long.append(d_)
r = table(ws, r, pd.concat(long, ignore_index=True), widths=[20, 12] + [10] * 12)
ws, r = sheet("10_Step3_Sep_CWT", "Step 3 · Threshold separability of every wavelet coefficient (scales 3–8)", ["Scale j corresponds to the dyadic scale 2^j (Cheng et al., 2010)."])
r = table(ws, r, T["sep_cwt"], widths=[8, 12] + [11] * 5)
ws, r = sheet("11_Step3_Scalograms", "Step 3 · Classification scalograms (analogues of Tian et al. Figs. 6 and 7)",
              ["Shade = balanced accuracy of the feature; orange ticks = top 5 % of the date; bottom row = features common to k* dates; ▼ = region representative kept in the DSSF pool."])
for f in ["R", "D1", "D2", "D3"]:
    r = image(ws, r, f"{FIG}/fig04_scalogram_{f}.png", 980)
r = image(ws, r, f"{FIG}/fig05_scalogram_CWT.png", 980)
r = image(ws, r, f"{FIG}/fig03_threshold_concept.png", 560)

# ================================================================ 12-14 DSSF / VIF
ws, r = sheet("12_Step4_DSSF_Selection", "Step 4 · Extraction of the disease-specific spectral features (DSSFs)",
              ["Spectral families (R, D1–D3, CWT): top 5 % per date → features common to k* dates (k* = largest of 4, 3, 2 with ≥ 1 common feature; C3) → contiguous regions → one representative per region (highest mean BA).",
               "Scalar families (VI, CR, RE): top 10 % per date → union over the four dates (as the union of the yearly top-10 SIs in Tian et al.). Exact duplicates (|r| > 0.9999) removed before the VIF."])
r = section(ws, r, "Selection summary per family")
r = table(ws, r, T["selection_summary"])
r = section(ws, r, "Candidate DSSFs (before VIF)")
dc = T["dssf_candidates"][["Family", "feature", "region", "region_size", "dates_in_top", "mean_BA", "Duplicate_removed", "In_final_pool", "In_plain_VIF_pool"]]
dc.columns = ["Family", "Feature", "Region", "Region size", "Dates in top", "Mean BA", "Duplicate removed", "In final pool (relevance-aware VIF)", "In pool of plain VIF rule"]
r = table(ws, r, dc, widths=[28, 50, 22, 10, 10, 10, 12, 16, 16])
vl = T["vif_log"]
ws, r = sheet("13_Step5_VIF", "Step 5 · Collinearity reduction with the variance inflation factor",
              ["Concept: VIF_j = 1/(1 − R²_j), R²_j from regressing feature j on all other features; VIF ≥ 10 signals strong collinearity (O'Brien, 2007). Computed on the features z-scored within each date and stacked over the four dates.",
               f"Rule (C5): while any VIF ≥ 10, remove the least separable of the features with VIF ≥ 10. Result: {len(T['dssf_candidates']) - len(T['dssf_candidates'][T['dssf_candidates'].Duplicate_removed])} candidates → {len(POOL)} DSSFs (plain highest-VIF rule: {int(T['dssf_candidates'].In_plain_VIF_pool.sum())})."])
r = table(ws, r, vl, widths=[10, 12, 12, 50, 14, 14])
r = image(ws, r, f"{FIG}/fig07_vif.png", 700)
ws, r = sheet("14_Step5_Final_DSSF", "Step 5 · Final disease-specific spectral features (analogue of Tian et al. Table 2)",
              ["Codes F1–F16 are used in all later sheets. BA per date = threshold separability (Step 3)."])
r = table(ws, r, T["dssf_final"], widths=[7, 46, 30, 22, 9, 9, 9, 9, 90, 11, 11, 11, 11], wrap_cols=("Meaning",))

# ================================================================ 17 repetitions (written before 15 for formulas)
rep = F[["date", "clf", "rep", "OFC codes", "k", "OA", "BA", "kappa", "sens", "spec", "all_OA", "all_BA", "params"]].copy()
rep["date"] = rep.date.map(DLAB); rep["params"] = rep.params.apply(lambda p: f"C={p[0]:g}, γ·p={p[1]:g}" if p and p != (1.0, 1.0) else "–")
rep.columns = ["Date", "Classifier", "Repetition", "OFC", "No. features", "OA", "BA", "kappa", "Sensitivity", "Specificity", "All-DSSF OA", "All-DSSF BA", "SVM parameters"]
ws17 = wb.create_sheet("S17_Step6A_Repetitions")
ws17["A1"] = "Step 6A · Every repetition of the paper-faithful ML-SFFS (4-fold CV partition re-drawn in each repetition)"; ws17["A1"].font = F_TITLE
table(ws17, 3, rep, widths=[13, 10, 10, 26, 10, 9, 9, 9, 10, 10, 11, 11, 18])
NREP = len(rep)

fa = T["faithful"]
ws, r = sheet("15_Step6A_Faithful", "Step 6A · Paper-faithful ML-SFFS (analogue of Tian et al. Table 3 and Fig. 8)",
              ["Concept: SFFS (Pudil et al., 1994) adds the most significant feature, then conditionally removes the least significant one while this improves the criterion; the criterion is the 4-fold CV accuracy of LDA, k-NN (k = 3) or SVM-RBF (grid search of C and γ). Repeated 100 times with new random partitions; OA = mean of the 100 optimal accuracies; OFC = most frequent optimal feature combination.",
               "Green columns are live formulas (AVERAGEIFS over sheet 17) and must equal the Python values beside them.",
               "Note: this is the paper's estimate. Because features and parameters are selected on the same plots whose accuracy is reported, it is optimistic – see sheets 18 (nested CV) and 19 (permutation null)."])
start = r
cols = ["Date", "Classifier", "Most frequent OFC", "OFC frequency (%)", "Distinct OFCs", "Mean no. of features", "Mean OA", "SD OA", "Mean BA", "Mean kappa", "Mean sensitivity", "Mean specificity"]
r = table(ws, r, fa[cols], fmt={"OFC frequency (%)": pct, "Mean no. of features": "0.00"}, widths=[13, 10, 26, 11, 10, 11, 9, 9, 9, 9, 11, 11, 12, 12, 12, 12])
hdr = ["Mean OA (formula)", "Mean BA (formula)", "Mean kappa (formula)", "Δ OA (formula − Python)"]
for j, h in enumerate(hdr):
    c_ = ws.cell(start, 13 + j, h); c_.font = F_H; c_.fill = FILL_H; c_.alignment = Alignment(wrap_text=True)
rng = lambda col: f"S17_Step6A_Repetitions!${col}$4:${col}${3 + NREP}"
for i in range(len(fa)):
    rr = start + 1 + i
    crit = f'{rng("A")},$A{rr},{rng("B")},$B{rr}'
    ws.cell(rr, 13, f"=AVERAGEIFS({rng('F')},{crit})")
    ws.cell(rr, 14, f"=AVERAGEIFS({rng('G')},{crit})")
    ws.cell(rr, 15, f"=AVERAGEIFS({rng('H')},{crit})")
    ws.cell(rr, 16, f"=M{rr}-G{rr}")
    for j in range(13, 17):
        ws.cell(rr, j).fill = FILL_F; ws.cell(rr, j).number_format = "0.000" if j < 16 else "0.0E+00"; ws.cell(rr, j).font = F_B; ws.cell(rr, j).border = BOX
r = section(ws, r, "SFFS-selected features vs all DSSFs (paper Fig. 8: SVM-SFFS vs SVM-all)")
r = table(ws, r, T["sffs_vs_all"], fmt={"Gain OA (pp)": pct, "Gain BA (pp)": pct})
r = image(ws, r, f"{FIG}/fig08_sffs_vs_all.png", 980)
ws, r = sheet("16_Step6A_OFC_Detail", "Step 6A · Optimal feature combinations and selection frequency",
              ["Top five OFCs of every date × classifier, and the percentage of the 100 repetitions in which each DSSF was part of the OFC."])
r = table(ws, r, T["faithful_top_ofc"], widths=[13, 10, 26, 12, 12, 12])
r = section(ws, r, "Selection frequency (%) of each DSSF")
r = table(ws, r, T["faithful_freq"], fmt={c_: "0" for c_ in T["faithful_freq"].columns[2:]})
r = image(ws, r, f"{FIG}/fig10_selection_frequency.png", 900)

# ================================================================ 18 nested
ne = T["nested"]; nr = T["nested_reps"]
ws, r = sheet("18_Step6B_Nested", "Step 6B · Corrected estimate: nested cross-validation (C4)",
              ["Concept: an accuracy estimate is unbiased only if the test plots play no part in choosing the features, the feature subset or the classifier parameters (Ambroise and McLachlan, 2002; Varma and Simon, 2006; Cawley and Talbot, 2010).",
               "Method: 10 repetitions × stratified 4-fold outer CV over plots (the same test plots on every date). Inside each outer training set the whole pipeline – threshold separability, DSSF pool, VIF, SFFS (balanced-accuracy criterion, inner 4-fold CV × 3) and SVM grid search – is re-run; the selected model then classifies the held-out plots. LDA uses equal priors, the SVM class-balanced weights.",
               "Green columns are live AVERAGEIFS formulas over the per-repetition table below."])
start = r
cols = ["Date", "Classifier", "OA", "SD OA", "BA", "SD BA", "kappa", "Sensitivity", "Specificity", "All-DSSF OA", "All-DSSF BA", "Mean no. of features", "Inner-CV BA (optimistic)"]
r = table(ws, r, ne[cols], widths=[13, 10, 9, 9, 9, 9, 9, 11, 11, 11, 11, 11, 12, 12, 12])
nr_start = r + 1
for j, h in enumerate(["BA (formula)", "OA (formula)"]):
    c_ = ws.cell(start, 14 + j, h); c_.font = F_H; c_.fill = FILL_H; c_.alignment = Alignment(wrap_text=True)
NN = len(nr)
for i in range(len(ne)):
    rr = start + 1 + i
    crit = f"$A${nr_start + 1}:$A${nr_start + NN},$A{rr},$B${nr_start + 1}:$B${nr_start + NN},$B{rr}"
    ws.cell(rr, 14, f"=AVERAGEIFS($E${nr_start + 1}:$E${nr_start + NN},{crit})")
    ws.cell(rr, 15, f"=AVERAGEIFS($D${nr_start + 1}:$D${nr_start + NN},{crit})")
    for j in (14, 15):
        ws.cell(rr, j).fill = FILL_F; ws.cell(rr, j).number_format = "0.000"; ws.cell(rr, j).font = F_B; ws.cell(rr, j).border = BOX
r = section(ws, r, "Per-repetition results (pooled predictions of the 4 outer folds)")
nr2 = nr[["Date", "Classifier", "Repetition", "OA", "BA", "kappa", "sens", "spec", "All-DSSF BA"]].rename(columns={"sens": "Sensitivity", "spec": "Specificity"})
assert r == nr_start
r = table(ws, r, nr2)
r = section(ws, r, "Optimism of the paper-faithful estimate (A − B)")
r = table(ws, r, T["optimism"], fmt={"Optimism OA (pp)": pct, "Optimism BA (pp)": pct})
r = section(ws, r, "Stability: how often each feature entered the DSSF pool / the OFC across the 40 outer folds")
r = table(ws, r, T["nested_stability"], fmt={"In pool of outer folds (%)": pct, "Selected by SFFS (% of fold x date x classifier fits)": pct})
r = image(ws, r, f"{FIG}/fig09_faithful_vs_nested_vs_null.png", 980)

# ================================================================ 19 permutation
ws, r = sheet("19_Step6C_Permutation", "Step 6C · Permutation null of the paper-faithful procedure",
              ["Method (Ojala and Garriga, 2010): the 48 plot labels were shuffled (the same shuffle on all dates, keeping 12 'healthy' and 36 'diseased'), and the complete paper-faithful pipeline (separability → DSSF pool → VIF → SFFS, 4-fold CV) was re-run. LDA and k-NN: 20 permutations × 5 repetitions; SVM: 4 permutations × 2 repetitions (computational cost; indicative only).",
               "p = (1 + number of null values ≥ observed)/(1 + number of permutations); with 20 permutations the smallest attainable p is 0.048, with 4 permutations 0.20.",
               "Reading: with labels that carry no information the paper-faithful procedure still reports OA ≈ 85–92 % and BA ≈ 73–84 %; this is the optimism that nested CV (sheet 18) removes."])
r = table(ws, r, T["perm"], fmt={"p (OA)": "0.000"})
pr = T["perm_raw"].copy(); pr["date"] = pr.date.map(DLAB)
r = section(ws, r, "All permutation runs")
r = table(ws, r, pr.rename(columns={"perm": "Permutation", "date": "Date", "clf": "Classifier", "rep": "Repetition", "k": "No. features", "pool_size": "Null pool size"}))
r = image(ws, r, f"{FIG}/fig11_permutation_null.png", 980)

# ================================================================ 20 stage
ws, r = sheet("20_Step6D_Stage", "Step 6D · Detection by infection stage (date T1–T4 × inoculation timing D1–D4)",
              ["A (paper-faithful, analogue of Table 3 by DAI): 12 control plots vs the 9 plots of one inoculation timing (3 varieties × 3 replications), fixed DSSF pool F1–F16, ML-SFFS with 4-fold CV × 100 repetitions.",
               "B (unbiased): share of plots of each group called 'diseased' by the nested-CV models of Step 6B (controls: false-alarm rate; inoculated groups: detection rate), and correct classification per variety.",
               "Note: in A the DSSF pool was selected on all plots, so A is optimistic here as well."])
r = section(ws, r, "A · Paper-faithful ML-SFFS per stage")
r = table(ws, r, T["stageA"].assign(Date=T["stageA"].Date.astype(str)), fmt={"OFC frequency (%)": pct, "Mean no. of features": "0.00"}, widths=[13, 12, 10, 9, 9, 9, 9, 10, 24, 10, 10])
r = section(ws, r, "B · Detection rate per group (nested CV)")
r = table(ws, r, T["stage_detection"].assign(Date=T["stage_detection"].Date.astype(str)), fmt={"Called diseased (%)": pct, "Correct (%)": pct})
r = section(ws, r, "B · Correct classification per variety (nested CV)")
r = table(ws, r, T["variety_detection"].assign(Date=T["variety_detection"].Date.astype(str)), fmt={"Correct (%)": pct})
r = image(ws, r, f"{FIG}/fig12_stage_wise.png", 980)

# ================================================================ 21 DSSF values + 22 formula check
vals = []
for d in DATES:
    X = date_matrix(feats, d, POOL)
    for p in X.index:
        row = {"Date": DLAB[d], "Plot": int(p), "Variety": TR.loc[p, "Variety_t"], "Inoculation": TR.loc[p, "Inoculation"],
               "Block": TR.loc[p, "Replication"], "Class": "Healthy" if y[d][list(X.index).index(p)] == 0 else "Diseased"}
        for n in POOL:
            row[CODE[n]] = float(X.loc[p, n])
        vals.append(row)
V = pd.DataFrame(vals)
ws, r = sheet("21_Step7_DSSF_Values", "Step 7 · Values of the 16 DSSFs for every plot and date", ["Feature codes as in sheet 14. Used by the formula checks in sheet 22."])
v_start = r
r = table(ws, r, V, fmt={CODE[n]: "0.000000" for n in POOL}, widths=[13, 7, 8, 11, 7, 10] + [12] * 16)
NV = len(V)
ws, r = sheet("22_Step7_Formula_Check", "Step 7 · Excel recomputes the class means and the threshold balanced accuracy of every DSSF",
              ["Python values (white) are compared with live Excel formulas (green): AVERAGEIFS for the class means, COUNTIFS at the Python threshold for sensitivity and specificity. Differences must be ≈ 0.",
               "Diseased plots are those on the side of the threshold where the diseased mean lies (> threshold if Mi > Mh, otherwise <)."])
hdr = ["Date", "Code", "Feature", "Python Mh", "Python Mi", "Python threshold", "Python BA", "Mh (formula)", "Mi (formula)", "Sensitivity (formula)",
       "Specificity (formula)", "BA (formula)", "Δ BA", "Δ Mi"]
for j, h in enumerate(hdr):
    c_ = ws.cell(r, 1 + j, h); c_.font = F_H; c_.fill = FILL_H; c_.alignment = Alignment(wrap_text=True); c_.border = BOX
setw(ws, [13, 7, 40, 12, 12, 12, 10, 12, 12, 12, 12, 10, 10, 10])
fam_of = T["dssf_candidates"].set_index("feature").family
vs = "S21_Step7_DSSF_Values!"
rngA = f"{vs}$A${v_start + 1}:$A${v_start + NV}"; rngF = f"{vs}$F${v_start + 1}:$F${v_start + NV}"
POOLP = pickle.load(open(f"{O}/pool_full.pkl", "rb"))
i = 0
for d in DATES:
    for n in POOL:
        i += 1; rr = r + i
        s_ = POOLP["sep"][(d, fam_of[n])].loc[n]
        col = get_column_letter(7 + POOL.index(n))
        rv = f"{vs}${col}${v_start + 1}:${col}${v_start + NV}"
        vals_ = [DLAB[d], CODE[n], n, float(s_.Mh), float(s_.Mi), float(s_.threshold), float(s_.BA)]
        for j, v in enumerate(vals_):
            c_ = ws.cell(rr, 1 + j, v); c_.font = F_B; c_.border = BOX
            if j >= 3: c_.number_format = "0.000000" if j < 6 else "0.000"
        ws.cell(rr, 8, f'=AVERAGEIFS({rv},{rngA},$A{rr},{rngF},"Healthy")')
        ws.cell(rr, 9, f'=AVERAGEIFS({rv},{rngA},$A{rr},{rngF},"Diseased")')
        ws.cell(rr, 10, f'=IF(I{rr}>H{rr},COUNTIFS({rngA},$A{rr},{rngF},"Diseased",{rv},">"&F{rr}),COUNTIFS({rngA},$A{rr},{rngF},"Diseased",{rv},"<"&F{rr}))/COUNTIFS({rngA},$A{rr},{rngF},"Diseased")')
        ws.cell(rr, 11, f'=IF(I{rr}>H{rr},COUNTIFS({rngA},$A{rr},{rngF},"Healthy",{rv},"<="&F{rr}),COUNTIFS({rngA},$A{rr},{rngF},"Healthy",{rv},">="&F{rr}))/COUNTIFS({rngA},$A{rr},{rngF},"Healthy")')
        ws.cell(rr, 12, f"=(J{rr}+K{rr})/2")
        ws.cell(rr, 13, f"=L{rr}-G{rr}")
        ws.cell(rr, 14, f"=I{rr}-E{rr}")
        for j in range(8, 15):
            c_ = ws.cell(rr, j); c_.fill = FILL_F; c_.font = F_B; c_.border = BOX
            c_.number_format = "0.000000" if j in (8, 9) else ("0.0E+00" if j >= 13 else "0.000")
last_check_row = r + i
ws.cell(last_check_row + 2, 1, "Largest |Δ BA|:").font = F_BB
ws.cell(last_check_row + 2, 3, f"=SUMPRODUCT(MAX(ABS(M{r + 1}:M{last_check_row})))").fill = FILL_F
ws.cell(last_check_row + 3, 1, "Largest |Δ Mi|:").font = F_BB
ws.cell(last_check_row + 3, 3, f"=SUMPRODUCT(MAX(ABS(N{r + 1}:N{last_check_row})))").fill = FILL_F

# ================================================================ 23 cross-checks
ws, r = sheet("23_Step7_Crosschecks", "Step 7 · Cross-checks of every step",
              ["Every step was checked by an independent route. FLAG = not an error of the analysis but a property of the data or of the paper's procedure that must be reported."])
r = table(ws, r, CK, widths=[7, 60, 50, 90, 9], wrap_cols=("Check", "Independent route", "Result"), flag_col="Status")
r = section(ws, r, "Formula checks inside this workbook")
r = table(ws, r, pd.DataFrame([
    ("Sheet 15", "Mean OA / BA / kappa of procedure A recomputed by AVERAGEIFS from sheet 17"),
    ("Sheet 18", "Mean OA / BA of procedure B recomputed by AVERAGEIFS from the per-repetition table"),
    ("Sheet 22", "Class means and threshold BA of all 16 DSSFs × 4 dates recomputed by AVERAGEIFS / COUNTIFS from sheet 21"),
], columns=["Where", "What"]), widths=[12, 100])

# ================================================================ 24 references
ws, r = sheet("24_References", "References (cross-checked)",
              ["Each reference was checked against a publisher, repository or indexing record (web search, October 2026). The status column states what was confirmed.",
               "Note: the remaining index and red-edge references are those of the source workbooks (sheet 28 there); the ones used for the proxies and final DSSFs are verified here."])
rf = pd.DataFrame(REFS, columns=["Citation", "Reference", "DOI / identifier", "Verification", "Used for"]).sort_values("Citation")
r = table(ws, r, rf, widths=[26, 100, 30, 40, 45], wrap_cols=("Reference", "Verification", "Used for"))
r = section(ws, r, "Errors found in the reference list of Tian et al. (2021)")
r = table(ws, r, pd.DataFrame(TIAN_REF_ERRORS, columns=["Reference", "As printed in Tian et al. (2021)", "Correct"]), widths=[26, 70, 80], wrap_cols=("As printed in Tian et al. (2021)", "Correct"))

wb._sheets.sort(key=lambda w: w.title)
wb.save(OUT)
print("saved", OUT)
