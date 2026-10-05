"""Helpers for appending extracted rows to the meta-analysis workbook."""
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARIAL = Font(name="Arial", size=10)
RED = Font(name="Arial", size=10, color="FFFF0000")
NOTE_FILL = PatternFill("solid", fgColor="FFFFF3CD")


class Book:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path)
        self.hdr = {}

    def headers(self, sheet):
        if sheet not in self.hdr:
            ws = self.wb[sheet]
            h = {}
            for c in ws[1]:
                if c.value is not None and c.value != "UNIT FOR THIS SHEET:" and c.value not in h:
                    h[c.value] = c.column
            self.hdr[sheet] = h
        return self.hdr[sheet]

    def col(self, sheet, name):
        return get_column_letter(self.headers(sheet)[name])

    def next_row(self, sheet):
        ws = self.wb[sheet]
        r = 2
        while ws.cell(r, 1).value is not None or ws.cell(r, 2).value is not None:
            r += 1
        return r

    def add(self, sheet, data, red=False):
        """Append a row; data maps header -> value. Returns the row number."""
        ws = self.wb[sheet]
        h = self.headers(sheet)
        r = self.next_row(sheet)
        for k, v in data.items():
            if v is None:
                continue
            if k not in h:
                raise KeyError(f"{sheet}: no column '{k}'")
            c = ws.cell(r, h[k], v)
        # SD formula on data sheets
        if "SD" in h and "Obs" in h and "Rep" in h and data.get("SD") is None:
            o, p = self.col(sheet, "Obs"), self.col(sheet, "Rep")
            ws.cell(r, h["SD"], f'=IF(AND(N({o}{r})>0,N({p}{r})>0),SQRT(2*{o}{r}/{p}{r}),"")')
        last = max(h.values())
        for ci in range(1, last + 1):
            c = ws.cell(r, ci)
            c.font = copy(RED if red else ARIAL)
        if "Notes/Doubts" in h:
            ws.cell(r, h["Notes/Doubts"]).fill = copy(NOTE_FILL)
        return r

    def ref(self, sheet, name, row, absolute=False):
        q = f"'{sheet}'" if any(ch in sheet for ch in " ()-,.") else sheet
        return f"{q}!{self.col(sheet, name)}{row}"

    def save(self, path):
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(path)
