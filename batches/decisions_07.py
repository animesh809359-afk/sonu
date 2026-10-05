"""Batch 07 rules (2026-10-05): method column, Kukal 2003 exclusion, study-15 companion rows, factorial 2-way tables, duplicate yield means."""
import sys
from copy import copy
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
ROWS = [
    ("Workflow & data management",
     "METHODOLOGY: every data row records the paper's method for that parameter in 'Method used (from paper)' (instrument / extractant / procedure and its citation, "
     "sampling depth and timing where given); 'Not stated in paper' when the paper gives none; derived values state the formula.",
     "Checked 2026-10-05: all 495 rows of updated 06 already carry a method; batch 07 rows filled the same way.", "Author 2026-10-05", "IN FORCE"),
    ("Treatment codes",
     "Shallow / reduced PUDDLING followed by conventionally tilled wheat = rotation mismatch (as old-master 73 ReP-CTW); normal-depth puddling at different intensities = CT.",
     "Kukal & Aggarwal 2003 (170): all treatments puddled rice + CT wheat -> one usable code -> study excluded; data in EXCLUDED_rows.", "Default 2026-10-05 (rules 21, 22)", "PENDING CONFIRMATION"),
    ("Treatment codes",
     "Companion of an old-master trial that used one treatment per code: the paper's other rotation-matched treatments under the same code are added as rows b, c, d (rule 20, T counted in Obs); "
     "puddled rice + ZT wheat with only incidental stubble (residue not a treatment factor) -> pZT.",
     "15 (companion 2) Gathala 2011 SSSAJ: ZT rows a T5, b T6, c T3 (beds), d T4 (beds); T2 (puddled AWD rice + ZT wheat) -> pZT; old-master 15 itself unchanged (T1 vs T5).", "Default 2026-10-05 (rules 15, 20, 68)", "PENDING CONFIRMATION"),
    ("Pooled means",
     "Tillage x residue x N factorial printed only as two-way tables: soil properties from the tillage x residue cells (rows a/b, pooled over N, flagged); yields and NUE from the tillage x N cells "
     "(one row per N rate, pooled over residue, flagged); economics from the tillage main effect. The other two-way table goes to Notes - never both.",
     "40 (companion) Kader 2022: CTR vs MTR.", "Default 2026-10-05 (rule 33)", "PENDING CONFIRMATION"),
    ("Workflow & data management",
     "A multi-year MEAN that summarises year-wise values already entered for an old-master study is entered only with a 'DUPLICATE FLAG' in Notes (use it only for the new contrasts, or drop it at merge).",
     "15 (companion 2) Table 6 7-yr mean yields (T1/T5 year-wise yields are in old-master 15).", "Default 2026-10-05", "PENDING CONFIRMATION"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("rules", n0 + 1, "-", n0 + len(ROWS))
