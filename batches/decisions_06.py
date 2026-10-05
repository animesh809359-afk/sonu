"""Author instruction 2026-10-05: enter every crop-stage-wise value of any parameter (one row per stage) -> updated_06."""
import sys
from copy import copy
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 1).value == 38:
        ws.cell(r, 6).value = "IN FORCE - extended by rule %d (all parameters, soil and plant)" % (n0 + 1)
row = (n0 + 1, "Seasons, years & rows",
       "STAGE-WISE DATA: when any parameter (soil OR plant - e.g. enzymes, MBC, nutrients, moisture, BD, penetration resistance, roots, plant height, dry matter, "
       "nutrient uptake, LAI) is reported at several crop growth stages, EVERY stage is entered as its own row, with the stage (and DAS/DAT) named at the start of "
       "'Crop/season of sampling'.",
       "Applies to wheat-season stages by default and to rice-season stages under the rice-season fallback (red rows). Stage rows are not averaged; a paper's "
       "own season mean goes to Notes. Obs per row as rule 45 (each stage row has Y = 1 unless it pools years). Checked 2026-10-05: none of the 10 papers entered so far "
       "(serials 164-169, 23 (companion 3/4), 131 (companion), 165 (companion)) reports stage-wise values of an extracted parameter.",
       "Author 2026-10-05", "IN FORCE")
r = ws.max_row + 1
for j, v in enumerate(row, 1):
    ws.cell(r, j, v)._style = copy(style[j - 1])
rd = wb["README"]
r = rd.max_row + 1
rd.cell(r, 1, "Stage-wise data (2026-10-05)").font = openpyxl.styles.Font(name="Arial", size=10, bold=True)
c = rd.cell(r, 2, "updated 06: author instruction - every crop-stage-wise value of any soil or plant parameter is entered as its own row (rule %d). "
                  "Existing studies re-checked: no stage-wise data were left out. No data values changed." % (n0 + 1))
c.font = openpyxl.styles.Font(name="Arial", size=10)
c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("rule", n0 + 1)
