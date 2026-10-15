"""Small openpyxl helper layer: consistent styles, tables, text blocks,
images and formula-driven error matrices."""
from __future__ import annotations

import math

import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
DARK = "1F3D2B"
F_TITLE = Font(name=FONT, size=14, bold=True, color=DARK)
F_SUB = Font(name=FONT, size=10, italic=True, color="52514E")
F_H2 = Font(name=FONT, size=11, bold=True, color=DARK)
F_HEAD = Font(name=FONT, size=9, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=9)
F_BOLD = Font(name=FONT, size=9, bold=True)
F_NOTE = Font(name=FONT, size=8.5, italic=True, color="52514E")
F_INPUT = Font(name=FONT, size=9, color="0000FF")
FILL_HEAD = PatternFill("solid", fgColor=DARK)
FILL_KEY = PatternFill("solid", fgColor="E3F4EC")
FILL_FLAG = PatternFill("solid", fgColor="FDE2C5")
FILL_SEC = PatternFill("solid", fgColor="F0EFEC")
THIN = Side(style="thin", color="C9C8C1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


class Sheet:
    """Writes top-to-bottom; keeps the current row."""

    def __init__(self, wb, name, title, subtitle="", width=12, colw=None):
        self.ws = wb.create_sheet(name)
        self.name = name
        self.row = 1
        self.width = width
        self.ws.sheet_view.showGridLines = False
        for c in range(1, 40):
            self.ws.column_dimensions[get_column_letter(c)].width = 14
        if colw:
            for k, v in colw.items():
                self.ws.column_dimensions[k].width = v
        self.cell(1, 1, title, F_TITLE)
        self.row = 2
        if subtitle:
            self.para(subtitle, font=F_SUB)
        self.row += 1

    def cell(self, r, c, v, font=F_BODY, fmt=None, fill=None, align=None, border=False):
        if isinstance(v, (np.floating,)):
            v = float(v)
        if isinstance(v, (np.integer,)):
            v = int(v)
        if isinstance(v, (np.bool_,)):
            v = bool(v)
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            v = None
        x = self.ws.cell(row=r, column=c, value=v)
        x.font = font
        if fmt:
            x.number_format = fmt
        if fill:
            x.fill = fill
        if align:
            x.alignment = align
        if border:
            x.border = BORDER
        return x

    def h2(self, text):
        self.row += 1
        self.cell(self.row, 1, text, F_H2)
        self.row += 1

    def para(self, text, font=F_BODY, width=None, fill=None):
        """Paragraph merged across `width` columns with wrapping."""
        w = width or self.width
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=w)
        c = self.cell(self.row, 1, text, font, align=WRAP, fill=fill)
        chars_per_line = max(40, int(w * 14 * 1.15))
        lines = sum(max(1, math.ceil(len(seg) / chars_per_line)) for seg in str(text).split("\n"))
        self.ws.row_dimensions[self.row].height = max(13, 12.5 * lines + 2)
        self.row += 1
        return c

    def bullets(self, items, title=None):
        if title:
            self.cell(self.row, 1, title, F_BOLD)
            self.row += 1
        for it in items:
            self.para("• " + it)

    def table(self, df, title=None, fmts=None, note=None, col=1, widths=None, key_rows=None, flag_col=None,
              max_rows=None):
        """Write a DataFrame; returns (first data row, last data row, {colname: letter})."""
        fmts = fmts or {}
        if title:
            self.cell(self.row, col, title, F_BOLD)
            self.row += 1
        hr = self.row
        cols = list(df.columns)
        for j, cn in enumerate(cols):
            self.cell(hr, col + j, str(cn), F_HEAD, fill=FILL_HEAD, align=CENTER, border=True)
        self.ws.row_dimensions[hr].height = max(15, 12 * math.ceil(max(len(str(c)) for c in cols) / 14))
        r = hr + 1
        data = df if max_rows is None else df.head(max_rows)
        for i, (_, rec) in enumerate(data.iterrows()):
            for j, cn in enumerate(cols):
                v = rec[cn]
                f = fmts.get(cn)
                if f is None and isinstance(v, (float, np.floating)):
                    av = abs(v) if v == v else 0
                    f = "0.000E+00" if (0 < av < 1e-3 or av >= 1e6) else ("0.0000" if av < 10 else "0.00")
                fill = FILL_KEY if key_rows and i in key_rows else None
                if flag_col and cn == flag_col and str(v).upper().startswith("FLAG"):
                    fill = FILL_FLAG
                if isinstance(v, (list, tuple, dict)):
                    v = str(v)
                self.cell(r, col + j, v, F_BODY, fmt=f, fill=fill, border=True,
                          align=Alignment(wrap_text=isinstance(v, str) and len(v) > 18, vertical="top"))
            r += 1
        letters = {cn: get_column_letter(col + j) for j, cn in enumerate(cols)}
        first, last = hr + 1, r - 1
        self.row = r
        if note:
            self.para(note, font=F_NOTE)
        self.row += 1
        if widths:
            for k, v in widths.items():
                self.ws.column_dimensions[k].width = v
        return first, last, letters

    def image(self, path, width_px=900, col="A", rows_per_100px=5.2):
        img = XLImage(str(path))
        ratio = width_px / img.width
        img.width, img.height = int(img.width * ratio), int(img.height * ratio)
        self.ws.add_image(img, f"{col}{self.row}")
        self.row += int(math.ceil(img.height / 100 * rows_per_100px)) + 2

    def kv(self, rows, title=None, fmts=None):
        """Two-column key/value table; values may be formulas."""
        df = pd.DataFrame(rows, columns=["Item", "Value"])
        return self.table(df, title=title, fmts=fmts)


def error_matrix(S, m, title, cls=("Healthy", "Diseased"), note=None):
    """Das et al. Table 2 layout (rows = classified, columns = reference) with
    every total and accuracy computed by Excel formulas from the four counts."""
    ws = S.ws
    S.cell(S.row, 1, title, F_BOLD)
    S.row += 1
    r0 = S.row
    hdr = ["Classified \\ Reference", cls[0], cls[1], "Row total", "User's accuracy (%)"]
    for j, h in enumerate(hdr):
        S.cell(r0, 1 + j, h, F_HEAD, fill=FILL_HEAD, align=CENTER, border=True)
    a, b = r0 + 1, r0 + 2
    S.cell(a, 1, cls[0], F_BOLD, border=True)
    S.cell(a, 2, m["HH"], F_INPUT, border=True)
    S.cell(a, 3, m["DH"], F_INPUT, border=True)
    S.cell(b, 1, cls[1], F_BOLD, border=True)
    S.cell(b, 2, m["HD"], F_INPUT, border=True)
    S.cell(b, 3, m["DD"], F_INPUT, border=True)
    for rr in (a, b):
        S.cell(rr, 4, f"=B{rr}+C{rr}", border=True)
        S.cell(rr, 5, f"=IF(D{rr}=0,\"\",100*{'B' if rr == a else 'C'}{rr}/D{rr})", fmt="0.00", border=True)
    c = b + 1
    S.cell(c, 1, "Column total", F_BOLD, border=True)
    S.cell(c, 2, f"=B{a}+B{b}", border=True)
    S.cell(c, 3, f"=C{a}+C{b}", border=True)
    S.cell(c, 4, f"=D{a}+D{b}", border=True)
    pa = c + 1
    S.cell(pa, 1, "Producer's accuracy (%)", F_BOLD, border=True)
    S.cell(pa, 2, f"=100*B{a}/B{c}", fmt="0.00", border=True)
    S.cell(pa, 3, f"=100*C{b}/C{c}", fmt="0.00", border=True)
    rows = [("Overall accuracy (%)", f"=100*(B{a}+C{b})/D{c}"),
            ("Balanced accuracy (%)", f"=(B{pa}+C{pa})/2"),
            ("Cohen's κ", f"=((B{a}+C{b})/D{c}-(D{a}*B{c}+D{b}*C{c})/D{c}^2)/(1-(D{a}*B{c}+D{b}*C{c})/D{c}^2)"),
            ("F1 healthy", f"=2*E{a}*B{pa}/(E{a}+B{pa})/100"),
            ("F1 diseased", f"=2*E{b}*C{pa}/(E{b}+C{pa})/100")]
    r = pa + 1
    for k, f in rows:
        S.cell(r, 1, k, F_BOLD, border=True)
        S.cell(r, 2, f, fmt="0.000" if "κ" in k or "F1" in k else "0.00", border=True)
        r += 1
    S.row = r
    if note:
        S.para(note, font=F_NOTE)
    S.row += 1
    return {"OA": f"B{r - 5}", "BA": f"B{r - 4}", "kappa": f"B{r - 3}", "first": a}
