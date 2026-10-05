"""Build an EMPTY meta-analysis workbook from META_ANALYSIS_MASTER_updated_53.xlsx.

- Keeps every sheet, header, colour, unit note, data validation and rule sheet.
- Removes ALL previous data rows (author instruction: do not keep previous data).
- Adds partial-CA treatment columns <prefix>pCA, pZT, pMT, pMTR after every
  <prefix>CA ... <prefix>DTR group on every parameter sheet.
- Adds the new author rules (partial-CA codes, porosity PD fallback, WSA, etc.)
  and a LAT_LONG sheet.

Usage: python build_new_workbook.py <updated_53.xlsx> <output.xlsx>
"""
import sys
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
TODAY = "2026-10-05"
CODES = ["CA", "CT", "ZT", "DT", "MTR", "MT", "CTR", "DTR"]
PCODES = ["pCA", "pZT", "pMT", "pMTR"]
P_FILL = PatternFill("solid", fgColor="FF6B8E23")  # olive green = partial-CA columns
ARIAL = Font(name="Arial", size=10)
ARIAL_B = Font(name="Arial", size=10, bold=True)

# Unit notes that were blank in updated 53 (taken from the units used in its rows / author list)
MISSING_UNITS = {
    "PR": "MPa (soil penetration resistance; 10 cm depth classes 0-10 ... 50-60 CM, >60 CM)",
    "P": "kg/ha (available P; mg/kg x 2.24 when converted - state basis in UNIT column)",
    "total P": "g/kg (total soil P)",
    "ALKP": "ug PNP/g soil/h (alkaline phosphatase)",
    "YIELD": "t/ha (grain and straw; state moisture basis in UNIT column)",
    "P uptake": "kg P/ha",
    "SQI": "unitless (0-1 soil quality index; state method in UNIT column)",
}

wb = openpyxl.load_workbook(SRC)
names = wb.sheetnames
first_data = names.index("BD")
data_sheets = [n for n in names[first_data:] if n != "EXCLUDED_rows"]


def find_groups(hdr):
    """Return [(index_of_DTR_col_1based, base)] for every complete CA..DTR group."""
    groups, i = [], 0
    while i <= len(hdr) - 8:
        h = hdr[i]
        if isinstance(h, str) and h.endswith("CA"):
            base = h[:-2]
            if all(hdr[i + k] == base + CODES[k] for k in range(8)):
                groups.append((i + 8, base))
                i += 8
                continue
        i += 1
    return groups


summary = []
for name in data_sheets:
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    unit_col = hdr.index("UNIT FOR THIS SHEET:") + 1 if "UNIT FOR THIS SHEET:" in hdr else None

    # Save the unit-note column (rows >= 2) - it holds the sheet unit and 'what this sheet holds'
    note_cells = []
    if unit_col:
        for r in range(2, ws.max_row + 1):
            c = ws.cell(r, unit_col)
            if c.value is not None:
                note_cells.append((r, c.value, copy(c.font), copy(c.fill), copy(c.alignment)))
        if name in MISSING_UNITS and not any(r == 2 for r, *_ in note_cells):
            ref = ws.cell(1, unit_col)
            note_cells.insert(0, (2, MISSING_UNITS[name], Font(name="Arial", size=10, bold=True,
                                  color=ref.font.color.rgb if ref.font.color else None),
                                  PatternFill(), Alignment(wrap_text=True, vertical="top")))

    # Remove every previous data row
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)

    widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width}

    # Insert partial-CA columns right-to-left so earlier indices stay valid
    groups = find_groups(hdr)
    for dtr_idx, base in sorted(groups, reverse=True):
        ws.insert_cols(dtr_idx + 1, len(PCODES))
        src = ws.cell(1, dtr_idx)
        for k, pc in enumerate(PCODES):
            c = ws.cell(1, dtr_idx + 1 + k, base + pc)
            c.font = copy(src.font)
            c.alignment = copy(src.alignment)
            c.border = copy(src.border)
            c.fill = copy(P_FILL)
        # shift widths of columns to the right of the insertion
        widths = {
            (get_column_letter(openpyxl.utils.column_index_from_string(k) + len(PCODES))
             if openpyxl.utils.column_index_from_string(k) > dtr_idx else k): w
            for k, w in widths.items()
        }
        if unit_col and unit_col > dtr_idx:
            unit_col += len(PCODES)

    for k in list(ws.column_dimensions.keys()):
        del ws.column_dimensions[k]
    for k, w in widths.items():
        ws.column_dimensions[k].width = w

    # Restore the unit note
    if unit_col:
        for r, v, f, fl, al in note_cells:
            c = ws.cell(r, unit_col, v)
            c.font, c.fill, c.alignment = f, fl, al

    last = ws.max_column
    hdr_last = max(i for i, c in enumerate(ws[1], 1) if c.value not in (None, "UNIT FOR THIS SHEET:"))
    ws.auto_filter.ref = f"A1:{get_column_letter(hdr_last)}1"
    ws.freeze_panes = "D2"
    summary.append((name, [g[1] for g in groups]))

# ---- Study_Info / Treatment_Mapping / EXCLUDED_rows: header only ----------------
for name in ["Study_Info", "Treatment_Mapping", "EXCLUDED_rows"]:
    ws = wb[name]
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)

# ---- Index sheets ---------------------------------------------------------------
ws = wb["PARAMETERS_BY_STUDY"]
ws["A1"] = f"PARAMETERS OF EACH STUDY - auto-generated from the data sheets ({TODAY}); regenerated at every build"
if ws.max_row >= 4:
    ws.delete_rows(4, ws.max_row - 3)

ws = wb["STUDIES_BY_PARAMETER"]
ws["A1"] = (f"STUDIES OF EACH PARAMETER SHEET - auto-generated ({TODAY}); "
            "'NO DATA' = none of the entered studies reports this parameter")
style_row = [copy(c._style) for c in ws[4]]
if ws.max_row >= 4:
    ws.delete_rows(4, ws.max_row - 3)
for i, name in enumerate(data_sheets, 1):
    dws = wb[name]
    hdr = [c.value for c in dws[1]]
    unit = dws.cell(2, hdr.index("UNIT FOR THIS SHEET:") + 1).value if "UNIT FOR THIS SHEET:" in hdr else None
    for j, v in enumerate([i, name, unit, 0, 0, "NO DATA"], 1):
        c = ws.cell(3 + i, j, v)
        c._style = copy(style_row[j - 1])

# ---- Codes: partial-CA codes ------------------------------------------------------
ws = wb["Codes"]
ws.insert_rows(11, 4)
p_rows = [
    ("pCA", "PARTIAL CA: puddled / conventional-till rice + ZERO-TILL wheat WITH residue retained",
     "PTR-ZTW+R, TPR fb ZTW (Happy Seeder / Turbo Happy Seeder into rice residue), puddled rice + ZT wheat with mulch"),
    ("pZT", "PARTIAL ZT: puddled / conventional-till rice + ZERO-TILL wheat, residue REMOVED / burnt",
     "PTR-ZTW, TPR fb ZTW (no residue), puddled rice + ZT drill wheat"),
    ("pMT", "PARTIAL MT: puddled / conventional-till rice + MINIMUM / REDUCED-till wheat, NO residue",
     "PTR + rotary till-drill / Super Seeder / strip-till wheat (rotary <= 8 cm or single pass), no residue"),
    ("pMTR", "PARTIAL MTR: puddled / conventional-till rice + MINIMUM / REDUCED-till wheat WITH residue",
     "PTR + Super Seeder / Roto Seeder / strip-till wheat with rice residue retained"),
]
for k, row in enumerate(p_rows):
    for j, v in enumerate(row, 1):
        c = ws.cell(11 + k, j, v)
        c.font = ARIAL_B if j == 1 else ARIAL
        c.alignment = Alignment(wrap_text=True, vertical="top")

# ---- RULES: new author rules ------------------------------------------------------
ws = wb["RULES"]
ws["A1"] = f"RULES - consolidated extraction, coding and analysis rules of this workbook (author-confirmed unless marked). Updated {TODAY}"
hdr_style = [copy(c._style) for c in ws[4]]
body_style = [copy(c._style) for c in ws[ws.max_row]]
r = ws.max_row + 1
for j in range(1, 7):
    ws.cell(r, j)._style = copy(hdr_style[j - 1])
ws.cell(r, 1, f"NEW WORKBOOK SERIES ({TODAY})")
new_rules = [
    ("Workflow & data management",
     "NEW workbook: built on the updated 53 structure with NO previous data carried over. Serial numbers start at 1. Each upload is saved as a new numbered version (META_ANALYSIS_NEW_updated_NN.xlsx).",
     "README, RULES, Codes and FORMULAS are kept as the rulebook; worked examples there cite the OLD master's serials (updated 53), not studies of this workbook.",
     f"Author {TODAY}", "IN FORCE"),
    ("Treatment codes",
     "PARTIAL CA CODES: pCA = puddled / conventional rice + ZT wheat WITH residue; pZT = puddled / conventional rice + ZT wheat, residue removed; pMT = puddled / conventional rice + MT wheat, no residue; pMTR = puddled / conventional rice + MT wheat WITH residue.",
     "Every parameter sheet carries <prefix>pCA, pZT, pMT, pMTR right after <prefix>DTR (olive-green headers). These treatments are NO LONGER excluded under the rotation-matching rule (rule 22) and are NOT entered in the CA / ZT / MT / MTR columns; they go to the p-columns. The reverse mismatch (ZT / no-till rice + conventional wheat, RNT-WCT) stays EXCLUDED. MT vs MTR for the wheat phase follows the rotary rule (rule 13).",
     f"Author {TODAY}", "IN FORCE"),
    ("Treatment codes",
     "Rotary rule restated: full-width rotary to >= 10 cm, or >= 2 full-width passes = CT / CTR whatever the paper calls it; rotary <= 8 cm, single-pass till-drill, Super / Roto Seeder, strip or zone rotary = MT / MTR.",
     "8-10 cm or depth unstated -> ask the author.", f"Author {TODAY}", "IN FORCE"),
    ("Treatment codes",
     "Same treatment type in rice and wheat = include (e.g. puddled transplanted rice + CT wheat = CT). Rice tillage not stated -> classify on the wheat tillage alone.",
     "Rules 22-23 unchanged except for the partial-CA codes above.", f"Author {TODAY}", "IN FORCE"),
    ("Cropping system & design",
     "Third crop / green manure (mungbean, Sesbania, etc.) in a rice-wheat system: study included; green manuring counts as part of CA (supports CA).",
     "Flag the row 'GREEN MANURE / 3rd CROP' in Notes. If the green manure / third crop is present only in the CA treatments, the row is entered flagged so it can be dropped in a sensitivity analysis.",
     f"Author {TODAY}", "IN FORCE"),
    ("Cropping system & design",
     "Inclusion: any study with ANY TWO of CT, ZT, CA, DT, MT, MTR, CTR, DTR, pCA, pZT, pMT, pMTR is included. Not restricted to South Asia.",
     "Rule 21 extended to the partial-CA codes.", f"Author {TODAY}", "IN FORCE"),
    ("Parameter-specific",
     "Porosity (%) = (1 - BD / PD) x 100 from the TREATMENT BD (changed BD). PD = the treatment's PD if reported, else the paper's initial PD, else 2.65 Mg/m3 (flagged).",
     "Live formula linking the BD sheet (and PD sheet where used). A printed total porosity is entered as printed.", f"Author {TODAY}", "IN FORCE"),
    ("Parameter-specific",
     "WSA: when the paper gives no WSA but gives water-stable aggregate fractions, WSA = sum of the water-stable macro + micro (meso) aggregate fractions (DERIVED, flagged).",
     "Live formula referencing MACRO / MICRO rows where possible.", f"Author {TODAY}", "IN FORCE"),
    ("Parameter-specific",
     "Penetration resistance: 10 cm classes 0-10, 10-20, 20-30, 30-40, 40-50, 50-60 CM (and >60 CM) - do NOT stop at 30 cm. C stock and sequestration: cumulative 0-10 ... 0-60 CM.",
     "Rule 43 restated.", f"Author {TODAY}", "IN FORCE"),
    ("Parameter-specific",
     "SOC (Walkley-Black / 'SOC' / OM/1.724) of every paper goes to SOC(active C pool); every other carbon pool reported (TOC, TC, POXC, PSOC, MOC, WSC, DOC, LFOC, HFOC, Chan fractions, macro/micro-aggregate C, stocks) goes to its own sheet.",
     "Rule 47 restated.", f"Author {TODAY}", "IN FORCE"),
    ("Parameter-specific",
     "Soil water: AWC (FC - PWP) on BOTH AWC sheets (converted with BD); water content at sampling -> GWC (weight basis; v/v / BD); FC (-33 kPa) -> FC (weight basis); PWP (-1500 kPa) -> PWP (weight basis); Keen box -> WHC; saturated WC -> SAT WC.",
     "Rule 49 restated.", f"Author {TODAY}", "IN FORCE"),
]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
for k, rule in enumerate(new_rules, 1):
    rr = r + k
    for j, v in enumerate((n0 + k,) + rule, 1):
        c = ws.cell(rr, j, v)
        c._style = copy(body_style[j - 1])

# ---- FORMULAS: porosity with PD fallback -----------------------------------------
ws = wb["FORMULAS"]
for row in ws.iter_rows(min_row=2):
    if row[0].value == "TOTAL POROSITY":
        row[1].value = "Porosity (%) = (1 - BD / PD) x 100; PD = treatment PD, else initial PD of the paper, else 2.65 Mg/m3"
        row[2].value = "Excel: =ROUND((1-BD!cell/PD)*100,2) (live link to the BD sheet; PD cell or value named in Notes)"
        row[4].value = "POROSITY sheet (all treatment columns incl. pCA / pZT / pMT / pMTR)"
        row[5].value = f"Author rule {TODAY}"
    if row[0].value == "SD":
        row[4].value = "every data sheet (SD column)"

# ---- README -----------------------------------------------------------------------
ws = wb["README"]
ws["B1"] = "Conservation agriculture meta-analysis - NEW data extraction workbook (structure of META_ANALYSIS_MASTER updated 53, no previous data)"
ws.insert_rows(2, 3)
ws["A2"], ws["B2"] = "NEW SERIES", (f"Started {TODAY}. All data rows of the old master were removed at the author's request; "
                                    "serials start at 1. The history entries further down describe the OLD master (updated 53) and are kept only as the rulebook.")
ws["A3"], ws["B3"] = "PARTIAL CA COLUMNS", ("Every parameter sheet has <prefix>pCA, pZT, pMT, pMTR (olive-green headers) after <prefix>DTR: "
                                             "puddled / conventional rice + ZT wheat with residue (pCA) / without residue (pZT), + MT wheat without (pMT) / with residue (pMTR). See Codes and RULES.")
ws["A4"], ws["B4"] = "LAT_LONG", "One row per study x site: coordinates as reported and in decimal degrees, source (paper / Google Maps), climate class."
for rr in (2, 3, 4):
    ws.cell(rr, 1).font = ARIAL_B
    ws.cell(rr, 2).font = ARIAL
    ws.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top")

# README example row: its SD formula pointed at the old row positions - write it as text
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("=SQRT(2*"):
            c.value = "SD = SQRT(2*Obs/Rep)  (live formula in the data sheets)"
            c.data_type = "s"

# ---- LAT_LONG sheet ---------------------------------------------------------------
ll = wb.create_sheet("LAT_LONG", index=names.index("Treatment_Mapping") + 1)
si_hdr = wb["Study_Info"][1]
ll_cols = ["No.", "SERIAL NO", "Authors", "Year", "Country", "Site/Location",
           "latitude (as reported)", "longitude (as reported)", "latitude", "longitude",
           "Coordinates source", "CLIMATE", "Notes"]
for j, v in enumerate(ll_cols, 1):
    c = ll.cell(1, j, v)
    c._style = copy(si_hdr[0]._style)
for j, w in enumerate([6, 10, 28, 8, 18, 34, 16, 16, 11, 11, 30, 10, 40], 1):
    ll.column_dimensions[get_column_letter(j)].width = w
ll.freeze_panes = "D2"
ll.auto_filter.ref = f"A1:{get_column_letter(len(ll_cols))}1"
ll.row_dimensions[1].height = wb["Study_Info"].row_dimensions[1].height

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)

print(f"saved {OUT}: {len(wb.sheetnames)} sheets, {len(data_sheets)} parameter sheets")
for n, g in summary:
    if not g:
        print("  no treatment group:", n)
