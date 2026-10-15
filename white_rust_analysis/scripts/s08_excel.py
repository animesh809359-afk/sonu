"""Build the analysis workbook: one sheet per step, editable native Excel charts, formula-driven summaries."""
import os, sys
import numpy as np, pandas as pd
sys.path.insert(0, os.path.dirname(__file__))
from common import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, BarChart, LineChart, Reference, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.comments import Comment

OUT = os.path.join(ROOT, 'outputs', 'WhiteRust_Mustard_Hyperspectral_Analysis.xlsx')
T = lambda n: pd.read_csv(os.path.join(TAB, n))
have = lambda n: os.path.exists(os.path.join(TAB, n))
FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F4E2C')
HDR_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=13, color='1F4E2C')
NOTE_FONT = Font(name=FONT, italic=True, size=9, color='555555')
BODY = Font(name=FONT, size=10)
THIN = Side(style='thin', color='D0D0D0')
COL_H, COL_D, COL_3, COL_4 = '2E7D32', 'D08A00', '2A6FB0', 'B03A2E'
DATEC = [COL_3, COL_H, COL_D, COL_4]

wb = Workbook()
wb.remove(wb.active)
INDEX = []


def sheet(name, title, note, ref=''):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = TITLE_FONT
    ws['A2'] = note; ws['A2'].font = NOTE_FONT
    if ref:
        ws['A3'] = 'Method reference(s): ' + ref; ws['A3'].font = NOTE_FONT
    INDEX.append((name, title, note))
    ws.sheet_view.showGridLines = False
    return ws


def write_df(ws, df, r0=5, c0=1, fmt=None, widths=None, index=False):
    if index:
        df = df.reset_index()
    for j, col in enumerate(df.columns):
        c = ws.cell(r0, c0 + j, str(col)); c.font = HDR_FONT; c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, (np.floating, float)) and not np.isfinite(v):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            c = ws.cell(r0 + i, c0 + j, v); c.font = BODY
            c.border = Border(bottom=THIN)
            if isinstance(v, float):
                c.number_format = (fmt or {}).get(df.columns[j], '0.000')
    for j, col in enumerate(df.columns):
        w = (widths or {}).get(col, min(max(len(str(col)) + 2, 10), 28))
        ws.column_dimensions[get_column_letter(c0 + j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, c0)
    return r0, c0, r0 + len(df), c0 + len(df.columns) - 1


def scatter_lines(ws, anchor, title, xcol, ycols, r0, r1, colors, xt='Wavelength (nm)', yt='', names=None, w=24, h=11, dashes=None):
    ch = ScatterChart(); ch.title = title; ch.style = 13; ch.x_axis.title = xt; ch.y_axis.title = yt
    ch.width, ch.height = w, h
    xs = Reference(ws, min_col=xcol, min_row=r0 + 1, max_row=r1)
    for k, yc in enumerate(ycols):
        ys = Reference(ws, min_col=yc, min_row=r0, max_row=r1)
        s = Series(ys, xs, title_from_data=True)
        s.marker.symbol = 'none'; s.smooth = False
        s.graphicalProperties.line.solidFill = colors[k % len(colors)]; s.graphicalProperties.line.width = 22000
        if dashes and dashes[k]:
            s.graphicalProperties.line.dashStyle = 'dash'
        ch.series.append(s)
    ch.x_axis.scaling.min = 400; ch.x_axis.scaling.max = 2400
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.legend.position = 'b'
    ws.add_chart(ch, anchor)


def bar(ws, anchor, title, catcol, valcols, r0, r1, colors, yt='', horizontal=False, w=18, h=10, ymin=None, ymax=None):
    ch = BarChart(); ch.type = 'bar' if horizontal else 'col'; ch.title = title; ch.style = 10
    ch.width, ch.height = w, h; ch.y_axis.title = yt
    data = Reference(ws, min_col=valcols[0], max_col=valcols[-1], min_row=r0, max_row=r1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=catcol, min_row=r0 + 1, max_row=r1))
    for k, s in enumerate(ch.series):
        s.graphicalProperties.solidFill = colors[k % len(colors)]; s.graphicalProperties.line.noFill = True
    if ymin is not None: ch.y_axis.scaling.min = ymin
    if ymax is not None: ch.y_axis.scaling.max = ymax
    ch.x_axis.delete = False; ch.y_axis.delete = False
    if len(valcols) == 1: ch.legend = None
    else: ch.legend.position = 'b'
    ws.add_chart(ch, anchor)


def heat(ws, rng, lo='FFFFE5', mid='78C679', hi='004529'):
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type='min', start_color=lo, mid_type='percentile', mid_value=50,
                                                      mid_color=mid, end_type='max', end_color=hi))


# ===================================================================== 00 README (filled at end)
readme = wb.create_sheet('00_README')

# ===================================================================== 01 data design
M = pd.read_csv(os.path.join(CACHE, 'meta.csv'))
ws = sheet('01_Data_Design', 'Step 1 - Data assembly and experimental design',
           'Four ASD contact-probe workbooks (11 Feb, 24 Feb, 03 Mar, 13 Mar 2026) assembled; Excel AVERAGE columns ignored; raw spectra = 2339.')
tab = M.assign(Date=M['date'].map(DATE_LABEL)).pivot_table(index='Date', columns='label', values='spec_uid', aggfunc='count').reindex([DATE_LABEL[d] for d in DATE_ORDER])[['Healthy', 'Diseased']].reset_index()
r0, c0, r1, c1 = write_df(ws, tab, 5, 1)
n = r1
ws.cell(n + 1, 1, 'Total').font = Font(name=FONT, bold=True)
for j, col in enumerate(['B', 'C']):
    ws.cell(n + 1, 2 + j, f'=SUM({col}6:{col}{n})').font = Font(name=FONT, bold=True)
ws.cell(5, 4, 'Total').font = HDR_FONT; ws.cell(5, 4).fill = HDR_FILL
ws.cell(5, 5, 'Healthy : Diseased').font = HDR_FONT; ws.cell(5, 5).fill = HDR_FILL
for r in range(6, n + 2):
    ws.cell(r, 4, f'=B{r}+C{r}'); ws.cell(r, 5, f'="1 : "&TEXT(C{r}/B{r},"0.0")')
bar(ws, 'H5', 'Spectra per date and class', 1, [2, 3], 5, n, [COL_H, COL_D], 'Number of spectra', w=16, h=8)
for k, fac in enumerate(['variety', 'stage']):
    t2 = M.pivot_table(index=fac, columns='label', values='spec_uid', aggfunc='count')[['Healthy', 'Diseased']].reset_index()
    rr = n + 4 + k * 7
    ws.cell(rr - 1, 1, f'Spectra by {fac}').font = Font(name=FONT, bold=True)
    write_df(ws, t2, rr, 1)
pl = M.groupby(['plot', 'date']).agg(variety=('variety', 'first'), label=('label', 'first'), stage=('stage', 'first'), n=('spec_uid', 'count')).reset_index()
plw = pl.assign(cell=pl['stage'].str[:3] + ' (' + pl['n'].astype(str) + ')').pivot(index=['plot', 'variety', 'label'], columns='date', values='cell')
plw.columns = [DATE_LABEL[c] for c in plw.columns]
ws.cell(n + 18, 1, 'Plot x date layout: growth stage (number of spectra)').font = Font(name=FONT, bold=True)
write_df(ws, plw.reset_index(), n + 19, 1)
ws.freeze_panes = None

# ===================================================================== 02 QC
ws = sheet('02_QC', 'Step 2 - Quality control: noise, splice correction, outlier screening',
           'Noise = SD of high-frequency residual (raw spectrum minus 9-nm moving average). Outliers removed when flagged by >= 3 of 4 label-free rules.',
           'Hueni & Bialek (2017); Liu et al. (2008); Selvaprakash et al. (2026)')
NZ = T('qc_noise.csv')
r0, c0, r1, c1 = write_df(ws, NZ, 5, 1, fmt={'Mean_reflectance': '0.000', 'HF_noise_SD': '0.00000', 'SNR': '0.0'})
bar(ws, 'G5', 'Signal-to-noise ratio by spectral region (raw data)', 1, [4], r0, r1, [COL_H], 'SNR', w=17, h=8)
QS = T('qc_summary.csv'); QS.columns = ['Quantity', 'Value']
write_df(ws, QS, r1 + 3, 1, fmt={'Value': '0.00000'})
OL = T('qc_outliers.csv')
summ = pd.DataFrame({'Rule': ['Hotelling T2 (99%)', 'Q residual (99%)', 'Isolation Forest (2%)', 'SAM to plot-date median (robust z > 3.5)', 'Removed (>= 3 votes)'],
                     'Flagged spectra': [OL['T2_flag'].sum(), OL['Q_flag'].sum(), OL['IsoForest_flag'].sum(), OL['SAM_flag'].sum(), OL['removed'].sum()]})
rr = r1 + 5 + len(QS)
write_df(ws, summ, rr, 1)
ws.cell(rr + 7, 1, 'Removed spectra (all were from diseased plots):').font = Font(name=FONT, bold=True)
write_df(ws, OL[OL['removed']].drop(columns=['removed']), rr + 8, 1)
ws.freeze_panes = None

# ===================================================================== 03 spectra by class
ws = sheet('03_Spectra_Class', 'Step 3a - Mean spectra of healthy and white-rust-diseased leaves',
           'Splice-corrected, Savitzky-Golay smoothed (11 nm, 2nd order), 10-nm resampled reflectance; first/second derivatives from SG filter.',
           'Savitzky & Golay (1964); Zhang et al. (2012); Mandal et al. (2022)')
S = T('spectra_mean_class.csv')
r0, c0, r1, c1 = write_df(ws, S, 5, 1, fmt={c: '0.0000' for c in S.columns} | {'FD_Healthy': '0.000000', 'FD_Diseased': '0.000000', 'SD2_Healthy': '0.0000000', 'SD2_Diseased': '0.0000000'})
cols = list(S.columns)
scatter_lines(ws, 'O5', 'Mean reflectance: healthy vs diseased', 1, [cols.index('Healthy_mean') + 1, cols.index('Diseased_mean') + 1], r0, r1, [COL_H, COL_D], yt='Reflectance')
scatter_lines(ws, 'O28', 'Difference spectrum (Diseased - Healthy)', 1, [cols.index('Difference_D-H') + 1], r0, r1, ['333333'], yt='Reflectance difference')
scatter_lines(ws, 'O51', 'Ratio spectrum (Diseased / Healthy)', 1, [cols.index('Ratio_D/H') + 1], r0, r1, ['333333'], yt='Ratio')
scatter_lines(ws, 'O74', 'First-derivative spectra', 1, [cols.index('FD_Healthy') + 1, cols.index('FD_Diseased') + 1], r0, r1, [COL_H, COL_D], yt='dR/dλ')

# ===================================================================== 04 spectra by date / variety / stage
for fac, nm, title in [('date', '04_Spectra_Date', 'Step 3b - Mean spectra and difference spectra by measurement date'),
                       ('variety', '05_Spectra_Variety', 'Step 3c - Mean spectra by variety'),
                       ('stage', '06_Spectra_Stage', 'Step 3d - Mean spectra by growth stage')]:
    ws = sheet(nm, title, 'H = healthy plots, D = diseased plots, Diff = D - H (spectrum-level means).')
    S = T(f'spectra_mean_by_{fac}.csv')
    r0, c0, r1, c1 = write_df(ws, S, 5, 1, fmt={c: '0.0000' for c in S.columns})
    cols = list(S.columns)
    diffs = [cols.index(c) + 1 for c in cols if c.endswith('_Diff')]
    hs = [cols.index(c) + 1 for c in cols if c.endswith('_H')]
    ds = [cols.index(c) + 1 for c in cols if c.endswith('_D')]
    col_anchor = get_column_letter(c1 + 2)
    scatter_lines(ws, f'{col_anchor}5', f'Difference spectra (D - H) by {fac}', 1, diffs, r0, r1, DATEC, yt='Reflectance difference')
    scatter_lines(ws, f'{col_anchor}28', f'Healthy (solid) vs diseased (dashed) by {fac}', 1, [x for p in zip(hs, ds) for x in p], r0, r1,
                  [c for c in DATEC[:len(hs)] for _ in (0, 1)], yt='Reflectance', dashes=[k % 2 == 1 for k in range(2 * len(hs))])

# ===================================================================== 07 band-wise statistics
ws = sheet('07_Bandwise_Stats', 'Step 3e - Band-wise statistics: Welch t-test, Mann-Whitney, FDR, Cohen d, correlation, J-M distance',
           'Tests on plot-date mean spectra (n = 192: 48 healthy, 144 diseased) to avoid pseudo-replication; J-M on individual spectra.',
           'Benjamini & Hochberg (1995); Mandal et al. (2022, 2023); Richards & Jia (2006)')
B = T('bandwise_statistics.csv')
r0, c0, r1, c1 = write_df(ws, B, 5, 1, fmt={'p_value': '0.00E+00', 'q_FDR': '0.00E+00', 'MannWhitney_p': '0.00E+00', 'FD_q_FDR': '0.00E+00'})
cols = list(B.columns)
an = get_column_letter(c1 + 2)
scatter_lines(ws, f'{an}5', "Cohen's d per band (Diseased - Healthy)", 1, [cols.index('Cohen_d') + 1], r0, r1, ['333333'], yt="Cohen's d")
scatter_lines(ws, f'{an}28', 'Point-biserial correlation spectrum', 1, [cols.index('Point_biserial_r_spectrum') + 1, cols.index('Point_biserial_r_plotdate') + 1], r0, r1, [COL_3, COL_4], yt='r')
scatter_lines(ws, f'{an}51', 'Jeffries-Matusita distance per band', 1, [cols.index('JM_distance') + 1, cols.index('FD_JM') + 1], r0, r1, [COL_H, COL_D], yt='J-M (0-2)')
ws2 = sheet('08_Sensitive_Regions', 'Step 3f - Contiguous sensitive regions (FDR q < 0.05) and effect size by date',
            'Regions of >= 1 consecutive 10-nm band significant after FDR correction; right: Cohen d per band within each date.')
RG = T('sensitive_regions.csv'); r0, c0, r1, c1 = write_df(ws2, RG, 5, 1)
CD = T('bandwise_cohen_d_by_date.csv')
r0b, c0b, r1b, c1b = write_df(ws2, CD, 5, 8)
scatter_lines(ws2, 'N5', "Cohen's d per band within each date", 8, list(range(9, 13)), r0b, r1b, DATEC, yt="Cohen's d")
ws2.freeze_panes = None

# ===================================================================== 09 red edge
ws = sheet('09_RedEdge_Params', 'Step 3g - Red-edge and edge parameters by class and date (plot-date means, Welch t-test)',
           'REP_LI linear interpolation; REP_FD/REV from first-derivative maximum 680-750 nm; Db/Dy/Dr edge amplitudes, positions, sums.',
           'Guyot & Baret (1988); Mandal et al. (2022, 2023); Zhang et al. (2012)')
ER = T('red_edge_parameters.csv')
write_df(ws, ER, 5, 1, fmt={'Welch_p': '0.0000', 'Healthy_mean': '0.0000', 'Diseased_mean': '0.0000', 'Healthy_SD': '0.0000', 'Diseased_SD': '0.0000'})
sub = ER[ER['Parameter'] == 'REP_LI'][['Date', 'Healthy_mean', 'Diseased_mean']]
r0, c0, r1, c1 = write_df(ws, sub, 5, 10)
bar(ws, 'N5', 'Red-edge position (REP, nm) by date', 10, [11, 12], r0, r1, [COL_H, COL_D], 'REP (nm)', ymin=705, ymax=722, w=15, h=8)

# ===================================================================== 10 PCA / t-SNE
ws = sheet('10_PCA_tSNE', 'Step 3h - PCA and t-SNE of SNV spectra; eta-squared of PC scores by factor',
           'Scores for every spectrum; eta2 = share of PC-score variance explained by each factor.',
           'van der Maaten & Hinton (2008); Selvaprakash et al. (2026)')
V = T('pca_variance.csv'); r0, c0, r1, c1 = write_df(ws, V, 5, 1, fmt={'Explained_variance_pct': '0.00', 'Cumulative_pct': '0.00'})
bar(ws, 'E5', 'PCA explained variance (%)', 1, [2], r0, r1, [COL_H], '%', w=14, h=7)
E = T('pca_eta2.csv'); write_df(ws, E, 19, 1)
P = T('pca_tsne_scores.csv')
P = P.sort_values('label')
r0, c0, r1, c1 = write_df(ws, P, 30, 1)
nh = (P['label'] == 'Diseased').sum()
cols = list(P.columns)
for ttl, xc, yc, anc in [('PCA score plot (PC1 vs PC2)', 'PC1', 'PC2', 'M5'), ('t-SNE embedding', 'tSNE1', 'tSNE2', 'M28')]:
    ch = ScatterChart(); ch.title = ttl; ch.style = 13; ch.width, ch.height = 16, 11
    ch.x_axis.title = xc; ch.y_axis.title = yc
    for lab, a, b, colr in [('Diseased', r0 + 1, r0 + nh, COL_D), ('Healthy', r0 + nh + 1, r1, COL_H)]:
        s = Series(Reference(ws, min_col=cols.index(yc) + 1, min_row=a, max_row=b), Reference(ws, min_col=cols.index(xc) + 1, min_row=a, max_row=b), title=lab)
        s.marker.symbol = 'circle'; s.marker.size = 3; s.graphicalProperties.line.noFill = True
        s.marker.graphicalProperties.solidFill = colr; s.marker.graphicalProperties.line.solidFill = colr
        ch.series.append(s)
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.legend.position = 'b'
    ws.add_chart(ch, anc)

# ===================================================================== 11 PERMANOVA + mixed models
ws = sheet('11_PERMANOVA', 'Step 3i - PERMANOVA variance partitioning of SNV spectra (plot-date means, 999 permutations)',
           'Marginal (Type II) SS for main effects; interactions tested in full model. Disease/variety permuted across plots, date/stage within plots.',
           'Anderson (2001)')
PM = T('permanova.csv'); r0, c0, r1, c1 = write_df(ws, PM, 5, 1, fmt={'p_perm': '0.000', 'R2_partial_pct': '0.00', 'Pseudo_F': '0.00', 'SS': '0.000'})
bar(ws, 'I5', 'Marginal variance explained (%)', 1, [5], r0, r1, [COL_H], '% of total SS', horizontal=True, w=15, h=8)
ws = sheet('12_Mixed_Models', 'Step 3j - Linear mixed models: feature ~ disease + date + variety + (1 | plot)',
           'Features standardised (SD units); p_disease_x_date = likelihood-ratio test of disease x date interaction.', 'Bates et al. (2015)')
LM = T('mixed_models_features.csv')
write_df(ws, LM, 5, 1, fmt={'p_disease': '0.00E+00', 'p_disease_x_date': '0.00E+00', 'q_disease': '0.00E+00'})

# ===================================================================== 13 VI
ws = sheet('13_VI_Ranking', 'Step 4a - Published vegetation indices: definitions, ranking (AUC), J-M, date-consistency',
           'Cohen d per date on plot-date means; Sign_consistent = same direction on all four dates. Welch p on plot-date means, FDR corrected.',
           'see Reference column; Mahlein et al. (2013); Mandal et al. (2023)')
VR = T('vi_ranking.csv')
r0, c0, r1, c1 = write_df(ws, VR, 5, 1, fmt={'Welch_p_plotdate': '0.00E+00', 'q_FDR': '0.00E+00'}, widths={'Reference': 28, 'Trait': 26})
top = VR.head(15)[['Index', 'AUC']].iloc[::-1]
r0b, c0b, r1b, c1b = write_df(ws, top, 5, c1 + 2)
bar(ws, f'{get_column_letter(c1 + 5)}5', 'Top-15 indices by AUC', c0b, [c0b + 1], r0b, r1b, [COL_H], 'AUC', horizontal=True, ymin=0.5, ymax=0.75, w=15, h=11)

# ===================================================================== 14 NDSI
ws = sheet('14_NDSI_MWRI', 'Step 4b - Exhaustive normalised-difference index (NDSI) search -> Mustard White Rust Index (MWRI)',
           'Score = mean |Cohen d| across the four dates (0 if the sign changes between dates). Heat map shows every 20 nm (upper triangle).',
           'Mandal et al. (2023); Selvaprakash et al. (2026); Mahlein et al. (2013)')
NT = T('ndsi_top_pairs.csv'); r0, c0, r1, c1 = write_df(ws, NT, 5, 1, widths={'Formula': 30})
NM = pd.read_csv(os.path.join(TAB, 'ndsi_matrix_20nm.csv'), index_col=0)
NM = NM.where(np.triu(np.ones(NM.shape), 1).astype(bool))
NM.index.name = 'Band_i \\ Band_j'
r0m, c0m, r1m, c1m = write_df(ws, NM.reset_index(), r1 + 4, 1, fmt={c: '0.00' for c in NM.columns})
for j in range(2, c1m + 1):
    ws.column_dimensions[get_column_letter(j)].width = 5
ws.column_dimensions['A'].width = 12
heat(ws, f'B{r0m + 1}:{get_column_letter(c1m)}{r1m}')
ws.freeze_panes = None

# ===================================================================== 15 CWT
ws = sheet('15_CWT_Scalogram', 'Step 4c - Continuous wavelet transform sensitivity: |Cohen d| per wavelength x scale',
           'Mexican-hat mother wavelet, dyadic scales 2^1..2^8; computed on plot-date means.', 'Torrence & Compo (1998); Cheng et al. (2011); Tian et al. (2021); Lv et al. (2023)')
CW = pd.read_csv(os.path.join(TAB, 'cwt_abs_d_scalogram.csv'), index_col=0); CW.index.name = 'Wavelength_nm'
r0, c0, r1, c1 = write_df(ws, CW.reset_index(), 5, 1)
heat(ws, f'B6:{get_column_letter(c1)}{r1}')
best = CW.stack().sort_values(ascending=False).head(15).reset_index(); best.columns = ['Wavelength_nm', 'Scale', 'abs_d']
write_df(ws, best, 5, c1 + 2)
ws.freeze_panes = 'A6'

# ===================================================================== 16 features / VI definitions
ws = sheet('16_Feature_Families', 'Step 4d - Feature families used for modelling',
           'All features computed from splice-corrected, SG-smoothed 1-nm reflectance (400-2400 nm).')
fam = pd.DataFrame([
    ['F1', 'Reflectance bands (SNV, 10 nm)', 201, 'Mandal et al., 2023; Terentev et al., 2023'],
    ['F2', 'Published vegetation indices + red-edge parameters', 46, 'Mandal et al., 2023; Tian et al., 2021; Zhang et al., 2012; Lv et al., 2023'],
    ['F3', 'Novel NDSI (exhaustive band-pair search, in-fold)', 3, 'Mandal et al., 2023; Selvaprakash et al., 2026'],
    ['F4', 'First-derivative bands + derivative indices + edge parameters', 221, 'Lv et al., 2023; Zhang et al., 2012; Tsai & Philpot, 1998'],
    ['F5', 'Continuum-removal depth/width/area/position (6 absorption features)', 24, 'Zhang et al., 2012; Kokaly & Clark, 1999'],
    ['F6', 'CWT wavelet features (Mexican hat, scales 1-8, every 10 nm)', 1608, 'Tian et al., 2021; Lv et al., 2023; Cheng et al., 2011'],
    ['-', 'Image texture/colour features (GLCM, colour moments)', 0, 'Not applicable to point spectra (Huang et al., 2020; Zhu et al., 2017; Nguyen et al., 2023)'],
], columns=['Family', 'Description', 'n_features', 'Source'])
write_df(ws, fam, 5, 1, widths={'Description': 60, 'Source': 70})
VD = T('vi_definitions.csv'); ws.cell(15, 1, 'Vegetation index definitions (F2)').font = Font(name=FONT, bold=True)
write_df(ws, VD, 16, 1, widths={'Reference': 30, 'Trait': 34})
ws.freeze_panes = None

# ===================================================================== 17 preprocessing E1
if have('E1_preprocessing_reps.csv'):
    ws = sheet('17_Preprocessing_E1', 'Step 5 - E1: preprocessing comparison (full spectrum, 10 x 4-fold plot-grouped nested CV)',
               'Summary cells are AVERAGEIFS formulas over the repeat-level table (right). MSC and date normalisation fitted on training plots only.',
               'Barnes et al. (1989); Geladi et al. (1985); Lv et al. (2023)')
    E1 = T('E1_preprocessing_reps.csv')
    rr0, cc0, rr1, cc1 = write_df(ws, E1[['variant', 'clf', 'rep', 'BA', 'MCC', 'Kappa', 'MacroF1', 'ROC_AUC']], 5, 10)
    vs = ['Reflectance (SG)', 'SNV', 'MSC', 'First derivative (SG)', 'Date normalisation', 'Date normalisation + SNV']
    clfs = ['LDA', 'PLS-DA', 'LR', 'SVM']
    for j, h in enumerate(['Preprocessing'] + [f'BA {c}' for c in clfs] + ['Mean BA']):
        c = ws.cell(5, 1 + j, h); c.font = HDR_FONT; c.fill = HDR_FILL
    for i, v in enumerate(vs, start=6):
        ws.cell(i, 1, v).font = BODY
        for j, c in enumerate(clfs, start=2):
            ws.cell(i, j, f'=AVERAGEIFS($M${rr0 + 1}:$M${rr1},$J${rr0 + 1}:$J${rr1},$A{i},$K${rr0 + 1}:$K${rr1},"{c}")').number_format = '0.000'
        ws.cell(i, 6, f'=AVERAGE(B{i}:E{i})').number_format = '0.000'
    ws.column_dimensions['A'].width = 26
    heat(ws, 'B6:F11')
    bar(ws, 'A14', 'Balanced accuracy by preprocessing and classifier', 1, [2, 3, 4, 5], 5, 11, [COL_3, COL_H, COL_D, COL_4], 'Balanced accuracy', ymin=0.5, ymax=1.0, w=20, h=9)

# ===================================================================== 18-19 main model results
if have('cv_main_reps.csv'):
    REP = T('cv_main_reps.csv')
    ws_r = sheet('19_Model_Reps', 'Step 7 - Repeat-level performance of every pipeline (25 repeats x 10 scenarios x 10 classifiers)',
                 'Each row = predictions pooled over the 4 outer folds of one repeat (every spectrum predicted once). Source for the formula-driven summaries.')
    keep = ['scenario', 'clf', 'rep', 'OA', 'BA', 'Sensitivity', 'Specificity', 'Precision', 'F1_Diseased', 'F1_Healthy', 'MacroF1', 'Kappa', 'MCC', 'ROC_AUC', 'PR_AUC_Healthy', 'TN', 'FP', 'FN', 'TP']
    rr0, cc0, rr1, cc1 = write_df(ws_r, REP[keep], 5, 1)
    colL = {k: get_column_letter(1 + i) for i, k in enumerate(keep)}
    ws = sheet('18_Model_Matrix', 'Step 7 - Model comparison: balanced accuracy of 10 feature scenarios x 10 classifiers',
               "Cells are AVERAGEIFS formulas over sheet '19_Model_Reps' (mean of 25 repeats of nested, plot-grouped, stratified 4-fold CV).",
               'Tian et al. (2021); Lv et al. (2023); Varma & Simon (2006); Roberts et al. (2017)')
    from s04_cv import SCENARIOS
    from models import CLF_NAMES
    ws.cell(5, 1, 'Scenario').font = HDR_FONT; ws.cell(5, 1).fill = HDR_FILL
    ws.cell(5, 2, 'Description').font = HDR_FONT; ws.cell(5, 2).fill = HDR_FILL
    for j, c in enumerate(CLF_NAMES, start=3):
        cc = ws.cell(5, j, c); cc.font = HDR_FONT; cc.fill = HDR_FILL
    cc = ws.cell(5, 13, 'Row max'); cc.font = HDR_FONT; cc.fill = HDR_FILL
    rng = lambda k: f"'19_Model_Reps'!${colL[k]}${rr0 + 1}:${colL[k]}${rr1}"
    for i, (sc, nm) in enumerate(SCENARIOS.items(), start=6):
        ws.cell(i, 1, sc); ws.cell(i, 2, nm)
        for j, c in enumerate(CLF_NAMES, start=3):
            ws.cell(i, j, f'=AVERAGEIFS({rng("BA")},{rng("scenario")},$A{i},{rng("clf")},{get_column_letter(j)}$5)').number_format = '0.000'
        ws.cell(i, 13, f'=MAX(C{i}:L{i})').number_format = '0.000'
    ws.column_dimensions['B'].width = 44
    for j in range(3, 14): ws.column_dimensions[get_column_letter(j)].width = 9
    heat(ws, 'C6:L15')
    ws.cell(17, 1, 'Matthews correlation coefficient (MCC)').font = Font(name=FONT, bold=True)
    for j, c in enumerate(CLF_NAMES, start=3):
        cc = ws.cell(18, j, c); cc.font = HDR_FONT; cc.fill = HDR_FILL
    for i, sc in enumerate(SCENARIOS, start=19):
        ws.cell(i, 1, sc); ws.cell(i, 2, SCENARIOS[sc])
        for j, c in enumerate(CLF_NAMES, start=3):
            ws.cell(i, j, f'=AVERAGEIFS({rng("MCC")},{rng("scenario")},$A{i},{rng("clf")},{get_column_letter(j)}$18)').number_format = '0.000'
    heat(ws, 'C19:L28')
    bar(ws, 'A31', 'Best balanced accuracy per feature scenario', 1, [13], 5, 15, [COL_H], 'Balanced accuracy', ymin=0.5, ymax=1.0, w=20, h=9)
    ws.freeze_panes = None

    # summary with SD and CI via formulas
    ws = sheet('20_Model_Summary', 'Step 7 - Full metric summary (mean, SD, 95% CI over 25 repeats) for every pipeline',
               "Mean = AVERAGEIFS; SD = SUMPRODUCT formula; CI = mean +/- t(0.975, n-1) x SD / sqrt(n). All read from '19_Model_Reps'.",
               'Chicco & Jurman (2020); Cohen (1960)')
    hdr = ['Scenario', 'Classifier', 'n', 'BA mean', 'BA SD', 'BA CI95 low', 'BA CI95 high', 'OA', 'Sensitivity', 'Specificity', 'MacroF1', 'Kappa', 'MCC', 'ROC_AUC', 'PR_AUC_Healthy']
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(5, j, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = Alignment(wrap_text=True, horizontal='center')
    i = 6
    for sc in SCENARIOS:
        for c in CLF_NAMES:
            ws.cell(i, 1, sc); ws.cell(i, 2, c)
            cond = f'{rng("scenario")},$A{i},{rng("clf")},$B{i}'
            ws.cell(i, 3, f'=COUNTIFS({cond})')
            ws.cell(i, 4, f'=AVERAGEIFS({rng("BA")},{cond})').number_format = '0.000'
            ws.cell(i, 5, f'=SQRT(SUMPRODUCT(({rng("scenario")}=$A{i})*({rng("clf")}=$B{i})*({rng("BA")}-D{i})^2)/(C{i}-1))').number_format = '0.000'
            ws.cell(i, 6, f'=D{i}-TINV(0.05,C{i}-1)*E{i}/SQRT(C{i})').number_format = '0.000'
            ws.cell(i, 7, f'=D{i}+TINV(0.05,C{i}-1)*E{i}/SQRT(C{i})').number_format = '0.000'
            for j, k in enumerate(['OA', 'Sensitivity', 'Specificity', 'MacroF1', 'Kappa', 'MCC', 'ROC_AUC', 'PR_AUC_Healthy'], start=8):
                ws.cell(i, j, f'=AVERAGEIFS({rng(k)},{cond})').number_format = '0.000'
            i += 1
    for j in range(1, 16): ws.column_dimensions[get_column_letter(j)].width = 11
    heat(ws, f'D6:D{i - 1}')
    ws.freeze_panes = 'C6'

    # best pipelines, Friedman, t-tests
    ws = sheet('21_Best_Pipelines', 'Step 8 - Best pipelines, Friedman/Nemenyi ranking and corrected resampled t-tests',
               'Friedman blocks: classifiers ranked within each scenario x repeat; scenarios ranked within each classifier x repeat.',
               'Demsar (2006); Nadeau & Bengio (2003)')
    BP = T('cv_best_per_scenario.csv')[['scenario', 'Scenario_name', 'clf', 'BA_mean', 'BA_SD', 'BA_CI95_low', 'BA_CI95_high', 'MCC_mean', 'Kappa_mean', 'MacroF1_mean', 'ROC_AUC_mean', 'Sensitivity_mean', 'Specificity_mean']]
    r0, c0, r1, c1 = write_df(ws, BP, 5, 1, widths={'Scenario_name': 40})
    bar(ws, f'{get_column_letter(c1 + 2)}5', 'Best balanced accuracy per scenario', 1, [4], r0, r1, [COL_H], 'Balanced accuracy', ymin=0.5, ymax=1.0, w=17, h=8)
    FT = T('friedman_tests.csv'); r0b, _, r1b, _ = write_df(ws, FT, r1 + 3, 1)
    FC = T('friedman_classifiers.csv'); r0c, _, r1c, _ = write_df(ws, FC, r1b + 3, 1)
    bar(ws, f'{get_column_letter(c1 + 2)}24', 'Average Friedman rank of classifiers (lower = better)', 1, [2], r0c, r1c, [COL_3], 'Average rank', horizontal=True, w=17, h=8)
    FS = T('friedman_scenarios.csv'); r0d, _, r1d, _ = write_df(ws, FS, r1c + 3, 1)
    CT = T('corrected_ttests.csv'); write_df(ws, CT, r1d + 3, 1, fmt={'p_value': '0.0000'}, widths={'Comparison': 40})
    ws.freeze_panes = None

    # confusion matrices with formula metrics
    ws = sheet('22_Confusion_Matrix', 'Step 8 - Confusion matrices of the best pipelines (representative repeat) with formula-derived metrics',
               'Edit the four counts and every metric recalculates (Diseased = positive class).', 'Cohen (1960); Chicco & Jurman (2020); Zhang et al. (2012)')
    CM = T('confusion_matrices.csv')
    r = 5
    for _, rw in CM.iterrows():
        ws.cell(r, 1, f"{rw['Role']} - {rw['Pipeline']} - {rw['Level']} (repeat {rw['Repeat']})").font = Font(name=FONT, bold=True)
        for j, h in enumerate(['', 'Predicted Healthy', 'Predicted Diseased', 'Total'], start=1):
            c = ws.cell(r + 1, j, h); c.font = HDR_FONT; c.fill = HDR_FILL
        ws.cell(r + 2, 1, 'True Healthy'); ws.cell(r + 2, 2, int(rw['TN'])); ws.cell(r + 2, 3, int(rw['FP'])); ws.cell(r + 2, 4, f'=B{r+2}+C{r+2}')
        ws.cell(r + 3, 1, 'True Diseased'); ws.cell(r + 3, 2, int(rw['FN'])); ws.cell(r + 3, 3, int(rw['TP'])); ws.cell(r + 3, 4, f'=B{r+3}+C{r+3}')
        for cc in (ws.cell(r + 2, 2), ws.cell(r + 2, 3), ws.cell(r + 3, 2), ws.cell(r + 3, 3)):
            cc.font = Font(name=FONT, color='0000FF')
        TN, FP, FN, TP = f'B{r+2}', f'C{r+2}', f'B{r+3}', f'C{r+3}'
        mets = [('Overall accuracy', f'=({TP}+{TN})/({TP}+{TN}+{FP}+{FN})'),
                ('Sensitivity (recall, diseased)', f'={TP}/({TP}+{FN})'),
                ('Specificity (healthy)', f'={TN}/({TN}+{FP})'),
                ('Balanced accuracy', f'=(G{r+2}+G{r+3})/2'),
                ("Producer's accuracy healthy", f'={TN}/({TN}+{FP})'),
                ("User's accuracy healthy", f'={TN}/({TN}+{FN})'),
                ('Precision (diseased)', f'={TP}/({TP}+{FP})'),
                ('F1 diseased', f'=2*{TP}/(2*{TP}+{FP}+{FN})'),
                ('F1 healthy', f'=2*{TN}/(2*{TN}+{FN}+{FP})'),
                ('Macro F1', f'=(G{r+8}+G{r+9})/2'),
                ("Cohen's kappa", f'=(G{r+1}-(({TP}+{FP})*({TP}+{FN})+({TN}+{FN})*({TN}+{FP}))/({TP}+{TN}+{FP}+{FN})^2)/(1-(({TP}+{FP})*({TP}+{FN})+({TN}+{FN})*({TN}+{FP}))/({TP}+{TN}+{FP}+{FN})^2)'),
                ('MCC', f'=({TP}*{TN}-{FP}*{FN})/SQRT(({TP}+{FP})*({TP}+{FN})*({TN}+{FP})*({TN}+{FN}))')]
        for k, (nm_, f_) in enumerate(mets):
            ws.cell(r + 1 + k, 6, nm_).font = BODY; c = ws.cell(r + 1 + k, 7, f_); c.number_format = '0.000'
        r += 15
    ws.column_dimensions['A'].width = 16; ws.column_dimensions['F'].width = 30
    for col in 'BCD': ws.column_dimensions[col].width = 18

    ws = sheet('23_ROC', 'Step 8 - ROC curves of the best pipelines (representative repeat)', 'TPR interpolated on a 0.02 FPR grid.')
    ROC = T('roc_curves.csv'); r0, c0, r1, c1 = write_df(ws, ROC, 5, 1)
    ch = ScatterChart(); ch.title = 'ROC curves'; ch.style = 13; ch.width, ch.height = 14, 12
    ch.x_axis.title = 'False-positive rate'; ch.y_axis.title = 'True-positive rate'
    for k in range(2, c1 + 1):
        s = Series(Reference(ws, min_col=k, min_row=r0, max_row=r1), Reference(ws, min_col=1, min_row=r0 + 1, max_row=r1), title_from_data=True)
        s.marker.symbol = 'none'; s.graphicalProperties.line.solidFill = [COL_H, COL_D][k - 2]; s.graphicalProperties.line.width = 25000
        ch.series.append(s)
    ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = 1; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.legend.position = 'b'
    ws.add_chart(ch, 'F5')

    ws = sheet('24_Plot_Level', 'Step 8 - Plot-date level performance (majority vote of spectra in each plot x date, n = 192)',
               'Mean over 25 repeats.')
    PDL = T('cv_main_summary_plotdate.csv')[['scenario', 'clf', 'BA_mean', 'BA_SD', 'OA_mean', 'Sensitivity_mean', 'Specificity_mean', 'MCC_mean', 'Kappa_mean']]
    write_df(ws, PDL.sort_values('BA_mean', ascending=False), 5, 1)

    ws = sheet('25_Selection_Freq', 'Step 6 - Feature-selection stability across the 100 outer training folds',
               'Frequency = % of outer folds in which a feature was retained by the robust selection (ReliefF per date -> stability >= 3/4 dates -> |r| < 0.98 -> VIF < 10 -> SFFS-LDA), SPA or GA-PLS.',
               'Robnik-Sikonja & Kononenko (2003); Pudil et al. (1994); Araujo et al. (2001); Leardi & Lupianez Gonzalez (1998); Tian et al. (2021); Lv et al. (2023)')
    FR = T('selection_frequency.csv')
    r0, c0, r1, c1 = write_df(ws, FR, 5, 1, widths={'Scenario_name': 38, 'Feature': 22})
    s8 = FR[FR['scenario'] == 'S8'].head(15)[['Feature', 'Frequency_pct']].iloc[::-1]
    r0b, c0b, r1b, c1b = write_df(ws, s8, 5, c1 + 2)
    bar(ws, f'{get_column_letter(c1 + 5)}5', 'S8 combined: most frequently selected features (%)', c0b, [c0b + 1], r0b, r1b, [COL_H], '% of folds', horizontal=True, w=16, h=11)
    BF = T('band_selection_frequency.csv'); r0c, c0c, r1c, c1c = write_df(ws, BF, r1b + 4, c1 + 2)
    scatter_lines(ws, f'{get_column_letter(c1 + 5)}30', 'Band selection frequency: SPA (S9) and GA-PLS (S10)', c0c, [c0c + 1, c0c + 2], r0c, r1c, [COL_3, COL_D], yt='% of folds', w=18, h=9)
    ws.freeze_panes = None

    ws = sheet('26_Final_Features', 'Step 9 - Final robust feature set (selected on all data) and its interpretation',
               'Permutation importance from 4-fold plot-grouped CV with the fixed feature set (SVM); mean |SHAP| from XGBoost TreeSHAP; JM = multivariate Jeffries-Matusita.',
               'Breiman (2001); Lundberg & Lee (2017); Richards & Jia (2006)')
    FI = T('final_features.csv'); r0, c0, r1, c1 = write_df(ws, FI, 5, 1, widths={'Feature': 22})
    bar(ws, f'{get_column_letter(c1 + 2)}5', 'Permutation importance (drop in BA)', 1, [6], r0, r1, [COL_H], 'Drop in BA', horizontal=True, w=16, h=9)
    JM = T('final_features_JM.csv'); write_df(ws, JM, r1 + 3, 1)
    LC = T('lr_coefficients.csv'); r0c, c0c, r1c, c1c = write_df(ws, LC, r1 + 12, 1)
    scatter_lines(ws, f'{get_column_letter(c1 + 2)}26', 'Full-spectrum logistic regression: standardised coefficients', 1, [2], r0c, r1c, ['333333'], yt='Coefficient', w=20, h=9)
    ws.freeze_panes = None

    ws = sheet('27_Subgroups', 'Step 9 - Performance of the best pipelines within each date, variety and growth stage',
               'From the main CV predictions (models trained on the pooled data), averaged over 25 repeats.')
    SG = T('subgroup_performance.csv'); r0, c0, r1, c1 = write_df(ws, SG, 5, 1)
    piv = SG.pivot_table(index=['Factor', 'Level'], columns='Role', values='BA', sort=False).reset_index()
    piv['Group'] = piv['Factor'] + ': ' + piv['Level']
    piv = piv[['Group'] + [c for c in piv.columns if c not in ('Factor', 'Level', 'Group')]]
    r0b, c0b, r1b, c1b = write_df(ws, piv, 5, c1 + 2)
    bar(ws, f'{get_column_letter(c1b + 2)}5', 'Balanced accuracy by sub-group', c0b, list(range(c0b + 1, c1b + 1)), r0b, r1b, [COL_H, COL_D], 'Balanced accuracy', ymin=0.5, ymax=1.0, w=20, h=9)
    ws.freeze_panes = None

if have('E7_ofc_reps.csv'):
    ws = sheet('27b_OFC_Tian2021', 'Step 6b - E7: optimal feature combinations (OFCs) from wrapper-based SFFS (Tian et al., 2021)',
               "Disease-specific feature (DSF) pool: top-5% ReliefF on >= 3 of 4 dates, |r| < 0.98, VIF < 10 (all families F1-F6). SFFS wrapped around LDA, kNN or SVM gives OFC-LDA/kNN/SVM; "
               "each set is classified by LDA, kNN and SVM (25 x 4-fold nested plot-grouped CV). Matrix cells are AVERAGEIFS formulas over the repeat table (right).",
               'Tian et al. (2021); Pudil et al. (1994); Robnik-Sikonja & Kononenko (2003)')
    E7 = T('E7_ofc_reps.csv')
    rr0, cc0, rr1, cc1 = write_df(ws, E7[['feature_set', 'clf', 'rep', 'median_n_features', 'BA', 'MCC', 'Kappa', 'MacroF1', 'ROC_AUC', 'Sensitivity', 'Specificity']], 5, 14)
    sets = ['OFC-LDA', 'OFC-kNN', 'OFC-SVM', 'All DSFs', 'Full spectrum (SNV)']
    L = lambda k: get_column_letter(14 + ['feature_set', 'clf', 'rep', 'median_n_features', 'BA', 'MCC', 'Kappa', 'MacroF1', 'ROC_AUC', 'Sensitivity', 'Specificity'].index(k))
    R_ = lambda k: f'${L(k)}${rr0 + 1}:${L(k)}${rr1}'
    hdr = ['Feature set', 'Median n features', 'BA LDA', 'BA kNN', 'BA SVM', 'MCC LDA', 'MCC kNN', 'MCC SVM']
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(5, j, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = Alignment(wrap_text=True, horizontal='center')
    for i, s_ in enumerate(sets, start=6):
        ws.cell(i, 1, s_).font = BODY
        ws.cell(i, 2, f'=AVERAGEIFS({R_("median_n_features")},{R_("feature_set")},$A{i})').number_format = '0'
        for j, (met, c_) in enumerate([(m_, c2) for m_ in ('BA', 'MCC') for c2 in ('LDA', 'kNN', 'SVM')], start=3):
            ws.cell(i, j, f'=AVERAGEIFS({R_(met)},{R_("feature_set")},$A{i},{R_("clf")},"{c_}")').number_format = '0.000'
    ws.column_dimensions['A'].width = 22
    heat(ws, 'C6:E10'); heat(ws, 'F6:H10')
    bar(ws, 'A13', 'E7: balanced accuracy of OFCs vs all DSFs vs full spectrum', 1, [3, 4, 5], 5, 10, [COL_3, COL_H, COL_D], 'Balanced accuracy', ymin=0.6, ymax=1.0, w=20, h=9)
    FS = T('E7_ofc_final_sets.csv'); ws.cell(33, 1, 'Final OFCs selected on all data (cf. Tian et al., 2021, Table 4)').font = Font(name=FONT, bold=True)
    write_df(ws, FS, 34, 1, widths={'Feature': 22})
    BD = T('E7_ofc_by_date.csv'); r_bd = 36 + len(FS)
    ws.cell(r_bd, 1, 'Accuracy of each OFC model within each date').font = Font(name=FONT, bold=True)
    write_df(ws, BD, r_bd + 1, 1)
    FF = T('E7_ofc_feature_frequency.csv'); r_ff = r_bd + 4 + len(BD)
    ws.cell(r_ff, 1, 'Selection frequency of features in each OFC across the 100 outer folds').font = Font(name=FONT, bold=True)
    write_df(ws, FF, r_ff + 1, 1)
    ws.freeze_panes = None

for fn, nm, title, note, ref in [
    ('E2_transfer_reps.csv', '28_Transfer_E2', 'Step 10 - E2: transfer to an unseen date / variety / growth stage',
     'Leave-one-date-out and leave-one-stage-out keep plots separated (train: other levels & other plots; test: held-out level & held-out plots). Leave-one-variety-out: train 2 varieties, test the 3rd.', 'Roberts et al. (2017); Terentev et al. (2023); Lv et al. (2023)'),
    ('E3_metadata_reps.csv', '29_Metadata_E3', 'Step 10 - E3: spectra only vs spectra + metadata (date, variety, stage one-hot)', 'If the two are equal, the spectral model is context-independent.', ''),
    ('E4_learning_curve_reps.csv', '30_LearningCurve_E4', 'Step 10 - E4: learning curve (number of training plots)', 'Training plots sub-sampled per variety x class (1:3 healthy:diseased).', ''),
    ('E5_sensor_reps.csv', '31_Sensors_E5', 'Step 10 - E5: simulated multispectral sensors (Sentinel-2 MSI, UAV 5-band)', 'Rectangular band responses applied to smoothed reflectance; band reflectance + NDVI/NDRE/GNDVI(/NDII).', 'Zhang et al. (2012); Nguyen et al. (2023)'),
    ('E6_outlier_sensitivity_reps.csv', '32_Outliers_E6', 'Step 10 - E6: sensitivity to outlier removal (all 2339 spectra)', 'Compare with the main CV results of the same pipelines.', '')]:
    if not have(fn):
        continue
    ws = sheet(nm, title, note, ref)
    D_ = T(fn)
    gcols = [c for c in ['kind', 'level', 'scenario', 'clf', 'inputs', 'n_train_plots', 'sensor', 'data'] if c in D_.columns]
    g = D_.groupby(gcols, sort=False)
    S_ = g[['BA', 'OA', 'Sensitivity', 'Specificity', 'MacroF1', 'Kappa', 'MCC', 'ROC_AUC']].mean().join(g['BA'].std().rename('BA_SD')).reset_index()
    if 'level' in S_.columns:
        S_['level'] = S_['level'].map(lambda v: DATE_LABEL.get(v, v))
    r0, c0, r1, c1 = write_df(ws, S_, 5, 1)
    lab = S_[gcols].astype(str).agg(' | '.join, axis=1)
    ws.cell(5, c1 + 2, 'Label').font = HDR_FONT; ws.cell(5, c1 + 2).fill = HDR_FILL
    ws.cell(5, c1 + 3, 'BA').font = HDR_FONT; ws.cell(5, c1 + 3).fill = HDR_FILL
    for i, (l, b) in enumerate(zip(lab, S_['BA']), start=6):
        ws.cell(i, c1 + 2, l); ws.cell(i, c1 + 3, f'={get_column_letter(list(S_.columns).index("BA") + 1)}{i}').number_format = '0.000'
    ws.column_dimensions[get_column_letter(c1 + 2)].width = 34
    bar(ws, f'{get_column_letter(c1 + 5)}5', title.split(': ')[-1], c1 + 2, [c1 + 3], 5, r1, [COL_H], 'Balanced accuracy', horizontal=True, ymin=0.4, ymax=1.0, w=18, h=max(8, 0.45 * len(S_)))
    ws.cell(r1 + 3, 1, 'Repeat-level results').font = Font(name=FONT, bold=True)
    write_df(ws, D_, r1 + 4, 1)
    ws.freeze_panes = None

# ===================================================================== data sheet
d = np.load(os.path.join(CACHE, 'clean.npz'))
Mc = pd.read_csv(os.path.join(CACHE, 'meta_clean.csv'))
ws = sheet('33_Data_10nm', 'Cleaned dataset: splice-corrected, SG-smoothed reflectance at 10-nm resolution (2290 spectra after QC)',
           'Rows = spectra; metadata columns followed by reflectance 400-2400 nm.')
DD = pd.concat([Mc[['spec_uid', 'date', 'plot', 'variety', 'stage', 'label']], pd.DataFrame(d['B10'], columns=[f'R{w}' for w in WL10])], axis=1)
write_df(ws, DD, 5, 1, fmt={f'R{w}': '0.0000' for w in WL10})

# ===================================================================== references
ws = sheet('34_References', 'References cited in the workbook and presentation', 'Papers provided by the researcher are marked [P].')
REFS = pd.read_csv(os.path.join(ROOT, 'scripts', 'references.csv'))
write_df(ws, REFS, 5, 1, widths={'Reference': 160})
ws.freeze_panes = None

# ===================================================================== README
ws = readme
ws.sheet_view.showGridLines = False
ws['A1'] = 'Hyperspectral detection of white rust (Albugo candida) in Indian mustard - analysis workbook'; ws['A1'].font = Font(name=FONT, bold=True, size=15, color='1F4E2C')
lines = [
    'Data: ASD FieldSpec contact-probe spectra (350-2500 nm), 48 plots (3 varieties x 16 plots: 4 healthy + 12 diseased), 4 dates (11 Feb - 13 Mar 2026), 2339 spectra.',
    'Goal: one pooled model that separates healthy from white-rust leaves irrespective of date, variety, growth stage and leaf layer.',
    'Unit of independence = plot. All validation is plot-grouped; every learning step (MSC, date normalisation, feature selection, NDSI search, tuning) is fitted inside training folds only.',
    'Main validation: nested 4-fold plot-grouped CV stratified by class and variety (each test fold = 1 healthy + 3 diseased plots per variety), repeated 25 times (100 outer folds); inner 3-fold plot-grouped CV for selection and tuning.',
    "Colour code: blue font = editable input counts (sheet 22); black = formulas or Python-computed values. Sheets 17, 18, 20, 22 contain live formulas; other numbers were computed by the Python scripts in white_rust_analysis/scripts (named in each sheet).",
    'Assumption: the Healthy/Diseased label is the plot label provided in the source workbooks (plots 1-12 healthy, 13-48 diseased); leaf-level severity scores were empty, so severity models were not fitted.',
]
for i, l in enumerate(lines, start=3):
    ws.cell(i, 1, l).font = BODY
ws.cell(10, 1, 'Sheet').font = HDR_FONT; ws.cell(10, 1).fill = HDR_FILL
ws.cell(10, 2, 'Content').font = HDR_FONT; ws.cell(10, 2).fill = HDR_FILL
ws.cell(10, 3, 'Notes').font = HDR_FONT; ws.cell(10, 3).fill = HDR_FILL
for i, (n_, t_, no_) in enumerate(INDEX, start=11):
    c = ws.cell(i, 1, n_); c.hyperlink = f"#'{n_}'!A1"; c.font = Font(name=FONT, color='0563C1', underline='single')
    ws.cell(i, 2, t_).font = BODY; ws.cell(i, 3, no_).font = NOTE_FONT
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 95; ws.column_dimensions['C'].width = 120
wb.move_sheet('00_README', offset=-len(wb.sheetnames))
wb.save(OUT)
print('saved', OUT)
