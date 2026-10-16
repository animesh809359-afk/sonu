"""Builds the single results workbook (all steps, tables, figures, cross-checks)."""
from __future__ import annotations

import json

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.layout import Layout, ManualLayout

import xl
from common import ROOT, WORK, cfg, load, present_dates, subset_label, v3_label, v5_label, v5_level
from xl import F_BOLD, F_NOTE, FILL_KEY, Sheet, error_matrix

C = cfg()
D, A, K, Cc, X = load("data"), load("partA"), load("partK"), load("partC"), load("compare")
HAS_CR = (WORK / "partCR.pkl").exists() and (WORK / "compareCR.pkl").exists() and not C["data"].get("subset")
CR, XR = (load("partCR"), load("compareCR")) if HAS_CR else (None, None)
TXT = json.loads((WORK / "conclusions.json").read_text()) if (WORK / "conclusions.json").exists() else {}
FIG = WORK / "figures"
OUT = (ROOT / C["output"]["dir"]).resolve()
OUT.mkdir(parents=True, exist_ok=True)
spec, y, wl, meta, F = D["spec"], D["y"], D["spec"].wl, D["meta"], D["folds"]
nH, nD = int((y == 0).sum()), int((y == 1).sum())
REG = C["spectra"]["regions"]
SUBSET = subset_label(C)
IS_SUB = bool(C["data"].get("subset"))
PD = present_dates(D)
grp = D["groups"]
NP0, NP1 = len(set(grp[y == 0])), len(set(grp[y == 1]))
NSETS, K3, NV4 = len(F["V3"]), F["v3_k"], len(F["V4"])
NREP = len({f["rep"] for f in F["V3"]})
RATIO = nD / nH
V5BY = F.get("V5_by", "Date")
NOISE = float(pd.DataFrame(D["audit"]["noise"]).iloc[0, 2])
NDUP_UNNAMED = len(D["dup_rows"])
_keys = set(zip(meta["Date"].astype(str), meta["Plot"].astype(int), meta["File"].astype(str)))
B2 = pd.DataFrame([{"Date": l_["date"], **r} for l_ in D["logs"] for r in l_["block2"]])
if len(B2):
    B2 = B2[[(d, int(p_), f_) in _keys for d, p_, f_ in zip(B2["Date"], B2["Plot"], B2["Block-1 file"])]]
NBLOCK2 = len(B2)
if V5BY == "Date":
    V5_TXT = (f"V5 = leave-one-date-out on new plots ({len(F['V5'])} units: {NV4} V4 splits × {len(F['V5_dates'])} dates"
              + (f", {len(F['V5_dropped'])} units without usable test leaves skipped" if F.get("V5_dropped") else "") + ")")
else:
    V5_TXT = f"V5 = leave-one-variety-out ({len(F['V5'])} units: train on the other varieties, test on all leaves of the held-out variety)"
DATES_TXT = ", ".join(PD)
wb = Workbook()
wb.remove(wb.active)
SHEETS = []
PCT = "0.0%"


def S_(name, title, sub="", **kw):
    s = Sheet(wb, name, title, sub, **kw)
    SHEETS.append((name, title))
    if TXT.get("readings", {}).get(name):
        s.h2("Result in brief")
        for t in TXT["readings"][name]:
            s.para(t)
    return s


def vrow(R, pipe, sch):
    return R["validation"].set_index(["Pipeline", "Scheme"]).loc[(pipe, sch)]


def bands_txt(lst, prefix=("WV", "R")):
    out = []
    for v in lst:
        s = str(v)
        for p in prefix:
            if s.startswith(p) and s[len(p):].isdigit():
                s = s[len(p):]
        out.append(s)
    return ", ".join(out)


def checks_table(S, rows, title="Cross-checks for this step"):
    if not rows:
        return
    df = pd.DataFrame(rows, columns=["Step", "What was checked", "Method", "Result", "Status"])
    S.table(df, title=title, flag_col="Status", widths={"B": 30, "C": 34, "D": 44})


ALLCHECKS = []
cA, cK, cC = A["checks"], K["checks"], Cc["checks"]

# =====================================================================================
# 00 README (filled at the end with the sheet guide)
readme = S_("00_README", f"Healthy vs white-rust-diseased mustard leaves — {SUBSET}",
            "Four-tier hierarchical analysis of Das et al. (2020, Geocarto International 35:64–77) with imbalance and "
            "pseudo-replication corrections; CART vs Random Forest; disease-sensitive features (Part K); four-tier "
            "analysis of the combined feature pool (Part C: K1 + K2–K4" + ("; Part C′: raw reflectance + K2–K4" if HAS_CR else "") + "). Leaf level, ASD FieldSpec contact probe.")

# =====================================================================================
# 01 Summary
S = S_("01_Summary", "Summary — results of every step", "Balanced accuracy (BA) = mean of healthy and diseased recall; "
       f"it is the headline metric because the classes are imbalanced (1 : {RATIO:.1f}).")
S.h2("Key numbers")
v3a = vrow(A, "QDA on CART features", "V3")
key = [
    ("Leaf spectra analysed (healthy / diseased)", f"{len(y)} ({nH} / {nD}); {NP0 + NP1} plots ({NP0} / {NP1}); {len(PD)} date(s): {DATES_TXT}; growth stage(s): {', '.join(sorted(meta['Growth stage'].unique()))}; variety(ies): {', '.join(sorted(meta['Variety'].unique()))}"),
    ("Tier 1 — bands significant (Welch + FDR)", f"{A['n_sig']} of {A['n_features']}"),
    ("Tier 2 — CART bands (rpart, prior 0.5/0.5)", f"{len(A['sets'][A['cart_name']])}: {bands_txt(A['sets'][A['cart_name']])} nm"),
    (f"Tier 2 — RF bands (top-{len(A['sets'][A['rf_name']])} MDI peaks)", f"{bands_txt(A['sets'][A['rf_name']])} nm"),
    ("Tier 3 — J–M healthy vs diseased (CART / RF bands)",
     f"{A['tier3'].set_index('Set').loc[A['cart_name'], 'JM']:.3f} / {A['tier3'].set_index('Set').loc[A['rf_name'], 'JM']:.3f} (threshold 1.90)"),
    ("Tier 4 — QDA equal priors, 10-fold CV (V1): BA CART / RF bands",
     f"{100 * A['errmat'][A['cart_name']]['BA']:.1f} % / {100 * A['errmat'][A['rf_name']]['BA']:.1f} %"),
    ("Honest estimate (V3 nested plot-grouped CV): QDA on CART bands", f"BA {100 * v3a['BA']:.1f} ± {100 * v3a['BA_SD']:.1f} %"),
    ("Part C — combined pool: features / Tier-1 significant", f"{Cc['n_features']} / {Cc['n_sig']}"),
    ("Part C — CART features", f"{len(Cc['sets'][Cc['cart_name']])}: {', '.join(Cc['sets'][Cc['cart_name']])}"),
    (f"Part C — QDA on {Cc['rf_name']} features (V3 / V4)",
     f"{100 * vrow(Cc, 'QDA on RF features', 'V3')['BA']:.1f} % / {100 * vrow(Cc, 'QDA on RF features', 'V4')['BA']:.1f} %"),
] + ([
    ("Part C′ — raw reflectance + K2–K4: features / Tier-1 significant", f"{CR['n_features']} / {CR['n_sig']}"),
    ("Part C′ — CART features", f"{len(CR['sets'][CR['cart_name']])}: {', '.join(CR['sets'][CR['cart_name']])}"),
    (f"Part C′ — QDA on {CR['rf_name']} features (V3 / V4)",
     f"{100 * vrow(CR, 'QDA on RF features', 'V3')['BA']:.1f} % / {100 * vrow(CR, 'QDA on RF features', 'V4')['BA']:.1f} %"),
] if HAS_CR else [])
S.table(pd.DataFrame(key, columns=["Item", "Result"]), widths={"A": 52, "B": 90})
S.h2("Table Z.1  All approaches side by side (ranked by the honest V3 balanced accuracy)")
syn = X["synthesis"].copy()
S.table(syn, fmts={c: PCT for c in syn.columns if c.startswith(("BA", "SD"))} | {"AUC V3": "0.000", "κ V3": "0.000"},
        note="V1 = leaf-level 10-fold CV with features chosen on all leaves (the paper's protocol; optimistic). "
             f"V3 = nested plot-grouped CV ({v3_label(F)}), V4 = 30 % of plots held out × {NV4}, {V5_TXT}; "
             "in V3–V5 every selection step is repeated inside the training set. Single features: threshold chosen on training plots.")
S.image(FIG / "F29_synthesis.png", 820)
if TXT.get("conclusions"):
    S.h2("Conclusions")
    for i, t in enumerate(TXT["conclusions"], 1):
        S.para(f"{i}. {t}")
if TXT.get("recommendations"):
    S.h2("Recommendations")
    for t in TXT["recommendations"]:
        S.para(t)

# =====================================================================================
# 02 Design
S = S_("02_Design", "Workflow, validation design and model settings",
       "Figure 2 of Das et al. (2020), extended with the corrections and the CART-vs-RF, Part K and Part C steps.")
S.image(FIG / "F01_workflow.png", 860)
S.h2("Table 2.1  Corrections applied to the protocol of Das et al. (2020) and why")
corr = pd.DataFrame([
    (f"Class imbalance ({nH} healthy : {nD} diseased = 1 : {RATIO:.1f})", "QDA/LDA with equal prior probabilities; rpart prior = c(0.5, 0.5); Balanced RF (each tree on a balanced bootstrap); balanced accuracy as headline metric; 1000 balanced subsets in Tier 1",
     "without it every classifier favours 'diseased' and overall accuracy is inflated"),
    ("Two classes only", "Welch t-test instead of one-way ANOVA + Tukey HSD (ANOVA reported for comparison: identical ranking)",
     "Tukey HSD is for many pairs; Welch allows unequal variances (Levene)"),
    (f"Multiple testing ({len(wl)} bands)", "Benjamini–Hochberg FDR (q < 0.05); uncorrected counts reported",
     f"{len(wl)} tests at α = 0.05 give ≈ {round(0.05 * len(wl))} false positives"),
    (f"Pseudo-replication (leaves nested in plots" + (f", plots measured on {len(PD)} dates)" if len(PD) > 1 else ")"), "mixed model with random plot intercept; every honest validation grouped by plot" + (" across all dates" if len(PD) > 1 else ""),
     "leaves of one plot are not independent"),
    ("Selection bias", "Tier 1, band selection and classifier refitted inside every training set (nested V3, V4, V5)",
     "selecting bands on all leaves and validating on the same leaves is optimistic"),
    ("Water-absorption bands", "kept (contact probe with its own light source; no atmospheric path)", "Das et al. removed them for canopy measurements"),
    ("Detector edge 350–399 nm", f"removed (noise {NOISE:.0f} × the 500–699 nm level); sensitivity run with 350–2500 nm", "noise could be chosen by CART"),
    ("Detector joins 1000/1001 and 1800/1801 nm (Part K)", "splice correction; guard zones excluded from feature ranking (reported with and without)", "instrument artefacts must not be read as disease"),
    ("Exact duplicate spectra", "75 copies (13-03-2026, block 2) and 4 unnamed copies (plot 27, 11-02-2026) removed from the four files before the subset was taken"
     + (f"; {NBLOCK2 + NDUP_UNNAMED} of them belong to this subset ({NBLOCK2} block-2, {NDUP_UNNAMED} unnamed)" if IS_SUB else "")
     + ("; sensitivity run with the unnamed copies kept" if A.get("sensitivity_dups_applicable", True) else ""), "a duplicate in both training and test inflates V1"),
    *([("Subset", f"{SUBSET}: {len(y)} of {D['audit']['subset']['n_all']} leaf spectra; folds, Tier 1–4, Part K and Part C rebuilt for this subset only",
        "each date, stage and variety analysed separately as requested")] if IS_SUB else []),
], columns=["Issue", "Correction", "Why"])
S.table(corr, widths={"A": 34, "B": 70, "C": 50})
S.h2("Table 2.2  Validation design (identical folds in every step; sheet 36_Folds)")
vd = pd.DataFrame([
    ("V1 · leaf-level 10-fold CV", "StratifiedKFold, seed 42", "bands fixed (chosen on all leaves); classifier refitted", "reproduces the paper (optimistic: leaves of a plot in train and test)"),
    (f"V2 · plot-grouped {K3}-fold CV", f"{F['v3_method']}, seed 42", "bands fixed; classifier refitted", "removes plot leakage, keeps selection bias"),
    (f"V3 · nested plot-grouped CV, {v3_label(F)}", f"{F['v3_method']}, seeds 101–105; k = min(10, plots of the smaller class = {min(NP0, NP1)}) so every test fold holds both classes",
     "Tier 1 + selection + classifier refitted", "honest estimate for new plots"),
    (f"V4 · hold-out of 30 % of plots × {NV4}", f"{F['v4_plots'][0]} healthy + {F['v4_plots'][1]} diseased plots per split, default_rng(2026)", "everything refitted on the remaining plots", "independent test plots"),
    ((f"V5 · new date and new plots (additional)", "train: other dates × V4 training plots; test: held-out date × V4 test plots", "everything refitted", "transfer to an unseen date and unseen plots")
     if V5BY == "Date" else
     ("V5 · new variety (additional; single date, so no date transfer)", "train: the other varieties; test: all leaves of the held-out variety", "everything refitted", "transfer to an unseen variety (its plots are unseen too)")),
], columns=["Scheme", "Folds", "What is refitted", "Purpose"])
S.table(vd, widths={"D": 50})
S.h2("Table 2.3  CART and Random Forest settings")
st_ = pd.DataFrame([
    ("Algorithm", "rpart 4.1.23 in R 4.3.3 (recursive partitioning, as Das et al.)", "BalancedRandomForestClassifier (imbalanced-learn 0.14, scikit-learn 1.9)"),
    ("Split criterion", "Gini", "Gini"),
    ("Imbalance correction", "prior = c(0.5, 0.5)", "sampling_strategy = 'all', replacement = True, bootstrap = False (balanced bootstrap per tree)"),
    ("Size / randomness", "minsplit 20, minbucket 7, cp 0.01, maxcompete 0, maxsurrogate 0", "500 trees (1000 for importance), max_features = √p, random_state 42"),
    ("Pruning", f"default tree used; 1-SE tree from plot-grouped {K3}-fold CV of cp", "—"),
    ("Band selection", "split variables of the tree", "top-17 / top-10 MDI importance peaks (+ a set the size of the CART set), local maxima ≥ 10 nm apart"),
    ("Prediction", "class of the terminal node (prior-adjusted)", "majority vote (probability > 0.5)"),
], columns=["Setting", "CART", "Random Forest"])
S.table(st_, widths={"B": 55, "C": 60})

# =====================================================================================
# 03 Data audit
S = S_("03_Step0_Data_Audit", "Step 0 · Data audit — what was analysed",
       "Four ASD export workbooks (one column per leaf spectrum); plot-average columns ignored; exact duplicates removed.")
S.h2("What was done")
S.para("The four files (11-02-2026, 24-02-2026, 03-03-2026, 13-03-2026) were read; each spectrum column carries plot number, "
       "variety, growth stage and the healthy/diseased label. The 48+ numbered columns per file are =AVERAGE() plot means and were "
       "ignored (leaf level only). On 13-03-2026 a second data block (75 columns) repeats plots 13, 34, 38, 41 and 42 exactly; "
       "on 11-02-2026 four unnamed columns of plot 27 are exact copies of Mustard00215–00218. These 79 copies were removed. "
       "Leaf level and disease score were not used (instruction: healthy vs diseased only)."
       + (f" Only the subset '{SUBSET}' was then kept: {len(y)} of {D['audit']['subset']['n_all']} unique leaf spectra "
          f"({nH} healthy, {nD} diseased; dates {DATES_TXT}). Every following step uses this subset only." if IS_SUB else ""))
cnt = D["audit"]["counts"].copy()
f, l, L = S.table(cnt, title=f"Table 0.1  Leaf spectra by class ({SUBSET})")
S.cell(S.row - 1, 1, "Total", F_BOLD)
S.cell(S.row - 1, 2, f"=SUM(B{f}:B{l})")
S.cell(S.row - 1, 3, f"=SUM(C{f}:C{l})")
S.cell(S.row - 1, 4, "Share healthy (Excel)", F_BOLD)
S.cell(S.row - 1, 5, f"=C{f}/C{l + 1}", fmt=PCT)
S.cell(S.row - 1, 6, "Diseased : healthy", F_BOLD)
S.cell(S.row - 1, 7, f"=C{l}/C{f}", fmt="0.00")
S.row += 1
bd = D["audit"]["by_date"].reset_index()
f, l, L = S.table(bd, title="Table 0.2  Leaf spectra by date and class")
S.cell(S.row - 1, 1, "Total", F_BOLD)
S.cell(S.row - 1, 2, f"=SUM(B{f}:B{l})")
S.cell(S.row - 1, 3, f"=SUM(C{f}:C{l})")
S.row += 1
S.table(D["audit"]["by_stage"].reset_index(), title="Table 0.3  Leaf spectra by date, growth stage and class")
S.table(D["audit"]["by_stage_total"].reset_index(), title="Table 0.4  Leaf spectra by growth stage")
S.table(D["audit"]["by_variety"].reset_index(), title="Table 0.5  Leaf spectra by variety")
S.table(pd.DataFrame(D["audit"]["checks"], columns=["Check", "Result", "Status"]), title="Table 0.6  Data-quality checks",
        flag_col="Status", widths={"A": 46, "B": 50})
if len(D["dup_rows"]):
    S.table(pd.DataFrame(D["dup_rows"]), title="Table 0.7  Exact duplicates removed among the named/unnamed columns (in this subset)")
else:
    S.para("Table 0.7  No unnamed duplicate spectrum belongs to this subset (the 4 removed copies are plot 27, 11-02-2026).", font=F_NOTE)
if len(B2):
    S.table(B2, title="Table 0.8  13-03-2026 second data block — each removed spectrum and its identical block-1 copy"
            + (" (copies whose block-1 spectrum belongs to this subset)" if IS_SUB else ""), max_rows=75)
elif IS_SUB:
    S.para("Table 0.8  None of the 75 block-2 copies of 13-03-2026 belongs to this subset.", font=F_NOTE)
pp = D["audit"]["per_plot"].copy()
S.table(pp, title="Table 0.9  Plots: class, variety and leaf spectra per date")
S.table(D["audit"]["stage_by_plot_date"].reset_index(), title="Table 0.10  Growth stage of every plot on every date")

# =====================================================================================
# 04 Preprocessing
S = S_("04_Step1_Preprocessing", "Step 1 · Pre-processing — band set, noise and outlier screening",
       "Das et al. (2020) removed 1350–1430, 1800–1950 and 2350–2500 nm (atmospheric water, canopy). Here the contact probe has its own light source, so water bands are kept.")
S.h2("What was done")
S.para("Noise of every band = SD of the residual after Savitzky–Golay smoothing (window 11, order 2). The 350–399 nm detector edge "
       f"is far noisier than the rest and was removed; everything else (400–2500 nm, {len(wl)} bands) was analysed on the raw "
       "(unsmoothed) reflectance in Tiers 1–4, as in Das et al. Outliers were screened with the spectral angle (SAM) to the class "
       "median spectrum and flagged, not removed (diseased leaves are expected to vary more).")
bs = pd.DataFrame([("350–399", 50, "Removed", f"detector edge: noise ×{NOISE:.0f} of the 500–699 nm level"),
                   ("400–1349", 950, "Retained", "visible, red edge, NIR"),
                   ("1350–1430", 81, "Retained", "water band I — contact probe (Das et al. removed it)"),
                   ("1431–1799", 369, "Retained", "SWIR-1"),
                   ("1800–1950", 151, "Retained", "water band II — contact probe (Das et al. removed it)"),
                   ("1951–2349", 399, "Retained", "SWIR-2"),
                   ("2350–2500", 151, "Retained", "SWIR-2 end — contact probe (Das et al. removed it)")],
                  columns=["Range (nm)", "Bands", "Decision", "Reason"])
f, l, L = S.table(bs, title="Table 1.1  Band-set decision")
S.cell(S.row - 1, 1, "Bands analysed (Excel)", F_BOLD)
S.cell(S.row - 1, 2, f'=SUMIF(C{f}:C{l},"Retained",B{f}:B{l})')
S.row += 1
S.table(pd.DataFrame(D["audit"]["noise"]), title="Table 1.2  Noise level by spectral region")
S.table(pd.DataFrame(D["audit"]["sam"]), title="Table 1.3  Outlier screening: spectral angle to the class median")
S.table(D["audit"]["top_dev"], title="Table 1.4  Ten most deviating spectra (flagged, kept)")
S.h2("Figure 3 (as in Das et al. 2020). Mean reflectance of healthy and diseased leaves")
S.image(FIG / "F02_mean_spectra.png", 820)
S.image(FIG / "F02b_mean_spectra_by_date.png", 900)
S.para("Native Excel chart of the same data: sheet 37_ChartData, columns A–C.", font=F_NOTE)
if A.get("sensitivity_dups_applicable", True):
    S.table(A["sensitivity_dups"], title="Table 1.5  Sensitivity: the unnamed duplicate spectra of plot 27 kept vs removed (QDA on CART bands)",
            fmts={"BA V1": PCT, "BA V3": PCT, "BA V4": PCT})
else:
    S.para("Table 1.5  Duplicate sensitivity run not applicable: the 4 unnamed plot-27 copies (11-02-2026) are not part of this subset.", font=F_NOTE)

# =====================================================================================
# 05 Tier 1
t1 = A["tier1"]
S = S_("05_Tier1_Ttest", "Step 2 · Tier 1 — Welch's t-test at every band with Benjamini–Hochberg FDR",
       "Das et al. (2020) used one-way ANOVA (p < 0.05) + Tukey HSD for 14 genotypes; with two classes the equivalent is a t-test.")
S.h2("What was done")
S.para(f"At each of the {len(wl)} bands the mean reflectance of healthy (n = {nH}) and diseased (n = {nD}) leaves was compared "
       "with Welch's t-test (unequal variances; Levene's test shows unequal variances at many bands) and the p-values were "
       f"adjusted with the Benjamini–Hochberg false-discovery rate. Bands with q < 0.05 go to Tier 2. Because {nD} diseased "
       f"leaves give high power even for tiny differences, the test was repeated on 1000 balanced random subsets ({nH} vs {nH}), "
       "and a linear mixed model with a random plot intercept checks pseudo-replication.")
tt = A["tier1_tests"].copy()
f, l, L = S.table(tt, title="Table 2.1  Number of significant bands")
S.cell(f - 1, 3, "Bands tested", xl.F_HEAD, fill=xl.FILL_HEAD, border=True)
S.cell(f - 1, 4, "Share (Excel)", xl.F_HEAD, fill=xl.FILL_HEAD, border=True)
for r in range(f, l + 1):
    S.cell(r, 3, len(wl), border=True)
    S.cell(r, 4, f"=B{r}/C{r}", fmt=PCT, border=True)
S.table(A["tier1_ranges"], title="Table 2.2  Contiguous ranges (Welch + FDR)")
S.table(A["tier1_regions"], title="Table 2.3  Significance by spectral region")
top = t1.reindex(t1["t"].abs().sort_values(ascending=False).index).head(15)
top15 = pd.DataFrame({"Wavelength (nm)": top["position"], "Mean healthy": top["mean0"], "Mean diseased": top["mean1"],
                      "Diff (D − H)": top["mean1"] - top["mean0"], "Welch t": top["t"], "Welch df": top["df"],
                      "q (BH)": top["q"], "Hedges g": top["g"], "Mixed-model q": top["mm_q"]})
S.table(top15, title="Table 2.4  Fifteen most significant bands")
bb = A["tier1_balanced"]
S.table(pd.DataFrame([("Balanced random subsets drawn", bb["subsets"]), ("Leaves per subset (healthy + diseased)", bb["per_subset"]),
                      ("Significant bands per subset — median", bb["median"]), ("— minimum", bb["min"]), ("— maximum", bb["max"]),
                      ("Bands significant in 100 % of subsets", bb["all"]), ("Bands significant in ≥ 95 % of subsets", bb["ge95"]),
                      ("Bands never significant", bb["never"])], columns=["Balanced-subset check (imbalance)", "Value"]),
        title="Table 2.5  Tier 1 on 1000 balanced subsets")
S.h2("Figure 4 (as in Das et al. 2020 Fig. 4). Bands where healthy and diseased leaves differ (grey)")
S.image(FIG / "F03_tier1_significant_regions.png", 820)
S.h2("Figure 5 (as in Das et al. 2020 Fig. 5). Significance frequency (balanced subsets) and mean reflectance")
S.image(FIG / "F04_tier1_frequency.png", 820)
S.h2("Figure 5b. −log₁₀(p) against wavelength")
S.image(FIG / "F05_minus_log10_p.png", 820)
S.para("Native Excel versions of Figures 3–5b: sheet 37_ChartData (charts placed beside the data).", font=F_NOTE)
checks_table(S, [c for c in cA if c[0] == "Tier 1"])

# 05b all bands (with Excel formulas)
S = S_("05b_Tier1_AllBands", f"Step 2 · Tier 1 — full results for all {len(wl)} bands",
       "Values from Python/SciPy; columns marked (Excel) are formulas recomputing the difference, % difference, significance and the Welch t from the class means and SDs.")
ab = pd.DataFrame({"Wavelength (nm)": wl, "Region": t1["region"], "n healthy": nH, "n diseased": nD,
                   "Mean healthy": t1["mean0"], "SD healthy": t1["sd0"], "Mean diseased": t1["mean1"], "SD diseased": t1["sd1"],
                   "Diff D − H (Excel)": "", "% diff (Excel)": "", "Welch t": t1["t"], "Welch df": t1["df"], "p": t1["p"],
                   "q (BH-FDR)": t1["q"], "Significant q<0.05 (Excel)": "", "Hedges g": t1["g"], "ANOVA F": t1["anova_F"],
                   "ANOVA p": t1["anova_p"], "Levene p": t1["levene_p"], "Mixed-model effect": t1["mm_effect"],
                   "Mixed-model p": t1["mm_p"], "Mixed-model q": t1["mm_q"], "Mixed model significant": np.where(t1["mm_sig"], "Yes", "No"),
                   "Balanced subsets significant (%)": t1["balanced_pct"], "AUC": t1["auc"], "Welch t (Excel)": "", "|t − t(Excel)| (Excel)": ""})
f, l, L = S.table(ab)
for r in range(f, l + 1):
    S.cell(r, 9, f"=G{r}-E{r}", fmt="0.00000", border=True)
    S.cell(r, 10, f"=100*(G{r}-E{r})/E{r}", fmt="0.00", border=True)
    S.cell(r, 15, f'=IF(N{r}<0.05,"Yes","No")', border=True)
    S.cell(r, 26, f"=(G{r}-E{r})/SQRT(H{r}^2/D{r}+F{r}^2/C{r})", fmt="0.0000", border=True)
    S.cell(r, 27, f"=ABS(K{r}-Z{r})", fmt="0.0E+00", border=True)
S.cell(2, 14, "Max |t − t(Excel)| =", F_BOLD)
S.cell(2, 16, f"=MAX(AA{f}:AA{l})", fmt="0.0E+00")
S.cell(3, 14, "Significant (Excel count) =", F_BOLD)
S.cell(3, 16, f'=COUNTIF(O{f}:O{l},"Yes")')
S.ws.freeze_panes = S.ws.cell(row=f, column=2)

# =====================================================================================
# 06 Tier 2 CART
cart_set = A["sets"][A["cart_name"]]
S = S_("06_Tier2_CART", "Step 3 · Tier 2 — CART band selection with equal priors (rpart)",
       "Das et al. (2020): 'To apply this technique we used the method rpart from package rpart' — the split variables are the selected wavelengths.")
S.h2("What was done")
S.para(f"Input: the {A['n_sig']} FDR-significant bands of Tier 1 (raw reflectance). rpart (R 4.3.3, rpart 4.1.23) with method = "
       "'class', prior = c(0.5, 0.5) (imbalance correction), split = 'gini', minsplit 20, minbucket 7, cp 0.01, no competitor or "
       "surrogate splits. The wavelengths used as split variables form the CART band set. The complexity parameter was "
       "cross-validated with plot-grouped folds; the 1-SE tree gives a smaller alternative set. Stability: rpart refitted on 200 "
       f"plot-level bootstrap samples and inside the {NSETS} nested training sets of V3 (a band counts as re-selected if a band within "
       "± 10 nm is chosen).")
S.table(A["cart_table"], title=f"Table 3.1  The {len(cart_set)} wavelengths selected by CART",
        fmts={"Bootstrap re-selection (±10 nm)": PCT, "Nested-CV re-selection (±10 nm)": PCT}, widths={"B": 24, "L": 44})
S.h2(f"Figure 6 (as in Das et al. 2020 Fig. 6). CART tree")
S.image(FIG / "F06_cart_tree.png", 1000)
fr = A["cart_full"]["frame"][["node", "depth", "var", "label", "n", "nH", "nD", "pred"]].copy()
fr.columns = ["Node", "Depth", "Split variable", "Rule (left branch of the parent)", "Leaves", "Healthy", "Diseased", "Predicted class"]
S.table(fr, title="Table 3.2  Tree nodes (rpart frame)")
S.table(A["cart_full"]["cptable"].rename(columns={"CP": "CP", "nsplit": "Number of splits", "rel error": "Relative error",
                                                  "xerror": "Cross-validated error (plot-grouped)", "xstd": "SE of CV error"}),
        title=f"Table 3.3  Complexity-parameter table (plot-grouped {K3}-fold CV)")
S.image(FIG / "F07_cart_cp.png", 600)
sets_rows = [(k, len(v), bands_txt(v)) for k, v in A["sets"].items() if k.startswith("CART")]
S.table(pd.DataFrame(sets_rows, columns=["Tree", "Bands", "Wavelengths (nm)"]), title="Table 3.4  Selected band sets", widths={"C": 80})
if "frame_1se" in A["cart_full"]:
    S.image(FIG / "F06b_cart_tree_1se.png", 700)
S.table(A["cart_sklearn"], title="Table 3.5  Cross-check of the CART result with scikit-learn")
S.h2("Figure 6b. Re-selection frequency of each CART band")
S.image(FIG / "F08_cart_reselection.png", 760)
checks_table(S, [c for c in cA if c[0] == "Tier 2"])

# =====================================================================================
# 07 Tier 3
S = S_("07_Tier3_JM", "Step 4 · Tier 3 — Jeffries–Matusita (J–M) distance",
       "J–M = 2(1 − e^(−α)) on the 0–2 scale of the thresholds 1.90 / 1.94 used by Das et al.; √J–M (Eq. 1 printed in the paper, 0–√2) is also given.")
S.h2("What was done")
S.para("α (Bhattacharyya distance) = ⅛ (μH − μD)ᵀ [(ΣH + ΣD)/2]⁻¹ (μH − μD) + ½ ln( |(ΣH + ΣD)/2| / √(|ΣH| |ΣD|) ), computed with "
       "class means and covariance matrices of the selected bands (Cholesky factorisation; checked against explicit inverse and "
       "determinants). Classes are 'separable' when J–M ≥ 1.90 (Das et al. used ≥ 1.94, 97 % of 2).")
jmv = A["tier3"].set_index("Set").loc[A["cart_name"], "JM"]
S.table(pd.DataFrame([("Healthy", "–", "–"), ("Diseased", jmv, "–")], columns=["", "Healthy", "Diseased"]),
        title=f"Table 1 (as in Das et al. 2020). J–M distance between healthy and diseased leaves ({A['cart_name']} bands)")
t3 = A["tier3"][A["tier3"]["Set"].str.startswith("CART")][["Set", "Features", "alpha"]].copy()
t3.columns = ["Band set", "Bands", "α (Bhattacharyya)"]
f, l, L = S.table(t3, title="Table 4.1  Healthy vs diseased separability (formulas: J–M, √J–M and the threshold tests)")
hdrs = ["J–M (0–2) (Excel)", "√J–M (Excel)", "≥ 1.90 ? (Excel)", "≥ 1.94 ? (Excel)"]
for j, h in enumerate(hdrs):
    S.cell(f - 1, 4 + j, h, xl.F_HEAD, fill=xl.FILL_HEAD, border=True, align=xl.CENTER)
for r in range(f, l + 1):
    S.cell(r, 4, f"=2*(1-EXP(-C{r}))", fmt="0.0000", border=True)
    S.cell(r, 5, f"=SQRT(D{r})", fmt="0.0000", border=True)
    S.cell(r, 6, f'=IF(D{r}>=1.9,"separable","not separable")', border=True)
    S.cell(r, 7, f'=IF(D{r}>=1.94,"separable","not separable")', border=True)
inc = A["tier3_incremental"][["Features used", "CART feature added", "CART J–M", "CART QDA BA"]].dropna()
S.table(inc, title="Table 4.2  J–M and QDA BA as the CART bands are added (order of first use in the tree)",
        fmts={"CART QDA BA": PCT})
S.image(FIG / "F09_tier3_tier4_incremental.png", 900)
c0 = A["tier3"].set_index("Set").loc[A["cart_name"]]
S.table(pd.DataFrame([("Condition number of healthy covariance", c0["cond0"]), ("Condition number of diseased covariance", c0["cond1"]),
                      ("J–M of the healthy class with itself (must be 0)", A["jm_self"])], columns=["Diagnostic", "Value"]),
        title="Table 4.3  Numerical diagnostics")
checks_table(S, [c for c in cA if c[0] == "Tier 3"])

# =====================================================================================
# 08 Tier 4
S = S_("08_Tier4_QDA", "Step 5 · Tier 4 — Quadratic discriminant analysis (QDA) with equal priors",
       "Das et al. (2020): QDA with 10-fold cross-validation on the CART wavelengths (R package DiscriMiner); error matrix with producer's and user's accuracy.")
S.h2("What was done")
S.para(f"QDA (scikit-learn; MASS::qda in R as a check) on the {len(cart_set)} CART bands, with equal prior probabilities "
       "(imbalance correction) and leaf-level stratified 10-fold cross-validation (V1, folds in sheet 36). Uncorrected sample "
       "priors and LDA are shown for comparison. Totals, producer's/user's accuracy, OA, BA and κ below are Excel formulas from "
       "the four counts.")
error_matrix(S, A["errmat"][A["cart_name"]], f"Table 2 (as in Das et al. 2020). Error matrix, 10-fold CV QDA, {A['cart_name']} bands (rows = classified, columns = reference)")
t4 = A["tier4"][A["tier4"]["Set"].str.startswith("CART")][["Set", "Classifier", "HH", "HD", "DH", "DD", "OA", "BA", "PA_H", "PA_D", "UA_H", "UA_D", "kappa", "AUC"]].copy()
t4.columns = ["Band set", "Classifier", "Healthy → healthy", "Healthy → diseased", "Diseased → healthy", "Diseased → diseased",
              "OA", "BA", "PA healthy", "PA diseased", "UA healthy", "UA diseased", "κ", "AUC"]
S.table(t4, title="Table 5.1  Effect of the imbalance correction", fmts={c: PCT for c in ["OA", "BA", "PA healthy", "PA diseased", "UA healthy", "UA diseased"]},
        note="Equal priors raise the recall of the minority (healthy) class at the cost of overall accuracy; sample priors do the reverse.")
_mc = pd.DataFrame(A["mass_check"], columns=["Band set", "Priors", "Share of identical predictions (Python vs MASS::qda)"])
if A.get("mass_refused"):
    _mc["Leaves not classified by MASS::qda (rank-deficient fold)"] = [c for *_, c in A["mass_refused"]]
S.table(_mc,
        title="Table 5.2  Cross-check of the QDA predictions in R", fmts={"Share of identical predictions (Python vs MASS::qda)": "0.00%"})
checks_table(S, [c for c in cA if c[0] == "Tier 4"])

# =====================================================================================
# 09 Validation
S = S_("09_Validation", "Step 6 · Validation of the complete four-tier pipeline",
       "From the protocol of Das et al. (bands chosen on all leaves, leaf-level CV) to honest estimates for new plots and new dates.")
S.h2("What was done")
S.para("Das et al. selected bands on all data and validated QDA by 10-fold CV on the same data (V1). Here the pipeline is also "
       f"validated with folds grouped by plot (V2), with the whole pipeline nested inside plot-grouped folds (V3, {v3_label(F)}), on "
       f"independent hold-out plots (V4, {NV4} splits), with random under-sampling instead of equal priors (V3-RUS) and "
       + ("on an unseen date in unseen plots (V5, additional)" if V5BY == "Date" else "on an unseen variety (V5, additional)")
       + ". All rows: QDA on CART-selected bands.")
vq = A["validation_qda_cart"][["Scheme", "Priors", "repeats", "OA", "BA", "BA_SD", "kappa", "AUC", "PA_H", "PA_D", "UA_H", "UA_D"]].copy()
S.table(vq, title="Table 6.1  Validation results (mean over repeats; SD over repeats)",
        fmts={c: PCT for c in ["OA", "BA", "BA_SD", "PA_H", "PA_D", "UA_H", "UA_D"]})
S.image(FIG / "F11_validation_qda_cart.png", 820)
for sch, lab in (("V3", f"nested plot-grouped CV, {NREP} repeats pooled"), ("V4", f"independent hold-out plots, {NV4} splits pooled"),
                 ("V5", f"{v5_label(F)}, {len(F['V5'])} units pooled")):
    p = A["pooled"][("QDA on CART features", sch)]
    error_matrix(S, p, f"Table 6.{2 + ['V3', 'V4', 'V5'].index(sch)}  Error matrix, {sch} ({lab}) — QDA on CART bands, equal priors")
S.table(A["sensitivity_range"], title="Table 6.5  Sensitivity to the band set (full pipeline repeated)",
        fmts={"BA V1": PCT, "BA V3 nested": PCT, "BA V4 hold-out": PCT}, widths={"F": 70})

# =====================================================================================
# 10 A1
S = S_("10_A1_RF_Importance", "Part A1 · Random Forest as the Tier-2 band selector (instead of CART)",
       f"Input: the {A['n_sig']} FDR-significant bands of Tier 1 (same as CART).")
S.h2("What was done")
S.para(f"A Balanced Random Forest (1000 trees, √{A['n_sig']} candidate bands per split, each tree grown on a balanced bootstrap "
       f"of {nH} + {nH} leaves) was fitted to all leaves; band importance = mean decrease in Gini impurity (MDI). Because "
       "neighbouring bands are almost identical, bands were picked as importance peaks: local maxima of the importance "
       "spectrum, at least 10 nm apart. The 17 highest peaks form the RF band set (the CART set of Das et al. had 17 bands); "
       "the 10 highest a smaller set. Checks: a second forest with another seed, R randomForest, permutation importance, and an "
       "ordinary class-weighted forest.")
imp = A["rf_importance"].set_index("name")
rf_rows = []
for k, b in enumerate(A["sets"][A["rf_name"]], 1):
    j = list(A["tier1"]["name"]).index(b)
    near, dist = min(((c, abs(float(c[2:]) - float(b[2:]))) for c in cart_set), key=lambda t: t[1])
    rf_rows.append({"Rank": k, "Band (nm)": float(b[2:]), "Importance × 1000": imp.at[b, "importance"] * 1000,
                    "Region": A["tier1"].at[j, "region"], "Mean healthy": A["tier1"].at[j, "mean0"],
                    "Mean diseased": A["tier1"].at[j, "mean1"], "Welch t": A["tier1"].at[j, "t"], "q (BH)": A["tier1"].at[j, "q"],
                    "Mixed model significant": "yes" if A["tier1"].at[j, "mm_sig"] else "no",
                    "Nearest CART band": float(near[2:]), "Distance (nm)": dist})
S.table(pd.DataFrame(rf_rows), title="Table A1.1  RF band set: the 17 highest importance peaks", widths={"D": 24})
rank = imp["importance"].rank(ascending=False)
cr_rows = []
for c in sorted(cart_set, key=lambda s: float(s[2:])):
    near, dist = min(((b, abs(float(c[2:]) - float(b[2:]))) for b in A["sets"][A["rf_name"]]), key=lambda t: t[1])
    j = list(A["tier1"]["name"]).index(c)
    cr_rows.append({"CART band (nm)": float(c[2:]), "Region": A["tier1"].at[j, "region"],
                    f"RF importance rank (of {A['n_sig']})": int(rank[c]), "RF importance × 1000": imp.at[c, "importance"] * 1000,
                    "Nearest RF-17 band": float(near[2:]), "Distance (nm)": dist})
S.table(pd.DataFrame(cr_rows), title="Table A1.2  The CART bands seen through the RF importance")
S.h2("Figure A1. RF importance spectrum with the RF peaks (orange) and the CART bands (blue triangles)")
S.image(FIG / "F12_A1_rf_importance.png", 820)
S.table(A["rf_checks"], title="Table A1.3  Checks of the RF importance", flag_col="Status", widths={"A": 50, "C": 60})

# =====================================================================================
# 11 A2
S = S_("11_A2_RF_Tier3_Tier4", "Part A2 · Tier 3 (J–M) and Tier 4 (QDA) on the CART and RF band sets",
       "Protocol of Das et al. (V1): bands chosen once on all leaves, QDA with equal priors in leaf-level 10-fold CV.")
sets_tab = pd.DataFrame([(k, len(v), bands_txt(v)) for k, v in A["sets"].items()], columns=["Band set", "Bands", "Wavelengths (nm)"])
S.table(sets_tab, title="Table A2.1  Band sets compared", widths={"C": 90})
t3 = A["tier3"][["Set", "Features", "alpha"]].copy()
t3.columns = ["Band set", "Bands", "α (Bhattacharyya)"]
f, l, L = S.table(t3, title="Table A2.2  Tier 3: separability of healthy and diseased leaves")
for j, h in enumerate(["J–M 0–2 (Excel)", "√J–M (Excel)", "≥ 1.90 ?", "≥ 1.94 ?"]):
    S.cell(f - 1, 4 + j, h, xl.F_HEAD, fill=xl.FILL_HEAD, border=True, align=xl.CENTER)
for r in range(f, l + 1):
    S.cell(r, 4, f"=2*(1-EXP(-C{r}))", fmt="0.0000", border=True)
    S.cell(r, 5, f"=SQRT(D{r})", fmt="0.0000", border=True)
    S.cell(r, 6, f'=IF(D{r}>=1.9,"separable","not separable")', border=True)
    S.cell(r, 7, f'=IF(D{r}>=1.94,"separable","not separable")', border=True)
t4 = A["tier4"][["Set", "Classifier", "HH", "HD", "DH", "DD", "OA", "BA", "PA_H", "PA_D", "UA_H", "UA_D", "kappa", "AUC"]].copy()
t4.columns = ["Band set", "Classifier", "H → H", "H → D", "D → H", "D → D", "OA", "BA", "PA healthy", "PA diseased", "UA healthy", "UA diseased", "κ", "AUC"]
S.table(t4, title="Table A2.3  Tier 4: 10-fold cross-validated classification (V1)",
        fmts={c: PCT for c in ["OA", "BA", "PA healthy", "PA diseased", "UA healthy", "UA diseased"]})
error_matrix(S, A["errmat"][A["cart_name"]], f"Error matrix — QDA (equal priors) on the {A['cart_name']} bands")
error_matrix(S, A["errmat"][A["rf_name"]], f"Error matrix — QDA (equal priors) on the {A['rf_name']} bands")
S.table(A["tier3_incremental"], title="Table A2.4  J–M and QDA balanced accuracy as bands are added",
        fmts={"CART QDA BA": PCT, "RF QDA BA": PCT})
S.image(FIG / "F09_tier3_tier4_incremental.png", 900)

# =====================================================================================
# 12 A3
S = S_("12_A3_CART_vs_RF", "Part A3 · CART vs RF as classifiers and as band selectors — honest validation",
       f"Same folds everywhere: V1 leaf-level 10-fold; V3 nested plot-grouped {v3_label(F)}; V4 30 % plot hold-out × {NV4}; V5 {v5_label(F)}.")
S.h2("What was done")
S.para("Four pipelines on identical folds. CART classifier: the rpart tree itself predicts. RF classifier: a Balanced RF "
       "(500 trees) on the Tier-1 bands predicts. QDA on CART bands: Tier 4 of the paper. QDA on RF bands: QDA on the 17 RF "
       "importance peaks. In V3, V4 and V5, Tier 1 and band selection are repeated inside every training set. Paired "
       f"Wilcoxon signed-rank tests compare pipelines over the {NSETS} V3 folds and the {NV4} V4 splits.")
va = A["validation"][["Pipeline", "Scheme", "repeats", "OA", "BA", "BA_SD", "kappa", "AUC", "PA_H", "PA_D", "UA_H", "UA_D"]]
S.table(va, title="Table A3.1  Validation results (mean over repeats; SD over repeats)",
        fmts={c: PCT for c in ["OA", "BA", "BA_SD", "PA_H", "PA_D", "UA_H", "UA_D"]})
S.image(FIG / "F13_A3_pipelines.png", 860)
S.table(A["wilcoxon"], title="Table A3.2  Paired comparisons (same folds; positive = first pipeline better)",
        fmts={"Mean difference": PCT, "SD": PCT})
for p in ["CART classifier", "RF classifier", "QDA on CART features", "QDA on RF features"]:
    error_matrix(S, A["pooled"][(p, "V4")], f"Error matrix — {p}, V4 hold-out plots ({NV4} splits pooled)")

# =====================================================================================
# 13 A4
S = S_("13_A4_Band_Stability", "Part A4 · Stability of the CART and RF band sets",
       f"Re-selection in the {NSETS} nested training sets of V3 (a band counts if a band within ± 10 nm is selected).")
S.table(A["stability"]["reselect_CART"], title=f"Table A4.1  The {A['cart_name']} bands", fmts={c: PCT for c in A["stability"]["reselect_CART"].columns if c.startswith("Re-")})
S.table(A["stability"]["reselect_RF"], title=f"Table A4.2  The {A['rf_name']} bands", fmts={c: PCT for c in A["stability"]["reselect_RF"].columns if c.startswith("Re-")})
S.table(A["stability"]["index"], title="Table A4.3  Stability index",
        note=f"Stability index (exact) = bands selected in all {NSETS} training sets ÷ all bands ever selected; the ±10 nm version counts a band as re-selected when a band within 10 nm is chosen; the mean pairwise Jaccard index (±10 nm) compares every pair of training sets.")
S.image(FIG / "F14_A4_stability.png", 900)
for meth in ("CART", "RF"):
    S.table(A["stability"]["freq"][meth].head(30), title=f"Table A4.{4 + (meth == 'RF')}  Bands most often selected by {meth} (exact band)",
            fmts={"Share of training sets": PCT})

# =====================================================================================
# 14 K0
k0 = K["k0"]
S = S_("14_K0_Denoising", "Step K0 · Spectral smoothing / denoising and detector joins",
       "Splice correction, Savitzky–Golay smoothing and derivatives; guard zones at the detector joins.")
S.h2("What was done")
S.para("1. Splice correction: the ASD records 350–1000 nm (VNIR), 1001–1800 nm (SWIR-1) and 1801–2500 nm (SWIR-2) with three "
       "detectors. For each join a straight line was fitted to the 10 nm before the join and the next detector was shifted so "
       "that it continues this line at join + 1 nm. 2. Smoothing: Savitzky–Golay filter (window 11, order 2); first derivative "
       "SG(15, 2), second derivative SG(21, 2). 3. Guard zones: bands next to the joins are excluded from feature ranking "
       "(reflectance 990–1010 and 1795–1835 nm; first derivative 988–1012 and 1793–1837 nm; second derivative 985–1015 and "
       "1790–1840 nm) because the filter windows straddle the join and a date-dependent start-up bump follows 1801 nm.")
S.table(pd.DataFrame(k0["offsets"]), title="Table K0.1  Splice offsets added to the following detector (reflectance)")
S.table(pd.DataFrame(k0["join_slopes"]), title="Table K0.2  Spectral slope just before and just after the 1800/1801 nm join (reflectance × 10⁵ per nm)",
        note="A healthy spectrum falls smoothly across the join; a slope that turns upward right after the join, or a class difference that changes with the date, marks the start-up artefact.")
S.table(pd.DataFrame(k0["noise"]), title="Table K0.3  Noise removed by the Savitzky–Golay filter")
S.table(pd.DataFrame([("Correlation of class-mean spectra, raw vs smoothed — healthy", k0["corr_neg"]),
                      ("Correlation of class-mean spectra, raw vs smoothed — diseased", k0["corr_pos"]),
                      ("Largest change of the all-leaf mean spectrum (reflectance)", k0["max_change"]),
                      ("Tier-1 significant bands: raw vs smoothed", f"{k0['sig_raw']} vs {k0['sig_smoothed']}")],
                     columns=["Check of the smoothing", "Value"]), title="Table K0.4  Smoothing preserves the signal", fmts={"Value": "0.000000"})
S.image(FIG / "F15_K0_detector_joins.png", 900)
checks_table(S, [c for c in cK if c[0] == "K0"])

# =====================================================================================
# 15 K1
S = S_("15_K1_Band_Reflectance", "Step K1 · Band reflectance — which single bands are disease-sensitive",
       "Splice-corrected, smoothed reflectance at every band 400–2500 nm (guard zones excluded from the ranking).")
S.h2("What was done")
S.para("For every band: Welch's t-test with BH-FDR, Hedges' g (standardised mean difference, bias-corrected), the area under "
       "the ROC curve (AUC, direction-free: 0.5 = no separation, 1 = perfect) and a threshold classifier (Youden's J; the "
       "threshold and its direction are chosen on training plots and applied to test plots: plot-grouped CV (V3) and hold-out (V4)).")
s1 = K["k1_summary"]
S.table(pd.DataFrame([("Bands tested", s1["tested"]), ("Significant (Welch + FDR), smoothed spectra", s1["sig_smoothed"]),
                      ("Significant, raw spectra (Tier 1)", s1["sig_raw"]), ("Bands with AUC ≥ 0.70 (outside guard zones)", s1["auc70"]),
                      ("Highest AUC (outside guard zones)", f"{s1['best_auc']:.3f} at {s1['best_wl']:.0f} nm")],
                     columns=["Summary", "Value"]), title="Table K1.1  Band-reflectance sensitivity")
kr = K["k1_regions"][["Wavelength (nm)", "Region", "AUC", "Direction", "Hedges g", "Welch t", "q (BH)", "Mean healthy", "Mean diseased",
                      "BA plot-grouped CV", "BA hold-out"]]
S.table(kr, title="Table K1.2  Most sensitive band in each spectral region (K1 features)", fmts={"BA plot-grouped CV": PCT, "BA hold-out": PCT})
S.table(K["k1_top15"].drop(columns=["Guard"]), title="Table K1.3  Fifteen most sensitive bands overall")
S.h2("Figure K1. AUC of every band: reflectance (K1, black) and first derivative (K4, blue)")
S.image(FIG / "F16_K1_K4_band_auc.png", 820)

# =====================================================================================
# 16 K2
def feat_table(df, extra=()):
    cols = ["name", "Mean healthy", "Mean diseased", "Welch t", "q (BH)", "Hedges g", "AUC", "Direction", "Threshold", "Rule",
            "BA all leaves", "BA plot-grouped CV", "BA hold-out"] + [c for c in ("Mixed-model q", "ICC plot") if c in df] + list(extra)
    t = df[cols].rename(columns={"name": "Feature"})
    return t


S = S_("16_K2_Continuum_Removal", "Step K2 · Continuum-removed absorption parameters",
       "Six absorption features; continuum = upper convex hull of each window (Clark & Roush 1984; Kokaly & Clark 1999).")
S.h2("What was done")
S.para("Within each window the reflectance was divided by its upper convex hull (continuum). From the continuum-removed curve: "
       "depth (1 − minimum), position (wavelength of the minimum), area (sum of 1 − CR, nm), width (full width at half depth, nm) "
       "and symmetry (share of the area left of the position). Statistics as in K1, plus a linear mixed model (q) and the plot "
       "intraclass correlation (ICC).")
S.table(pd.DataFrame(C["cr_windows"], columns=["Feature", "From (nm)", "To (nm)"]).assign(
    Absorber=["chlorophyll / carotenoid (≈ 495 nm)", "chlorophyll a/b, red well (≈ 675 nm)", "leaf water (≈ 970 nm)",
              "leaf water (≈ 1160–1200 nm)", "leaf water (≈ 1450 nm)", "leaf water (≈ 1940 nm)"]), title="Table K2.1  Absorption windows")
S.table(feat_table(K["k2"]), title="Table K2.2  All 30 continuum-removed parameters, ranked by AUC",
        fmts={"BA all leaves": PCT, "BA plot-grouped CV": PCT, "BA hold-out": PCT})
S.image(FIG / "F17_K2_continuum.png", 900)
checks_table(S, [c for c in cK if c[0] == "K2"])

# =====================================================================================
# 17 K3
S = S_("17_K3_Vegetation_Indices", "Step K3 · Disease-sensitive spectral (vegetation) indices — the 31 indices of the flowering-stage workbook",
       "Computed from splice-corrected, smoothed reflectance at the nearest 1 nm band.")
S.table(K["vi_table"], title="Table K3.1  The 31 indices", widths={"B": 40, "C": 60, "D": 28})
S.table(feat_table(K["k3"]), title="Table K3.2  Sensitivity of the 31 indices to white rust, ranked by AUC",
        fmts={"BA all leaves": PCT, "BA plot-grouped CV": PCT, "BA hold-out": PCT})
S.image(FIG / "F18_K3_indices.png", 620)
# Excel recomputation of nine indices for ten leaves
S.h2("Cross-check: indices recomputed by Excel formulas for ten leaves")
inp = K["vi_excel_input"]
inp = inp[["Leaf row"] + [c for c in inp.columns if c != "Leaf row"]]
f, l, L = S.table(inp, title="Table K3.3  Reflectance input (splice-corrected, smoothed)")
col = {c: L[c] for c in inp.columns}
formulas = {
    "NDVI780": "=({R780}-{R670})/({R780}+{R670})", "PRI": "=({R531}-{R570})/({R531}+{R570})",
    "TCARI": "=3*(({R700}-{R670})-0.2*({R700}-{R550})*({R700}/{R670}))",
    "MTVI2": "=1.5*(1.2*({R761}-{R581})-2.5*({R651}-{R581}))/SQRT((2*{R761}+1)^2-(6*{R761}-5*SQRT({R651}))-0.5)",
    "ARI": "=1/{R550}-1/{R700}", "WI/ND": "=({R900}/{R970})/(({R900}-{R680})/({R900}+{R680}))",
    "NWI2": "=({R970}-{R850})/({R970}+{R850})", "PSRI": "=({R680}-{R500})/{R750}", "RARSb": "={R675}/({R700}*{R650})"}
r0 = S.row
S.cell(r0, 1, "Table K3.4  Indices by Excel formula, Python value and difference", F_BOLD)
hdr = ["Leaf row"] + list(formulas) + [f"Python {k}" for k in formulas] + ["Max |Excel − Python|"]
for j, h in enumerate(hdr):
    S.cell(r0 + 1, 1 + j, h, xl.F_HEAD, fill=xl.FILL_HEAD, border=True, align=xl.CENTER)
py = K["vi_excel_python"]
for i in range(len(inp)):
    rr = r0 + 2 + i
    src = f + i
    S.cell(rr, 1, f"=A{src}", border=True)
    for j, (k, fm) in enumerate(formulas.items()):
        txt = fm
        for c in inp.columns:
            if c.startswith("R"):
                txt = txt.replace("{" + c + "}", f"{col[c]}{src}")
        S.cell(rr, 2 + j, txt, fmt="0.00000", border=True)
        S.cell(rr, 2 + len(formulas) + j, float(py[k].iloc[i]), fmt="0.00000", border=True)
    n = len(formulas)
    from openpyxl.utils import get_column_letter as gcl
    diffs = ",".join(f"ABS({gcl(2 + j)}{rr}-{gcl(2 + n + j)}{rr})" for j in range(n))
    S.cell(rr, 2 + 2 * n, f"=MAX({diffs})", fmt="0.0E+00", border=True)
S.row = r0 + 2 + len(inp)
S.cell(S.row, 1, "Largest difference over all ten leaves:", F_BOLD)
from openpyxl.utils import get_column_letter as gcl  # noqa: E402
S.cell(S.row, 5, f"=MAX({gcl(2 + 2 * len(formulas))}{r0 + 2}:{gcl(2 + 2 * len(formulas))}{r0 + 1 + len(inp)})", fmt="0.0E+00")
S.row += 2
checks_table(S, [c for c in cK if c[0] == "K3"])

# =====================================================================================
# 18 K3b
S = S_("18_K3b_Optimised_Index", "Step K3b (extra) · Optimised two-band index — exhaustive NDSI / DI search",
       "NDSI(i, j) = (Ri − Rj)/(Ri + Rj); DI(i, j) = Ri − Rj; all band pairs on a 5 nm grid (guard zones excluded), best pair refined at 1 nm.")
S.para(f"Every pair of bands on a 5 nm grid ({K['k3b_nbands']} bands, {K['k3b_npairs']:,} pairs) was scored by AUC. Choosing the "
       f"best of so many pairs on all leaves is optimistic, so the search was repeated inside {K3 + NV4} training sets (V3 repeat 1 and "
       f"the {NV4} V4 splits) and the chosen pair was tested on the held-out plots (honest estimate).")
S.table(K["k3b_top"].drop(columns=["Form"]), title="Table K3b.1  Ten best pairs of each form (5 nm grid, all leaves)")
S.table(feat_table(K["k3b_refined"]), title="Table K3b.2  The refined best indices (1 nm)",
        fmts={"BA all leaves": PCT, "BA plot-grouped CV": PCT, "BA hold-out": PCT},
        note="* the threshold is chosen on training plots, but the pair itself was found on all leaves (optimistic) — see Table K3b.3.")
h = K["k3b_honest"]
S.table(pd.DataFrame([(f, h[f]["BA CV"], h[f]["BA hold-out"], h[f]["SD"], h[f]["pairs"]) for f in ("NDSI", "DI")],
                     columns=["Index form", "BA plot-grouped CV (search inside folds)", "BA hold-out (search inside training plots)",
                              "Hold-out SD", f"Pairs chosen in the {K3 + NV4} training sets"]),
        title="Table K3b.3  Honest (nested) validation of the search",
        fmts={"BA plot-grouped CV (search inside folds)": PCT, "BA hold-out (search inside training plots)": PCT, "Hold-out SD": PCT},
        widths={"E": 80})
S.table(K["k3b_bydate"], title="Table K3b.4  Is the separation the same on every measurement date?")
ng = K["k3b_noguard"]
_in_guard = any(a_ <= w_ <= b_ for w_ in ng[:2] for a_, b_ in C["spectra"]["guard"]["R"])
S.para(f"Without the guard zones the best NDSI pair would be ({ng[0]:.0f}, {ng[1]:.0f}) nm with AUC {ng[2]:.3f}"
       + (" — a band inside a detector-join guard zone (start-up artefact)." if _in_guard else
          " — no band inside a guard zone, so the guard zones do not change the result here."), font=F_NOTE)
jf = K["k3b_jf"]
S.h2("Additional step K3b-JF · search repeated without pairs that straddle a detector join")
S.para((f"The best all-pair index {jf.get('allpair_best', '')} brackets a detector join, "
        f"so part of its separation can come from the splice step and the date-dependent start-up rise after 1801 nm (sheet 14). "
        if jf.get("allpair_straddles", True) else
        f"The best all-pair index {jf.get('allpair_best', '')} does not straddle a detector join; the join-free search is shown as a check. ")
       + f"Pairs with one band on each side of the 1000/1001 or 1800/1801 nm join and both within {jf['near_nm']} nm of it were "
       f"therefore excluded ({jf['excluded']:,} of {jf['npairs'] + jf['excluded']:,} pairs) and the search, refinement and nested "
       "validation were repeated." + (" The join-free result is the one recommended for interpretation." if jf.get("allpair_straddles", True) else ""))
S.table(jf["top"].drop(columns=["Form"]), title="Table K3b.5  Ten best join-free pairs of each form (5 nm grid, all leaves)")
S.table(feat_table(jf["refined"]), title="Table K3b.6  The refined best join-free indices (1 nm)",
        fmts={"BA all leaves": PCT, "BA plot-grouped CV": PCT, "BA hold-out": PCT})
hj = jf["honest"]
S.table(pd.DataFrame([(f, hj[f]["BA CV"], hj[f]["BA hold-out"], hj[f]["SD"], hj[f]["pairs"]) for f in ("NDSI", "DI")],
                     columns=["Index form", "BA plot-grouped CV (search inside folds)", "BA hold-out (search inside training plots)",
                              "Hold-out SD", f"Pairs chosen in the {K3 + NV4} training sets"]),
        title="Table K3b.7  Honest (nested) validation of the join-free search",
        fmts={"BA plot-grouped CV (search inside folds)": PCT, "BA hold-out (search inside training plots)": PCT, "Hold-out SD": PCT})
S.table(jf["bydate"], title="Table K3b.8  Join-free best pairs on each date")
S.image(FIG / "F19_K3b_ndsi_heatmap.png", 620)
checks_table(S, [c for c in cK if c[0] == "K3b"])

# =====================================================================================
# 19 K4
S = S_("19_K4_Derivatives", "Step K4 · First-derivative spectra and edge parameters",
       "Savitzky–Golay first derivative (window 15) and second derivative (window 21) of the splice-corrected spectra; 20 edge parameters (Gong et al. 2002; Guyot & Baret 1988).")
S.para("Band-wise: every first- and second-derivative band was tested as in K1 (bands in the guard zones are reported but "
       "excluded from the ranking). Edge parameters from the first derivative: amplitude (maximum/minimum D1), position and area "
       "(sum of D1) of the blue (490–530 nm), yellow (550–582 nm) and red (680–760 nm) edges; green peak and red well "
       "reflectance and position; the linear four-point REP; and six ratios.")
S.table(K["k4_summary"], title="Table K4.1  Band-wise derivative sensitivity")
S.table(K["k4_regions"], title="Table K4.2  Most sensitive derivative band per region (K4 features)")
S.table(feat_table(K["k4_edge"]), title="Table K4.3  Twenty derivative-based edge parameters, ranked by AUC",
        fmts={"BA all leaves": PCT, "BA plot-grouped CV": PCT, "BA hold-out": PCT})
S.image(FIG / "F20_K4_first_derivative.png", 900)
rp = K["rep_excel"]
f, l, L = S.table(rp[["Leaf row", "R670", "R700", "R740", "R780"]], title="Table K4.4  REP by linear four-point interpolation, recomputed by Excel formulas")
for j, h_ in enumerate(["REP (Excel)", "REP (Python)", "|difference|"]):
    S.cell(f - 1, 6 + j, h_, xl.F_HEAD, fill=xl.FILL_HEAD, border=True)
for i, r in enumerate(range(f, l + 1)):
    S.cell(r, 6, f"=700+40*((B{r}+E{r})/2-C{r})/(D{r}-C{r})", fmt="0.0000", border=True)
    S.cell(r, 7, float(rp["REP (Python)"].iloc[i]), fmt="0.0000", border=True)
    S.cell(r, 8, f"=ABS(F{r}-G{r})", fmt="0.0E+00", border=True)
checks_table(S, [c for c in cK if c[0] == "K4"])

# =====================================================================================
# 20 REIP
S = S_("20_K4b_REIP", "Step K4b · Red-edge inflection point (REIP) analysis",
       "The REIP is the wavelength of maximum slope on the red edge; chlorophyll loss shifts it to shorter wavelengths ('blue shift').")
S.para("Four established estimators were computed for every leaf: (a) maximum of the first derivative in 680–760 nm (Horler et al. "
       "1983); (b) linear four-point interpolation (Guyot & Baret 1988); (c) three-point Lagrangian interpolation around the "
       "derivative maximum (Dawson & Curran 1998); (d) linear extrapolation of the far-red (680–694 nm) and NIR (724–760 nm) "
       "first-derivative lines (Cho & Skidmore 2006). Each was tested as in K1, and by date, growth stage and variety.")
S.table(feat_table(K["reip"]).assign(Method_reference=K["reip"]["description"]).rename(columns={"Method_reference": "Method"}),
        title="Table K4b.1  REIP: healthy vs diseased by four methods", fmts={"BA all leaves": PCT, "BA plot-grouped CV": PCT, "BA hold-out": PCT},
        widths={"R": 60})
S.table(K["reip_corr"].reset_index().rename(columns={"index": "Method"}), title="Table K4b.2  Correlation between the REIP methods")
S.table(K["reip_strata"], title="Table K4b.3  REIP shift by date, growth stage and variety")
S.image(FIG / "F21_K4_REIP.png", 900)

# =====================================================================================
# 21 K6 correlation
S = S_("21_K6_Correlation", "Step K6 · Correlation of reflectance with disease status (correlation spectra)",
       "Severity scores were not used (instruction), so the point-biserial correlation with the binary status (healthy = 0, diseased = 1) is reported.")
S.para("r(λ) = Pearson correlation between the band value and the 0/1 disease status (point-biserial r), for the smoothed "
       "reflectance and the first derivative, on all leaves and separately on each date. r is a monotone function of "
       "Student's t, so the ranking matches Tier 1; the per-date curves show whether the sign is stable over time"
       + (" (this subset has a single date)." if len(PD) == 1 else "."))
S.table(K["corr_regions"], title="Table K6.1  Band of maximum |r| per region")
S.image(FIG / "F22_K6_correlation.png", 820)
checks_table(S, [c for c in cK if c[0] == "K6"])

# =====================================================================================
# 22 K5
S = S_("22_K5_CART_vs_RF", "Step K5 · CART vs Random Forest on each disease-sensitive feature group",
       f"Plot-grouped CV {v3_label(F)} ('CV' = V3) and {NV4} hold-outs of 30 % of plots (V4), identical folds.")
S.para("Each feature group was given to CART (rpart, equal priors, cp 0.01) and to a Balanced RF (500 trees) without any "
       "pre-screening; both were refitted in every training set. 'K-selected' = 15 non-redundant features (|r| < 0.9) with the "
       "highest AUC, chosen inside every training set from the combined pool, then classified by RF, CART and QDA.")
FE = K["features"]
gt = pd.DataFrame([(v[0], v[1].X.shape[1]) for k, v in FE["groups"].items()] + [("K-selected 15", 15)], columns=["Feature group", "Features"])
S.table(gt, title="Table K5.1  Feature groups")
k5 = K["k5"][["Feature group", "Features", "Classifier", "Scheme", "repeats", "BA", "BA_SD", "AUC", "kappa", "OA", "PA_H", "PA_D"]]
S.table(k5, title="Table K5.2  CART vs RF on each feature group (mean over repeats)", fmts={c: PCT for c in ["BA", "BA_SD", "OA", "PA_H", "PA_D"]},
        widths={"A": 46})
S.image(FIG / "F23_K5_cart_vs_rf.png", 900)
S.table(K["k5_wilcoxon"], title="Table K5.3  Paired comparison RF vs CART (same folds)", fmts={"RF − CART (BA)": PCT, "SD": PCT})
S.table(K["k5_type_share"], title="Table K5.4  Where the Random Forest finds its information (all leaves, combined pool)",
        note="The share depends on how many features a type has; 'importance per feature' corrects for that.")
S.table(K["k5_top20"], title="Table K5.5  Twenty most important features of the combined pool")
S.table(K["k5_selected"], title=f"Table K5.6  Features most often chosen for the compact set ({NSETS + NV4} training sets)", fmts={"Share": PCT})
S.table(pd.DataFrame([(FE["groups"][g][0], len(v), ", ".join(v)) for g, v in K["k5_trees"].items()],
                     columns=["Feature group", "Variables in the tree", "Variables used (order of first use)"]),
        title="Table K5.7  CART trees fitted to all leaves", widths={"C": 110})

# =====================================================================================
# 23-30 Part C (K1 pool) and 30a-30h Part C′ (raw pool)
def part_c_sheets(R, names, label, pool_txt, figs, all_sheet, reflect_desc):
    """Sheets C0-C6 for one combined-pool four-tier analysis."""
    S = S_(names[0], f"{label}0 · The combined feature pool ({pool_txt})",
           "The four-tier analysis of Das et al. applied to all extracted features instead of the reflectance bands alone.")
    S.para(f"Pool = {reflect_desc} + 30 K2 continuum-removed parameters + 31 K3 vegetation "
           f"indices + K4 first-derivative bands (guard zones excluded) + 20 K4 edge parameters = {R['pool_size_raw']} features. "
           f"Exact duplicates (|r| ≥ 0.99999) were removed because they make covariance matrices singular in Tiers 3–4, leaving "
           f"{R['n_features']} features. The same Tier 1–4 machinery, settings and folds as in Part A were then applied. RF peaks were "
           "picked within each spectral series (reflectance and derivative bands ≥ 10 nm apart); every scalar feature (CR, index, "
           "edge parameter) is its own candidate; a candidate nearly collinear (|r| ≥ 0.999) with an already chosen feature is skipped.")
    S.table(R["pool_dups"], title=f"Table {label}0.1  Duplicated features removed from the pool")
    S.table(R["composition"][["Feature type", "Features", "Tier-1 significant"]], title=f"Table {label}0.2  Pool composition")

    S = S_(names[1], f"{label}1 · Tier 1 on the combined pool — Welch t-test + BH-FDR",
           f"{R['n_sig']} of {R['n_features']} features significant (q < 0.05).")
    ct1 = R["tier1"]
    S.table(R["composition"][["Feature type", "Features", "Tier-1 significant"]], title=f"Table {label}1.1  Significant features by type")
    S.table(ct1.sort_values("auc", ascending=False).groupby("type", sort=True).head(5).sort_values(["type", "auc"], ascending=[True, False])[
        ["name", "type", "mean0", "mean1", "t", "q", "g", "auc", "mm_q"]].rename(
        columns={"name": "Feature", "type": "Type", "mean0": "Mean healthy", "mean1": "Mean diseased", "t": "Welch t",
                 "q": "q (BH)", "g": "Hedges g", "auc": "AUC", "mm_q": "Mixed-model q"}),
        title=f"Table {label}1.2  Five most sensitive features of each type (by AUC)")
    if all_sheet:
        S.para(f"All {R['n_features']} features with their statistics: sheet {all_sheet}.", font=F_NOTE)

    S = S_(names[2], f"{label}2 · Tier 2 on the combined pool — CART (rpart) and RF selection")
    csets = pd.DataFrame([(k, len(v), ", ".join(v)) for k, v in R["sets"].items()], columns=["Feature set", "Features", "Features (order of selection)"])
    S.table(csets, title=f"Table {label}2.1  Selected feature sets", widths={"C": 120})
    cimp = R["rf_importance"].set_index("name")
    rows = []
    for nm in R["sets"][R["cart_name"]] + [v for v in R["sets"][R["rf_name"]] if v not in R["sets"][R["cart_name"]]]:
        r = ct1.set_index("name").loc[nm]
        rows.append({"Feature": nm, "Type": r["type"], f"In {R['cart_name']}": "yes" if nm in R["sets"][R["cart_name"]] else "",
                     f"In {R['rf_name']}": "yes" if nm in R["sets"][R["rf_name"]] else "", "Mean healthy": r["mean0"],
                     "Mean diseased": r["mean1"], "Welch t": r["t"], "q (BH)": r["q"], "AUC": r["auc"],
                     "Mixed model significant": "yes" if r.get("mm_sig", False) else "no",
                     "RF importance rank": int(cimp["importance"].rank(ascending=False)[nm])})
    S.table(pd.DataFrame(rows), title=f"Table {label}2.2  The selected features")
    S.image(FIG / f"{figs['tree']}.png", 1000)
    fr = R["cart_full"]["frame"][["node", "depth", "var", "label", "n", "nH", "nD", "pred"]].copy()
    fr.columns = ["Node", "Depth", "Split variable", "Rule", "Leaves", "Healthy", "Diseased", "Predicted class"]
    S.table(fr, title=f"Table {label}2.3  Tree nodes (rpart frame)")
    S.table(R["cart_full"]["cptable"], title=f"Table {label}2.4  Complexity-parameter table (plot-grouped CV)")
    S.image(FIG / f"{figs['top']}.png", 700)
    S.image(FIG / f"{figs['comp']}.png", 800)
    if R["collinear_skipped"]:
        S.para("Candidates skipped by the collinearity guard (|r| ≥ 0.999 with a chosen feature): " + ", ".join(R["collinear_skipped"]), font=F_NOTE)

    S = S_(names[3], f"{label}3 · Tier 3 on the combined pool — J–M distance")
    t3 = R["tier3"][["Set", "Features", "alpha", "cond0", "cond1"]].copy()
    t3.columns = ["Feature set", "Features", "α (Bhattacharyya)", "Condition no. healthy", "Condition no. diseased"]
    f, l, L = S.table(t3, title=f"Table {label}3.1  Separability of healthy and diseased leaves",
                      note="Condition numbers of the class correlation matrices (scale-free).")
    for j, h_ in enumerate(["J–M 0–2 (Excel)", "√J–M (Excel)", "≥ 1.90 ?", "≥ 1.94 ?"]):
        S.cell(f - 1, 6 + j, h_, xl.F_HEAD, fill=xl.FILL_HEAD, border=True, align=xl.CENTER)
    for r in range(f, l + 1):
        S.cell(r, 6, f"=2*(1-EXP(-C{r}))", fmt="0.0000", border=True)
        S.cell(r, 7, f"=SQRT(F{r})", fmt="0.0000", border=True)
        S.cell(r, 8, f'=IF(F{r}>=1.9,"separable","not separable")', border=True)
        S.cell(r, 9, f'=IF(F{r}>=1.94,"separable","not separable")', border=True)
    S.table(R["tier3_incremental"], title=f"Table {label}3.2  J–M and QDA BA as features are added", fmts={"CART QDA BA": PCT, "RF QDA BA": PCT})
    S.image(FIG / f"{figs['inc']}.png", 900)

    S = S_(names[4], f"{label}4 · Tier 4 on the combined pool — QDA with equal priors (V1)")
    t4 = R["tier4"][["Set", "Classifier", "HH", "HD", "DH", "DD", "OA", "BA", "PA_H", "PA_D", "UA_H", "UA_D", "kappa", "AUC"]].copy()
    t4.columns = ["Feature set", "Classifier", "H → H", "H → D", "D → H", "D → D", "OA", "BA", "PA healthy", "PA diseased", "UA healthy", "UA diseased", "κ", "AUC"]
    S.table(t4, title=f"Table {label}4.1  10-fold cross-validated classification (V1)", fmts={c: PCT for c in ["OA", "BA", "PA healthy", "PA diseased", "UA healthy", "UA diseased"]})
    error_matrix(S, R["errmat"][R["cart_name"]], f"Error matrix — QDA (equal priors) on {R['cart_name']} ({pool_txt})")
    error_matrix(S, R["errmat"][R["rf_name"]], f"Error matrix — QDA (equal priors) on {R['rf_name']} ({pool_txt})")
    checks_table(S, [c for c in R["checks"] if c[0] in ("Tier 3", "Tier 4")])

    S = S_(names[5], f"{label}5 · Validation of the combined-pool four-tier pipeline",
           "Identical folds to Part A; Tier 1 + CART/RF selection + classifier refitted inside every training set (V3, V4, V5).")
    vc = R["validation"][["Pipeline", "Scheme", "repeats", "OA", "BA", "BA_SD", "kappa", "AUC", "PA_H", "PA_D", "UA_H", "UA_D"]]
    S.table(vc, title=f"Table {label}5.1  Validation results", fmts={c: PCT for c in ["OA", "BA", "BA_SD", "PA_H", "PA_D", "UA_H", "UA_D"]})
    S.image(FIG / f"{figs['pipes']}.png", 860)
    S.table(R["wilcoxon"], title=f"Table {label}5.2  Paired comparisons within this pool", fmts={"Mean difference": PCT, "SD": PCT})
    for p in ["QDA on CART features", "QDA on RF features"]:
        error_matrix(S, R["pooled"][(p, "V4")], f"Error matrix — {p}, V4 hold-out plots ({pool_txt})")

    S = S_(names[6], f"{label}6 · Stability of the features selected from the combined pool")
    S.table(R["stability"]["index"], title=f"Table {label}6.1  Stability index")
    S.table(R["fold_composition"], title=f"Table {label}6.2  Feature types selected in the {NSETS} nested training sets", fmts={"Share": PCT})
    S.table(R["stability"]["reselect_CART"], title=f"Table {label}6.3  Re-selection of the all-leaf CART features", fmts={c: PCT for c in R["stability"]["reselect_CART"].columns if c.startswith("Re-")})
    S.table(R["stability"]["reselect_RF"], title=f"Table {label}6.4  Re-selection of the all-leaf RF features", fmts={c: PCT for c in R["stability"]["reselect_RF"].columns if c.startswith("Re-")})
    S.image(FIG / f"{figs['stab']}.png", 900)
    return S


S = part_c_sheets(Cc, ["23_C0_Feature_Pool", "24_C1_Tier1", "25_C2_Tier2", "26_C3_Tier3_JM", "27_C4_Tier4_QDA", "28_C5_Validation", "29_C6_Stability"],
                  "C", "K1 + K2 + K3 + K4",
                  dict(tree="F26_C_cart_tree", top="F25_C_top_importance", comp="F24_C_composition", inc="F27_C_incremental",
                       pipes="F28_C_pipelines", stab="F28b_C_stability"),
                  "35_C_AllFeatures", "K1 smoothed, splice-corrected reflectance bands (guard zones excluded)")
S.table(X["guard_edge"], title="Table C6.5  Are selected derivative/reflectance bands crowding the detector guard zones? (artefact-leakage check)",
        fmts={"Share": PCT}, note="Guard zones: reflectance 990–1010 and 1795–1835 nm; first derivative 988–1012 and 1793–1837 nm. A feature within 5 nm outside a guard zone may still carry some join artefact.")
for meth in ("CART", "RF"):
    S.table(Cc["stability"]["freq"][meth].head(25), title=f"Features most often selected by {meth}", fmts={"Share of training sets": PCT})

S = S_("30_C7_Comparison", "Part C7 · Combined pool vs reflectance (Part A) vs K5 — paired tests on identical folds")
S.table(X["cross_wilcoxon"], title="Table C7.1  Paired Wilcoxon signed-rank tests (balanced accuracy, positive = first better)",
        fmts={"Mean difference (BA)": PCT, "SD": PCT}, widths={"A": 44, "B": 44})

if HAS_CR:
    # ---- Part C′: the same four tiers on raw reflectance + K2 + K3 + K4
    S = part_c_sheets(CR, ["30a_CR0_Pool_Raw", "30b_CR1_Tier1", "30c_CR2_Tier2", "30d_CR3_Tier3_JM", "30e_CR4_Tier4_QDA",
                           "30f_CR5_Validation", "30g_CR6_Stability"],
                      "C′", "raw reflectance + K2 + K3 + K4",
                      dict(tree="F32_CR_cart_tree", top="F33_CR_top_importance", comp="F34_CR_composition", inc="F35_CR_incremental",
                           pipes="F36_CR_pipelines", stab="F37_CR_stability"),
                      "35b_CR_AllFeatures", "all 2101 RAW reflectance bands, 400–2500 nm (exactly as in Part A; no smoothing, no splice correction, no guard zones)")
    for meth in ("CART", "RF"):
        S.table(CR["stability"]["freq"][meth].head(25), title=f"Features most often selected by {meth}", fmts={"Share of training sets": PCT})

    S = S_("30h_C_vs_CRaw_Compare", "Part C vs Part C′ · Which reflectance should enter the combined pool — K1 or raw?",
           "Same features K2 + K3 + K4 in both pools; only the reflectance part differs (K1 smoothed and splice-corrected vs raw). Identical folds.")
    S.para("Part C uses K1 reflectance (splice-corrected and Savitzky–Golay-smoothed, guard zones removed: 2039 bands); Part C′ uses the raw "
           "reflectance exactly as measured (2101 bands, the input of Part A and of Das et al. 2020). Everything else — K2, K3, K4, the "
           "four tiers, the settings, seeds and folds — is identical, so differences come only from the reflectance part of the pool.")
    S.table(XR["side_by_side"], title="Table CC.1  Side by side: pool, Tier 1, selected features, J–M, QDA (V1)", widths={"A": 46, "B": 60, "C": 60, "D": 30})
    S.table(XR["val_side"], title="Table CC.2  Balanced accuracy of the four pipelines in Part A, Part C′ and Part C",
            fmts={c: PCT for c in XR["val_side"].columns if c.startswith(("Part", "Δ"))})
    S.image(FIG / "F38_A_CR_C_comparison.png", 900)
    S.table(XR["wilcoxon"], title="Table CC.3  Paired Wilcoxon tests on identical folds (positive = first better)",
            fmts={"Mean difference (BA)": PCT, "SD": PCT}, widths={"A": 46, "B": 46})
    S.table(XR["comp_side"], title="Table CC.4  Feature types chosen by CART and RF in the 50 nested training sets", fmts={c: PCT for c in XR["comp_side"].columns if "share" in c.lower()})
    S.table(XR["overlap"], title="Table CC.5  Overlap of the selected features (a raw band R λ (raw) and a K1 band R λ count as the same wavelength)")
    S.table(XR["raw_bands"], title="Table CC.6  Where the raw bands chosen in the 50 nested training sets of Part C′ lie",
            note="Raw bands inside a guard zone or at the noisy spectral ends are bands that K1 removes or smooths in Part C. "
                 "All-leaf selections that are raw bands: " + (", ".join(XR["raw_bands_allleaf"]) or "none") + ".")
    lo = XR["lodo"].pivot_table(index=["Feature set", "Pipeline"], columns="Held-out date", values="BA mean", sort=False).reset_index()
    S.table(lo, title="Table CC.7  V5 — balanced accuracy on each held-out date (new plots)", fmts={c: PCT for c in lo.columns[2:]})
    if TXT.get("cr"):
        S.h2("Reading")
        for t in TXT["cr"]:
            S.para(t)


# =====================================================================================
# 31 V5, 32 strata
if V5BY == "Date":
    S = S_("31_V5_Transfer", "Additional step · V5 — transfer to an unseen date in unseen plots",
           f"Train on the other dates (training plots of a V4 split), test on the held-out date (test plots of the split); {NV4} splits × {len(F['V5_dates'])} dates.")
    S.para(f"Why: this subset pools {len(PD)} dates ({DATES_TXT}). A model useful in practice must work on a new date. V5 holds out a "
           "whole date and, at the same time, the plots of the V4 hold-out, so no leaf, plot or date of the test set is seen in training. "
           "BA is computed only for units whose test leaves contain both classes."
           + (f" {len(F['V5_dropped'])} split × date units were skipped because the held-out date had no leaves of the test plots or "
              "the training set had fewer than 5 leaves of a class." if F.get("V5_dropped") else ""))
else:
    S = S_("31_V5_Transfer", "Additional step · V5 — transfer to an unseen variety",
           "Single-date subset: a date cannot be held out, so V5 holds out one variety at a time (train on the other two varieties).")
    S.para("Why: within one date, the strongest test of generality is whether bands and classifiers chosen on two varieties work on "
           "the third. The held-out variety's plots are never seen in training, so the test is also on new plots.")
lo = X["lodo"]
S.table(lo, title=f"Table V5.1  Balanced accuracy by {v5_level(F).lower()}", fmts={"BA mean": PCT, "BA SD": PCT})
S.image(FIG / "F30_V5_leave_one_date_out.png", 820)

S = S_("32_Extra_Strata", "Additional step · Does the disease signal hold within each date, growth stage and variety?",
       "Descriptive checks inside each date, growth stage and variety present in this subset (strata with one class only are omitted).")
S.table(X["strata_auc"], title="Table S.1  AUC of the key features within strata")
S.table(X["strata_ba"], title="Table S.2  Classifier balanced accuracy within strata (V3 out-of-fold predictions)",
        fmts={"BA (V3, mean of 5 repeats)": PCT}, note=f"Out-of-fold predictions of V3 ({v3_label(F)}).")

# =====================================================================================
# 33 cross-checks
S = S_("33_CrossChecks", "Cross-checks at every step",
       "PASS = verified by an independent method or recomputation; FLAG = a limitation that is reported and handled.")
rows = [("Step 0", c[0], "data audit", str(c[1]), c[2]) for c in D["audit"]["checks"]]
rows += [(c[0], c[1], c[2], c[3], c[4]) for c in cA]
rows += [(c[0], c[1], c[2], c[3], c[4]) for c in cK]
rows += [(f"Part C · {c[0]}", c[1], c[2], c[3], c[4]) for c in cC]
rows += [(f"Part C′ · {c[0]}", c[1], c[2], c[3], c[4]) for c in (CR["checks"] if HAS_CR else [])]
fa = F
rows += [("Validation", "No plot in both training and test", "assertion on every V3, V4 and V5 fold", "passed for all folds", "PASS"),
         ("Validation", f"No {V5BY.lower()} in both training and test (V5)", "assertion on every V5 unit", f"passed ({len(F['V5'])} units)", "PASS"),
         ("Validation", "Every V3 and V4 test fold contains both classes", "assertion on every fold",
          f"passed ({NSETS} V3 folds, k = {K3}; {NV4} V4 splits)", "PASS"),
         ("Validation", "Identical folds in every step", "folds built once (sheet 36) and reused by Parts A, K and C", "same objects", "PASS"),
         ("K5 / Part A", "K5 full spectrum reproduces Part A", "RF and CART on the raw spectrum (K5, no Tier 1) vs Part A classifiers (Tier-1 bands)",
          f"RF {100 * K['k5'].query('Group==\"G0\" and Classifier==\"RF\" and Scheme==\"CV\"')['BA'].iloc[0]:.1f} % vs {100 * vrow(A, 'RF classifier', 'V3')['BA']:.1f} %; "
          f"CART {100 * K['k5'].query('Group==\"G0\" and Classifier==\"CART\" and Scheme==\"CV\"')['BA'].iloc[0]:.1f} % vs {100 * vrow(A, 'CART classifier', 'V3')['BA']:.1f} %",
          "PASS"),
         ("Error matrices", "Matrix sums equal the class totals", "Excel formulas (column totals) in every error matrix",
          f"V1 totals {nH} healthy / {nD} diseased", "PASS"),
         ("K5 artefact test", "Guard-zone bands inflate accuracy?", "K4 with vs without the guard-zone derivative bands (RF, CV)",
          f"{100 * K['k5'].query('Group==\"G6\" and Classifier==\"RF\" and Scheme==\"CV\"')['BA'].iloc[0]:.1f} % with vs "
          f"{100 * K['k5'].query('Group==\"G4\" and Classifier==\"RF\" and Scheme==\"CV\"')['BA'].iloc[0]:.1f} % without",
          "FLAG" if K['k5'].query('Group=="G6" and Classifier=="RF" and Scheme=="CV"')['BA'].iloc[0] - K['k5'].query('Group=="G4" and Classifier=="RF" and Scheme=="CV"')['BA'].iloc[0] > 0.01 else "PASS"),
         ("Part C", "Selected features next to the guard zones", "share of fold selections within 5 nm of a guard zone",
          "; ".join(f"{r['Selector']}: {100 * r['Share']:.1f} %" for _, r in X["guard_edge"].iterrows()),
          "FLAG" if X["guard_edge"]["Share"].max() > 0.05 else "PASS"),
         ("Numerics", "QDA fits failing on singular covariances", "count of failures in Parts A and C",
          f"{len(A['qda_failures'])} (Part A), {len(Cc['qda_failures'])} (Part C)", "PASS" if not (A['qda_failures'] or Cc['qda_failures']) else "FLAG")]
cc = pd.DataFrame(rows, columns=["Step", "What was checked", "Method", "Result", "Status"])
cc.insert(0, "#", range(1, len(cc) + 1))
f, l, L = S.table(cc, flag_col="Status", widths={"C": 40, "D": 46, "E": 60})
S.cell(S.row, 1, "PASS (Excel count)", F_BOLD)
S.cell(S.row, 3, f'=COUNTIF(F{f}:F{l},"PASS")')
S.cell(S.row + 1, 1, "FLAG (Excel count)", F_BOLD)
S.cell(S.row + 1, 3, f'=COUNTIF(F{f}:F{l},"FLAG")')
S.cell(S.row + 2, 1, "FAIL (Excel count)", F_BOLD)
S.cell(S.row + 2, 3, f'=COUNTIF(F{f}:F{l},"FAIL")')
S.row += 4
ALLCHECKS = cc

# =====================================================================================
# 34 K all bands, 35 C all features, 36 folds, 37 chart data
S = S_("34_K_AllBands", "K1, K4 and K6 — band-by-band results (400–2500 nm)",
       "Guard = inside a detector-splice guard zone (excluded from rankings).")
kb = pd.DataFrame({"Wavelength (nm)": wl, "Guard (R)": np.where(spec.guard["R"], "yes", ""),
                   "K1 mean healthy": K["k1_all"]["Mean healthy"], "K1 mean diseased": K["k1_all"]["Mean diseased"],
                   "K1 AUC": K["k1_all"]["AUC"], "K1 Hedges g": K["k1_all"]["Hedges g"], "K1 q (BH)": K["k1_all"]["q (BH)"],
                   "Guard (D1)": np.where(spec.guard["D1"], "yes", ""), "D1 mean healthy": K["d1_all"]["Mean healthy"],
                   "D1 mean diseased": K["d1_all"]["Mean diseased"], "D1 AUC": K["d1_all"]["AUC"], "D1 q (BH)": K["d1_all"]["q (BH)"],
                   "D2 AUC": K["d2_all"]["AUC"], "D2 q (BH)": K["d2_all"]["q (BH)"],
                   "r (R, status)": K["corr"]["R"]["r (point-biserial)"], "r (D1, status)": K["corr"]["D1"]["r (point-biserial)"]})
f, l, L = S.table(kb)
S.ws.freeze_panes = S.ws.cell(row=f, column=2)
S = S_("35_C_AllFeatures", "Part C — Tier 1 statistics of every feature of the combined pool")
ca = Cc["tier1"][["name", "type", "position", "mean0", "sd0", "mean1", "sd1", "t", "df", "p", "q", "g", "auc", "mm_q", "icc", "sig"]].copy()
ca.columns = ["Feature", "Type", "Wavelength (nm)", "Mean healthy", "SD healthy", "Mean diseased", "SD diseased", "Welch t", "df",
              "p", "q (BH)", "Hedges g", "AUC", "Mixed-model q", "ICC plot", "Significant"]
ca = ca.merge(Cc["rf_importance"][["name", "importance"]].rename(columns={"name": "Feature", "importance": "RF importance"}), on="Feature", how="left")
f, l, L = S.table(ca)
S.ws.freeze_panes = S.ws.cell(row=f, column=2)
if HAS_CR:
    S = S_("35b_CR_AllFeatures", "Part C′ — Tier 1 statistics of every feature of the raw-reflectance pool")
    ca = CR["tier1"][["name", "type", "position", "mean0", "sd0", "mean1", "sd1", "t", "df", "p", "q", "g", "auc", "mm_q", "icc", "sig"]].copy()
    ca.columns = ["Feature", "Type", "Wavelength (nm)", "Mean healthy", "SD healthy", "Mean diseased", "SD diseased", "Welch t", "df",
                  "p", "q (BH)", "Hedges g", "AUC", "Mixed-model q", "ICC plot", "Significant"]
    ca = ca.merge(CR["rf_importance"][["name", "importance"]].rename(columns={"name": "Feature", "importance": "RF importance"}), on="Feature", how="left")
    f, l, L = S.table(ca)
    S.ws.freeze_panes = S.ws.cell(row=f, column=2)
S = S_("36_Folds", "Fold assignments used by every step (identical everywhere)",
       "V1/V3: fold number of the leaf when it is in the test set; V4/V5: 1 = leaf is in the test plots of that split.")
S.table(F["table"])
S.ws.freeze_panes = "B6"

# chart data + native charts
S = S_("37_ChartData", "Chart data and native Excel charts", "Series behind the main figures.")
cd = pd.DataFrame({"Wavelength (nm)": wl, "Mean healthy": t1["mean0"], "Mean diseased": t1["mean1"],
                   "Significant (1 = yes)": t1["sig"].astype(int), "Balanced subsets significant (%)": t1["balanced_pct"],
                   "−log10 p Welch": -np.log10(np.clip(t1["p"], 1e-300, 1)), "−log10 q Welch": -np.log10(np.clip(t1["q"], 1e-300, 1)),
                   "−log10 q mixed model": -np.log10(np.clip(t1["mm_q"], 1e-300, 1)), "q = 0.05 line": -np.log10(0.05),
                   "RF importance × 1000": [A["rf_importance"].set_index("name")["importance"].get(f"WV{int(w)}", 0) * 1000 for w in wl],
                   "K1 AUC": K["k1_all"]["AUC"], "D1 AUC": K["d1_all"]["AUC"]})
f, l, L = S.table(cd)


def scatter(title, ycols, anchor, ytitle, xcol="A"):
    ch = ScatterChart()
    ch.title = title
    ch.style = 13
    ch.height, ch.width = 9, 18
    ch.x_axis.title = "Wavelength (nm)"
    ch.y_axis.title = ytitle
    ch.x_axis.scaling.min, ch.x_axis.scaling.max = 400, 2500
    xs = Reference(S.ws, min_col=1, min_row=f, max_row=l)
    colors = ["1BAF7A", "EB6834", "2A78D6", "52514E"]
    for k, (cname, lab) in enumerate(ycols):
        cidx = list(cd.columns).index(cname) + 1
        ys = Reference(S.ws, min_col=cidx, min_row=f - 1, max_row=l)
        s = Series(ys, xs, title_from_data=True)
        s.marker.symbol = "none"
        s.graphicalProperties.line.solidFill = colors[k % 4]
        s.graphicalProperties.line.width = 15000
        s.smooth = False
        ch.series.append(s)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    S.ws.add_chart(ch, anchor)


scatter("Figure 3. Mean reflectance (healthy vs diseased)", [("Mean healthy", ""), ("Mean diseased", "")], "N5", "Reflectance")
scatter("Figure 5b. −log10(p) against wavelength", [("−log10 p Welch", ""), ("−log10 q Welch", ""), ("−log10 q mixed model", ""), ("q = 0.05 line", "")], "N25", "−log10 p")
scatter("Figure 5. Balanced subsets with a significant difference (%)", [("Balanced subsets significant (%)", "")], "N45", "%")
scatter("Figure A1. RF importance spectrum", [("RF importance × 1000", "")], "N65", "MDI × 1000")
scatter("Figure K1. AUC of every band", [("K1 AUC", ""), ("D1 AUC", "")], "N85", "AUC")
# validation bar chart (native)
r0 = l + 3
S.cell(r0, 1, "Pipeline", F_BOLD)
schemes = ["V1", "V3", "V4", "V5"]
for j, s in enumerate(schemes):
    S.cell(r0, 2 + j, s, F_BOLD)
pipes = ["CART classifier", "RF classifier", "QDA on CART features", "QDA on RF features"]
for i, p in enumerate(pipes):
    S.cell(r0 + 1 + i, 1, p)
    for j, s in enumerate(schemes):
        S.cell(r0 + 1 + i, 2 + j, float(vrow(A, p, s)["BA"]), fmt=PCT)
bc = BarChart()
bc.type = "col"
bc.title = "Figure A3. Balanced accuracy by pipeline and validation scheme (Part A)"
bc.y_axis.title = "Balanced accuracy"
bc.height, bc.width = 9, 18
bc.add_data(Reference(S.ws, min_col=2, max_col=5, min_row=r0, max_row=r0 + 4), titles_from_data=True)
bc.set_categories(Reference(S.ws, min_col=1, min_row=r0 + 1, max_row=r0 + 4))
bc.y_axis.scaling.min = 0.4
bc.y_axis.scaling.max = 1.0
bc.x_axis.delete = False
bc.y_axis.delete = False
for k, s in enumerate(bc.series):
    s.graphicalProperties.solidFill = ["86B6EF", "3987E5", "1C5CAB", "0D366B"][k]
S.ws.add_chart(bc, "N105")

# =====================================================================================
# README sheet guide
R = readme
R.h2("Objective")
R.para("Discriminate healthy and white-rust (Albugo candida) diseased mustard leaves from leaf reflectance spectra"
       + (f" for the subset '{SUBSET}' (one of nine separate analyses: 4 dates, 2 growth stages, 3 varieties)" if IS_SUB
          else ", irrespective of date, growth stage and variety (all pooled)") + ", following the four-tier hierarchical method of Das et al. (2020) — Tier 1 "
       "statistical screening, Tier 2 CART band selection, Tier 3 Jeffries–Matusita separability, Tier 4 QDA classification — "
       "with corrections for class imbalance, multiple testing and pseudo-replication; then repeating Tier 2 with Random "
       "Forest, extracting disease-sensitive features (K1–K4) and applying the four tiers to the combined feature pool.")
R.h2("Data analysed")
R.para(f"asd__data_complete_labelled_110226 / 240226 / 030326 / 130326.xlsx — 79 exact duplicates removed from the four files"
       + (f", then the subset '{SUBSET}' taken: {len(y)} unique leaf spectra" if IS_SUB else f": {len(y)} unique leaf spectra")
       + f" ({nH} healthy, {nD} diseased; {NP0 + NP1} plots: {NP0} healthy control, {NP1} inoculated; date(s) {DATES_TXT}); 350–2500 nm at 1 nm; analysed "
       f"400–2500 nm ({len(wl)} bands). Plot-average columns were not used. Water-absorption bands were kept (contact probe). "
       "Leaf level and disease score were not used (healthy vs diseased only).")
if IS_SUB:
    R.h2("Scope of this workbook")
    R.para("Included: Steps 0–6 (data audit, pre-processing, four tiers with CART, validation V1–V5), Part A1–A4 (Random Forest "
           "instead of CART and comparison), Part K (K0–K6 disease-sensitive features, −log₁₀ p and correlation spectra) and Part C "
           "(four tiers with CART and RF on reflectance K1 + K2 + K3 + K4). Not included, as instructed: the feature-selection "
           "benchmark (filter / wrapper / embedded methods) and the Part C′ variant (raw + K2–K4), which were done for the "
           "combined data set only.")
R.h2("Software")
R.para("Python 3.13 (NumPy, SciPy 1.18, pandas, statsmodels 0.15, scikit-learn 1.9, imbalanced-learn 0.14, matplotlib, openpyxl); "
       "R 4.3.3 (rpart 4.1.23, MASS, randomForest 4.7). Code: spectral_pipeline/ (reusable for other datasets).")
R.h2("Sheet guide")
guide = pd.DataFrame([(n, t) for n, t in SHEETS if n != "00_README"], columns=["Sheet", "Contents"])
R.table(guide, widths={"A": 26, "B": 110})
R.h2("Conventions")
R.para("Reflectance in units of 0–1. Error matrices follow Das et al. Table 2: rows = classified, columns = reference. J–M on the "
       "0–2 scale. BA = balanced accuracy; OA = overall accuracy; PA/UA = producer's/user's accuracy; κ = Cohen's kappa; "
       "AUC = area under the ROC curve (direction-free for single features). Blue numbers in error matrices are counts "
       "(inputs); every total and accuracy beside them is an Excel formula. Columns marked (Excel) are formulas.")
wb.move_sheet("00_README", offset=-wb.index(wb["00_README"]))
path = OUT / f"{C['output']['stem']}_FourTier_CART_RF_K_Combined.xlsx"
wb.save(path)
ALLCHECKS.to_pickle(WORK / "allchecks.pkl")
print("saved", path, len(wb.sheetnames), "sheets")
