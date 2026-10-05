"""Author decision 2026-10-05: Chan C-pool concentrations derived from stock shares are scaled by TOC (not Walkley-Black SOC) -> updated_05."""
import sys
import openpyxl
p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 1).value == 87:
        ws.cell(r, 4).value = (str(ws.cell(r, 4).value) + " TOC (not SOC active pool) is the multiplier because the four Chan pools sum to total organic C "
                               "(NLC = TOC - C oxidised by 24 N H2SO4; Walkley-Black SOC ~ VLC + LC + LLC).")
        ws.cell(r, 6).value = "IN FORCE (author confirmed 2026-10-05: TOC basis)"
n = 0
for sh in ("VLC(Cfrac1)", "LC(Cfrac2)", "LLC(Cfrac3)", "NLC(Cfrac4)"):
    s = wb[sh]
    h = {c.value: c.column for c in s[1]}
    for r in range(2, s.max_row + 1):
        if s.cell(r, 2).value == "23 (companion 4)":
            c = s.cell(r, h["Notes/Doubts"])
            c.value = ("TOC BASIS CONFIRMED by author 2026-10-05 (rule 87): the pools sum to total organic C, so stock shares are scaled by TOC, not by "
                       "Walkley-Black SOC (active pool). " + str(c.value))
            n += 1
rd = wb["README"]
r = rd.max_row + 1
rd.cell(r, 1, "Decision on batch 114-116 (2026-10-05)").font = openpyxl.styles.Font(name="Arial", size=10, bold=True)
c = rd.cell(r, 2, "updated 05: author confirmed rule 87 - carbon-pool concentrations derived from printed stocks are scaled by TOC (pools sum to TOC). "
                  "No data values changed. Rules 85, 86, 88, 89 still pending confirmation.")
c.font = openpyxl.styles.Font(name="Arial", size=10)
c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("pool rows annotated:", n)
