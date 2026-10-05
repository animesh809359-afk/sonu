"""Batch 12 defaults (updated 12): new rules 103-106 (awaiting author confirmation)."""
import sys
from copy import copy

import openpyxl

p_in, p_out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p_in)
ws = wb["RULES"]
style = [copy(c._style) for c in ws[ws.max_row]]
n0 = max(v for v in (ws.cell(i, 1).value for i in range(1, ws.max_row + 1)) if isinstance(v, int))
ROWS = [
    ("Treatment codes",
     "In one paper, puddled rice + ZT wheat -> pZT / pCA (rule 68), while conventionally DRY-TILLED UNPUDDLED rice (DSR) + ZT wheat -> ZT / CA (rule 25, flagged); the two are "
     "never pooled.",
     "174 Gathala 2017: T1 pZT, T2 pCA, T3 / T4 (CT-DSR) ZT / CA row a, T5 / T6 (ZT-DSR) ZT / CA row b.", "Default 2026-10-05", "IN FORCE - awaiting author confirmation"),
    ("Workflow & data management",
     "A treatment that two rules code differently (e.g. conventional dry-tilled, non-puddled WATER-SEEDED rice + surface-seeded ZT wheat: ZT under rule 25 or pZT under rule 68) "
     "is NOT entered; its values are written into the Notes of the study's rows and the treatment is logged as PENDING until the author decides.",
     "176 / 176 (companion) Devkota 2015 WSRF-SSW-FI and WSRF-SSW-AWD.", "Default 2026-10-05 (rule 3)", "PENDING - author decision"),
    ("Parameter-specific",
     "B:C not defined in the paper: treated as gross / cost and converted (printed - 1, rule 57) when the printed numbers imply it (net/cost would give a total cost below the "
     "itemised costs); flagged on the row.",
     "175 Hossain 2022 (rice TA: gross 1133 US$/ha, BCR 1.07 -> cost 1059 as gross/cost vs 547 as net/cost < tillage + weeding 455 + fertiliser).", "Default 2026-10-05",
     "IN FORCE - awaiting author confirmation"),
    ("Parameter-specific",
     "Crop-model papers: MEASURED values printed only as points of a measured-vs-simulated plot are entered as measured data (read from the measured axis and checked against the "
     "printed measured means); the paper's simulated long-term means are entered as separate MODEL OUTPUT rows (rule 55); simulation-only scenarios never tested in the field "
     "are not entered.",
     "176 (companion) Devkota 2015 AFM: Figs 3A / 4 measured yields 2008-2009 (means reproduce Tables 4-5 within 3 kg/ha); Tables 6-7 DSSAT 1971-2010 yields, N uptake, "
     "N leached; scenarios 9-11 excluded.", "Default 2026-10-05", "IN FORCE - awaiting author confirmation"),
]
r = ws.max_row + 1
for i, row in enumerate(ROWS, 1):
    for j, v in enumerate((n0 + i,) + row, 1):
        ws.cell(r, j, v)._style = copy(style[j - 1])
    r += 1
wb.calculation.fullCalcOnLoad = True
wb.save(p_out)
print("new rules", n0 + 1, "-", n0 + len(ROWS))
