"""Plot-wise estimates of biophysical and biochemical parameters, 24/02/2026 (DI scored 26/02/2026).

Usage:
    python estimate_plots.py <asd_240226.xlsx> <plot_DI.xlsx> <treatment_detail.xlsx> \
        <model_workbook.xlsx> <prospect_d_spectra.txt> <output.xlsx>

Each plot's estimate combines
  1. the expected healthy value of its treatment on 24/02 (Simulation sheet of the model workbook),
  2. its measured leaf scores, converted to pustule leaf area and an effective severity s_eff,
  3. its plot-average ASD spectrum (chlorophyll from the MTCI red-edge index via PROSPECT-D).
The combination step is written as live formulas so the weights and couplings can be edited.
"""
import sys

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from scipy.optimize import brentq, minimize_scalar
from scipy.signal import savgol_filter
from scipy.special import expi

ASD, DIF, LAYOUT, MODEL, SPEC, OUT = sys.argv[1:7]
DATE = "24/02/2026"
AREA = {0: 0.0, 1: 2.5, 3: 7.5, 5: 18.0, 7: 38.0, 9: 75.0}   # leaf area (%) at each Bisht grade (class mid-points)
CLASSES = [(0, "Healthy"), (5, "Very low"), (10, "Low"), (25, "Moderate"), (50, "High"), (1e9, "Very high")]
LMA = {"V": 0.0038, "F": 0.0042, "P": 0.0046}   # g cm-2, as in the simulated dataset
N_LEAF = {"V": 1.50, "F": 1.60, "P": 1.75}
BROWN0 = {"V": 0.0, "F": 0.02, "P": 0.08}
WATER_SAT = 6.0
TOL = {"V1": 1.0, "V2": 1.1, "V3": 0.85}

# ---------------------------------------------------------------- inputs
ws = openpyxl.load_workbook(ASD, read_only=True).worksheets[0]
rows = list(ws.iter_rows(values_only=True))
hdr, data = rows[:7], rows[7:]
WL = np.array([r[0] for r in data], float)
idx = [j for j in range(1, len(hdr[0])) if isinstance(hdr[0][j], str) and hdr[0][j].endswith(".sco")]
X = np.array([[r[j] for j in idx] for r in data], float).T
spec_plot = [hdr[1][j] for j in idx]
spec_name = [hdr[0][j] for j in idx]

dws = openpyxl.load_workbook(DIF, data_only=True)["240226"]
scores, cur = [], None
for r in dws.iter_rows(min_row=2, values_only=True):
    if r[2] is not None:
        cur = r[2]
    scores.append((cur, r[4], r[5], r[8]))
assert len(scores) == len(idx) == 540
for k, (p, _, sid, _) in enumerate(scores):   # sample k <-> spectrum k
    assert p == spec_plot[k] and int(spec_name[k][7:12]) == sid, k

lay = {r[0]: dict(sow=r[1], var=r[2], rep="Control" if r[3] not in ("R1", "R2", "R3") else r[3])
       for r in openpyxl.load_workbook(LAYOUT).active.iter_rows(min_row=2, values_only=True)}

mws = openpyxl.load_workbook(MODEL, data_only=True, read_only=True)["Simulation"]
mit = mws.iter_rows(min_row=2, values_only=True)
mh = next(mit)
H = {}
for r in mit:
    d = dict(zip(mh, r))
    if d["Date code"] == DATE and d["Severity class"] == "Healthy":
        H[d["Treatment"]] = d
assert len(H) == 12

# ---------------------------------------------------------------- PROSPECT-D (Feret et al. 2017)
spec = np.loadtxt(SPEC, comments="#")
_, NR, KAB, KCAR, KANT, KBROWN, KW, KM = spec.T


def tav(alpha, nr):
    n2 = nr * nr
    npx, nm = n2 + 1, n2 - 1
    a = (nr + 1) ** 2 / 2.0
    k = -(n2 - 1) ** 2 / 4.0
    sa = np.sin(np.deg2rad(alpha))
    b1 = np.sqrt((sa * sa - npx / 2) ** 2 + k) if alpha != 90 else 0.0
    b = b1 - (sa * sa - npx / 2)
    ts = (k ** 2 / (6 * b ** 3) + k / b - b / 2) - (k ** 2 / (6 * a ** 3) + k / a - a / 2)
    tp = (-2 * n2 * (b - a) / npx ** 2 - 2 * n2 * npx * np.log(b / a) / nm ** 2 + n2 * (1 / b - 1 / a) / 2
          + 16 * n2 ** 2 * (n2 ** 2 + 1) * np.log((2 * npx * b - nm ** 2) / (2 * npx * a - nm ** 2)) / (npx ** 3 * nm ** 2)
          + 16 * n2 ** 3 * (1 / (2 * npx * b - nm ** 2) - 1 / (2 * npx * a - nm ** 2)) / npx ** 3)
    return (ts + tp) / (2 * sa ** 2)


T40, T90 = tav(40.0, NR), tav(90.0, NR)


def prospect(N, cab, car, ant, brown, cw, cm):
    k = (cab * KAB + car * KCAR + ant * KANT + brown * KBROWN + cw * KW + cm * KM) / N
    tau = np.where(k > 0, (1 - k) * np.exp(-k) + k ** 2 * (-expi(-np.maximum(k, 1e-12))), 1.0)
    talf, ralf, t12, r12 = T40, 1 - T40, T90, 1 - T90
    t21 = t12 / NR ** 2
    r21 = 1 - t21
    den = 1 - r21 ** 2 * tau ** 2
    Ta = talf * tau * t21 / den
    Ra = ralf + r21 * tau * Ta
    t = t12 * tau * t21 / den
    r = r12 + r21 * tau * t
    D = np.sqrt((1 + r + t) * (1 + r - t) * (1 - r + t) * (1 - r - t))
    a = (1 + r * r - t * t + D) / (2 * r)
    b = (1 - r * r + t * t + D) / (2 * t)
    bN1 = b ** (N - 1)
    d2 = a * a * bN1 * bN1 - 1
    Rs, Ts = a * (bN1 * bN1 - 1) / d2, bN1 * (a * a - 1) / d2
    return Ra + Ta * Rs * t / (1 - Rs * r)


def mtci(S, w0):
    R = lambda w: S[..., int(w - w0)]
    return (R(754) - R(709)) / (R(709) - R(681))


def cab_from_mtci(m, st, s_eff, cw, cm):
    """Chlorophyll a+b (ug cm-2) whose PROSPECT-D spectrum has this MTCI (affine-invariant index)."""
    f = lambda cab: mtci(savgol_filter(prospect(N_LEAF[st] * (1 - 0.15 * s_eff), cab, cab / 5.5, 0.4,
                                                 BROWN0[st] + 0.6 * s_eff, cw, cm), 21, 2), 400) - m
    return brentq(f, 5, 120)


# ---------------------------------------------------------------- plot features
Xs = savgol_filter(X, 21, 2, axis=1)
Rw = lambda S, w: S[..., int(w - WL[0])]
P = {}
for p in range(1, 49):
    k = np.array(spec_plot) == p
    S = Xs[k].mean(0)
    sc = [s[3] for s in scores if s[0] == p]
    lv = [s[1] for s in scores if s[0] == p]
    t = lay[p]["sow"] + lay[p]["var"]
    di = sum(sc) / (9 * len(sc)) * 100
    d1 = np.gradient(S)
    P[p] = dict(
        plot=p, t=t, sow=lay[p]["sow"], var=lay[p]["var"], rep=lay[p]["rep"], n=int(k.sum()),
        leaves=f'{lv.count("L3")} L3 / {lv.count("L5")} L5 / {lv.count("LL")} LL',
        mean_grade=float(np.mean(sc)), di=di, cls=next(c for lim, c in CLASSES if di <= lim),
        area=float(np.mean([AREA[int(v)] for v in sc])),
        stage=H[t]["Stage"], leaf=H[t]["Leaf sampled"],
        spectrum=X[k].mean(0),
        MTCI=float(mtci(S, WL[0])),
        NDVI=float((Rw(S, 800) - Rw(S, 670)) / (Rw(S, 800) + Rw(S, 670))),
        CIre=float(Rw(S, 780) / Rw(S, 710) - 1),
        REIP=float(680 + np.argmax(d1[int(680 - WL[0]):int(760 - WL[0])])),
        ARI=float(1 / Rw(S, 550) - 1 / Rw(S, 700)),
        NDWI=float((Rw(S, 860) - Rw(S, 1240)) / (Rw(S, 860) + Rw(S, 1240))),
        WBD1450=float(1 - Rw(S, 1450) / ((Rw(S, 1250) + Rw(S, 1650)) / 2)),
    )

# ---------------------------------------------------------------- severity calibration from the spectra
# Chlorophyll of each inoculated plot relative to the healthy plot of the same treatment (MTCI elasticity
# 1.13 from PROSPECT-D at 35-45 ug cm-2) is matched to the model's chlorophyll response to find kappa in
# s_eff = kappa x pustule leaf-area fraction.
ELAST = 1.13
hp = {P[q]["t"]: q for q in range(1, 13)}
dz = [p for p in range(13, 49)]
chl_rel = np.array([(P[p]["MTCI"] / P[hp[P[p]["t"]]]["MTCI"]) ** (1 / ELAST) for p in dz])
A = np.array([P[p]["area"] / 100 for p in dz])
tol = np.array([TOL[P[p]["var"]] for p in dz])


def model_tot(s, t):
    ra = -0.65 * (1 - np.exp(-3.2 * s))
    rb = -0.55 * (1 - np.exp(-3.0 * s))
    return (1.25 * (1 + ra * t) + 0.45 * (1 + rb * t)) / 1.70


KAPPA = minimize_scalar(lambda kp: np.sum((np.log(model_tot(kp * A, tol)) - np.log(chl_rel)) ** 2),
                        bounds=(0.05, 3), method="bounded").x
KAPPA = round(float(KAPPA), 2)
for i, p in enumerate(dz):
    P[p]["chl_rel"] = float(chl_rel[i])
for p in range(1, 13):
    P[p]["chl_rel"] = 1.0
for p in P:
    P[p]["s_eff"] = KAPPA * P[p]["area"] / 100


# ---------------------------------------------------------------- model priors (same equations as the model workbook)
def R(code, s):
    if code == "RWC":
        L = lambda x: 1 / (1 + np.exp(-(x - 0.35) / 0.12))
        return -0.30 * (L(s) - L(0)) / (1 - L(0))
    if code == "MSI":
        L = lambda x: 1 / (1 + np.exp(-(x - 0.28) / 0.10))
        return -0.70 * (L(s) - L(0)) / (1 - L(0))
    if code == "Car":
        return -0.40 * (1 - np.exp(-3 * s)) + 0.10 * (s / 0.10) * np.exp(1 - s / 0.10)
    sign, rmax, k = {"Pn": (-1, 0.75, 4.4), "gs": (-1, 0.65, 3.0), "Ci": (1, 0.30, 4.0), "LAI": (-1, 0.45, 2.5),
                     "ChlA": (-1, 0.65, 3.2), "ChlB": (-1, 0.55, 3.0), "Anth": (1, 0.60, 3.0),
                     "Prot": (-1, 0.15, 3.0)}[code]
    return sign * rmax * (1 - np.exp(-k * s))


CODES = ["RWC", "MSI", "Pn", "gs", "Ci", "LAI", "ChlA", "ChlB", "Car", "Anth", "Prot"]
for p, d in P.items():
    h, s, tl = H[d["t"]], d["s_eff"], TOL[d["var"]]
    pri = {c: h[f"{c} healthy"] * (1 + R(c, s) * tl) for c in CODES}
    pri["E"] = h["E healthy"] * pri["gs"] / h["gs healthy"] * (1 + 0.25 * s)
    d["prior"] = pri
    d["E_H"], d["gs_H"] = h["E healthy"], h["gs healthy"]
    # spectral chlorophyll in mg g-1 FW: MTCI -> Cab (ug cm-2) -> divide by fresh mass per area
    st = d["stage"]
    cw = WATER_SAT * LMA[st] * pri["RWC"] / 100
    cm = LMA[st] * (1 - 0.10 * s)
    d["Cab_area"] = cab_from_mtci(d["MTCI"], st, s, cw, cm)
    d["FW_area"] = cw + cm
    d["chl_spec"] = d["Cab_area"] / (d["FW_area"] * 1000)

# ---------------------------------------------------------------- workbook
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
FW = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FT = Font(name="Arial", size=12, bold=True, color="1F4D2B")
FR = Font(name="Arial", size=10, bold=True, color="C00000")
FIN = Font(name="Arial", size=10, color="0000FF")
HDR = PatternFill("solid", fgColor="1F4D2B")
YEL = PatternFill("solid", fgColor="FFFF00")
CF = {"Healthy": "E2EFDA", "Very low": "F2F7E8", "Low": "FFF8E1", "Moderate": "FFE9C7", "High": "FCD5C0", "Very high": "F4B6B6"}
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")


def put(ws, r, c, v, font=F, fill=None, fmt=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.border = font, BOX
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill) if isinstance(fill, str) else fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    return cell


def header(ws, r, cols, widths=None):
    for j, h in enumerate(cols, start=1):
        put(ws, r, j, h, FW, HDR, align=WRAP)
    ws.row_dimensions[r].height = 42
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[CL(j)].width = w


wb = openpyxl.Workbook()
rd = wb.active
rd.title = "README"
lines = [
    ("PLOT-WISE ESTIMATES OF BIOPHYSICAL AND BIOCHEMICAL PARAMETERS - 24/02/2026 (DI scored 26/02/2026)", FT),
    ("ESTIMATED VALUES, NOT LABORATORY MEASUREMENTS. Use them as expected plot values, for checking your lab data, as priors, and to test "
     "spectra-to-trait models. Report them as estimates; replace with your measured values where you have them.", FR),
    ("", F),
    ("HOW EACH PLOT VALUE IS BUILT", FB),
    ("1. Healthy baseline: the expected healthy value of the plot's treatment (sowing date x variety) on 24/02/2026 from the model workbook "
     "(stage, sampled leaf, variety, heat and VPD already included).", F),
    ("2. Severity: your leaf scores (0-9, Bisht et al. 2015) for the spectra of that plot -> pustule leaf area (grade mid-points 0, 2.5, 7.5, 18, "
     "38, 75%) -> plot mean area A -> effective severity s_eff = kappa x A.", F),
    (f"3. kappa = {KAPPA} was fitted so that the model's chlorophyll response matches the chlorophyll drop seen in your plot-average spectra "
     "(MTCI of each inoculated plot / MTCI of the healthy plot of the same treatment). Using DI directly (s = DI) would overstate the decline "
     "about 5-fold, because DI is a score index, not tissue area.", F),
    ("4. Prior: model response curves R(s_eff) x variety tolerance applied to the healthy baseline (Inputs sheet, 'prior' columns).", F),
    ("5. Spectral chlorophyll: plot-average ASD spectrum -> MTCI (insensitive to offset and gain) -> chlorophyll a+b per leaf area by PROSPECT-D "
     "-> divided by fresh mass per area -> mg g-1 FW.", F),
    ("6. Final (Estimates sheet, live formulas): total chlorophyll = weighted geometric mean of prior and spectral values; the ratio "
     "delta = final / prior chlorophyll then adjusts the other traits through the exponents in Settings (Pn, gs, MSI, protein follow the "
     "chlorophyll status; anthocyanin moves opposite). Chl a, Chl b and carotenoids keep the model's a/b and Car/Chl ratios.", F),
    ("7. Ranges: estimate x (1 -/+ 1.645 x CV) with the plant-to-plant and assay CVs of the model (about 90% interval).", F),
    ("", F),
    ("NOTES", FB),
    ("- The plot-average spectrum mixes the 3rd, 5th and lower leaves you scanned (L3/L5/LL); estimates refer to that plot-average leaf, on the "
     "stage-specific basis of the model (3rd leaf at vegetative, 5th leaf at flowering).", F),
    ("- Water indices showed no reliable difference between healthy and diseased plots on this date, so RWC follows the prior (small decline).", F),
    ("- LAI is a canopy trait that leaf spectra cannot see; it follows the prior (defoliation with severity).", F),
    ("- Each treatment has one healthy plot (6-9 spectra), so spectral chlorophyll of the healthy reference carries noise; the 50% weight on "
     "the prior limits its influence.", F),
    ("- Edit the yellow cells in Settings to change weights or couplings; Estimates and Treatment_Summary recalculate.", F),
    ("", F),
    ("SHEETS: Settings | Estimates (final values) | Ranges (low-high) | Inputs (severity, priors, spectral chlorophyll) | Spectral_Features | "
     "Treatment_Summary | Plot_Mean_Spectra (350-2500 nm).", F),
]
for i, (t, f) in enumerate(lines, start=1):
    c = rd.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
    rd.row_dimensions[i].height = 15 * max(1, len(t) // 150 + 1)
rd.column_dimensions["A"].width = 160

# Settings
se = wb.create_sheet("Settings")
se.cell(row=1, column=1, value="Combination settings (yellow = editable)").font = FT
SET = [
    ("w_spec", "Weight of spectral chlorophyll in total chlorophyll (0-1)", 0.5),
    ("e_Pn", "Exponent of chlorophyll ratio on Pn", 0.5),
    ("e_gs", "Exponent of chlorophyll ratio on gs (E follows gs)", 0.4),
    ("e_Ci", "Exponent of chlorophyll ratio on Ci", -0.15),
    ("e_MSI", "Exponent of chlorophyll ratio on MSI", 0.2),
    ("e_Prot", "Exponent of chlorophyll ratio on protein", 0.3),
    ("e_Anth", "Exponent of chlorophyll ratio on anthocyanin", -0.5),
    ("kappa", "s_eff = kappa x pustule leaf area (fitted from spectra; information only)", KAPPA),
]
header(se, 2, ["Name", "Meaning", "Value"], [10, 70, 10])
SREF = {}
for i, (n, mtxt, v) in enumerate(SET, start=3):
    put(se, i, 1, n, FB)
    put(se, i, 2, mtxt)
    put(se, i, 3, v, FIN, YEL, "0.00")
    SREF[n] = f"Settings!$C${i}"
CV = {"RWC": 0.04, "MSI": 0.07, "Pn": 0.12, "gs": 0.18, "E": 0.18, "Ci": 0.06, "LAI": 0.12, "ChlA": 0.10,
      "ChlB": 0.12, "Car": 0.12, "Anth": 0.30, "Prot": 0.12, "TotChl": 0.10, "WUE": 0.15}
r0 = len(SET) + 5
se.cell(row=r0 - 1, column=1, value="Coefficients of variation for ranges (model values)").font = FB
header(se, r0, ["Code", "Parameter", "CV"])
for i, (c, v) in enumerate(CV.items(), start=r0 + 1):
    put(se, i, 1, c, FB)
    put(se, i, 2, c)
    put(se, i, 3, v, FIN, YEL, "0.00")
    SREF["CV_" + c] = f"Settings!$C${i}"

ORDER = sorted(P, key=lambda p: (0 if P[p]["rep"] == "Control" else 1, P[p]["sow"], P[p]["var"], P[p]["rep"]))
ID = ["Plot", "Treatment", "Sowing", "Variety", "Replication", "Growth stage", "Plot DI (%)", "Severity class"]
IDW = [6, 10, 7, 7, 11, 12, 9, 11]


def id_cells(ws, r, d):
    vals = [d["plot"], d["t"], d["sow"], d["var"], d["rep"], {"V": "Vegetative", "F": "Flowering", "P": "Pod formation"}[d["stage"]],
            round(d["di"], 1), d["cls"]]
    for j, v in enumerate(vals, start=1):
        put(ws, r, j, v, FB if j in (1, 2) else F, CF[d["cls"]] if j == 8 else None, "0.0" if j == 7 else None)


# Inputs
ip = wb.create_sheet("Inputs")
ICOL = ID + ["Spectra (n)", "Leaves scanned", "Mean leaf grade", "Pustule leaf area A (%)", "s_eff", "MTCI", "Chl rel. to own healthy plot",
             "Cab spectral (ug cm-2)", "Fresh mass per area (g cm-2)", "Total chl spectral (mg g-1 FW)"] + \
           [f"{c} prior" for c in ["RWC", "MSI", "Pn", "gs", "E", "Ci", "LAI", "ChlA", "ChlB", "Car", "Anth", "Prot"]] + ["gs healthy", "E healthy"]
header(ip, 1, ICOL, IDW + [8, 18, 9, 10, 8, 8, 10, 10, 10, 11] + [9] * 12 + [9, 9])
IC = {h: j for j, h in enumerate(ICOL, start=1)}
for i, p in enumerate(ORDER, start=2):
    d = P[p]
    id_cells(ip, i, d)
    vals = [d["n"], d["leaves"], d["mean_grade"], d["area"], d["s_eff"], d["MTCI"], d["chl_rel"], d["Cab_area"], d["FW_area"], d["chl_spec"]]
    fm = ["0", None, "0.00", "0.0", "0.000", "0.000", "0.000", "0.0", "0.0000", "0.000"]
    for j, (v, f_) in enumerate(zip(vals, fm), start=len(ID) + 1):
        put(ip, i, j, v, fmt=f_)
    for c in ["RWC", "MSI", "Pn", "gs", "E", "Ci", "LAI", "ChlA", "ChlB", "Car", "Anth", "Prot"]:
        put(ip, i, IC[f"{c} prior"], d["prior"][c], fmt="0.000")
    put(ip, i, IC["gs healthy"], d["gs_H"], fmt="0.000")
    put(ip, i, IC["E healthy"], d["E_H"], fmt="0.000")
ip.freeze_panes = "C2"
INR = {p: i for i, p in enumerate(ORDER, start=2)}
ia = lambda col, r: f"Inputs!${CL(IC[col])}{r}"

# Estimates (formulas)
es = wb.create_sheet("Estimates")
EST = [
    ("RWC", "RWC (%)", "0.0"), ("MSI", "MSI (%)", "0.0"), ("Pn", "Pn (umol CO2 m-2 s-1)", "0.00"),
    ("gs", "gs (mol H2O m-2 s-1)", "0.000"), ("E", "E (mmol H2O m-2 s-1)", "0.00"), ("Ci", "Ci (umol mol-1)", "0"),
    ("LAI", "LAI (m2 m-2)", "0.00"), ("ChlA", "Chl a (ug mg-1 FW)", "0.000"), ("ChlB", "Chl b (ug mg-1 FW)", "0.000"),
    ("TotChl", "Total chl (ug mg-1 FW)", "0.000"), ("AB", "Chl a/b", "0.00"), ("Car", "Carotenoids (mg g-1 FW)", "0.000"),
    ("CarChl", "Car / total chl", "0.000"), ("Anth", "Anthocyanin (ug g-1 FW)", "0.0"), ("Prot", "Total protein (ug mg-1 FW)", "0.00"),
    ("WUE", "WUE Pn/E (umol mmol-1)", "0.00"), ("delta", "Chl final / prior (delta)", "0.000"),
]
header(es, 1, ID + [e[1] for e in EST], IDW + [10] * len(EST))
EC = {e[0]: len(ID) + 1 + j for j, e in enumerate(EST)}
ec = lambda code, r: f"{CL(EC[code])}{r}"
for i, p in enumerate(ORDER, start=2):
    d, r = P[p], INR[p]
    id_cells(es, i, d)
    pri_tot = f"({ia('ChlA prior', r)}+{ia('ChlB prior', r)})"
    f = {
        "TotChl": f"=EXP((1-{SREF['w_spec']})*LN({pri_tot})+{SREF['w_spec']}*LN({ia('Total chl spectral (mg g-1 FW)', r)}))",
        "delta": f"={ec('TotChl', i)}/{pri_tot}",
        "ChlA": f"={ec('TotChl', i)}*{ia('ChlA prior', r)}/{pri_tot}",
        "ChlB": f"={ec('TotChl', i)}*{ia('ChlB prior', r)}/{pri_tot}",
        "AB": f"={ec('ChlA', i)}/{ec('ChlB', i)}",
        "Car": f"={ec('TotChl', i)}*{ia('Car prior', r)}/{pri_tot}",
        "CarChl": f"={ec('Car', i)}/{ec('TotChl', i)}",
        "RWC": f"={ia('RWC prior', r)}",
        "MSI": f"={ia('MSI prior', r)}*{ec('delta', i)}^{SREF['e_MSI']}",
        "Pn": f"={ia('Pn prior', r)}*{ec('delta', i)}^{SREF['e_Pn']}",
        "gs": f"={ia('gs prior', r)}*{ec('delta', i)}^{SREF['e_gs']}",
        "E": f"={ia('E prior', r)}*{ec('gs', i)}/{ia('gs prior', r)}",
        "Ci": f"={ia('Ci prior', r)}*{ec('delta', i)}^{SREF['e_Ci']}",
        "LAI": f"={ia('LAI prior', r)}",
        "Anth": f"={ia('Anth prior', r)}*{ec('delta', i)}^{SREF['e_Anth']}",
        "Prot": f"={ia('Prot prior', r)}*{ec('delta', i)}^{SREF['e_Prot']}",
        "WUE": f"={ec('Pn', i)}/{ec('E', i)}",
    }
    for code, _, fm in EST:
        put(es, i, EC[code], f[code], FB if code in ("Pn", "TotChl", "RWC") else F, fmt=fm)
es.freeze_panes = "C2"
es.auto_filter.ref = f"A1:{CL(len(ID) + len(EST))}{len(ORDER) + 1}"
ESR = {p: i for i, p in enumerate(ORDER, start=2)}

# Ranges
rg = wb.create_sheet("Ranges")
RCODES = [e for e in EST if e[0] not in ("AB", "CarChl", "delta")]
rh = ID[:2]
for _, nm, _ in RCODES:
    rh += [f"{nm} low", f"{nm} high"]
header(rg, 1, rh, [6, 10] + [10] * (2 * len(RCODES)))
for i, p in enumerate(ORDER, start=2):
    put(rg, i, 1, p, FB)
    put(rg, i, 2, P[p]["t"], FB)
    for j, (code, _, fm) in enumerate(RCODES):
        src = f"Estimates!{ec(code, ESR[p])}"
        cv = SREF["CV_" + code]
        put(rg, i, 3 + 2 * j, f"={src}*(1-1.645*{cv})", fmt=fm)
        put(rg, i, 4 + 2 * j, f"={src}*(1+1.645*{cv})", fmt=fm)
rg.freeze_panes = "C2"

# Spectral features
sf = wb.create_sheet("Spectral_Features")
FEATS = [("MTCI", "MTCI (R754-R709)/(R709-R681)", "0.000"), ("CIre", "CI red-edge R780/R710-1", "0.000"),
         ("NDVI", "NDVI (R800,R670)", "0.000"), ("REIP", "Red-edge position, max 1st derivative (nm)", "0"),
         ("ARI", "ARI 1/R550-1/R700", "0.000"), ("NDWI", "NDWI (R860,R1240)", "0.000"),
         ("WBD1450", "1450 nm water band depth", "0.000")]
header(sf, 1, ID + [f[1] for f in FEATS], IDW + [14] * len(FEATS))
for i, p in enumerate(ORDER, start=2):
    id_cells(sf, i, P[p])
    for j, (k, _, fm) in enumerate(FEATS, start=len(ID) + 1):
        put(sf, i, j, P[p][k], fmt=fm)
sf.freeze_panes = "C2"

# Treatment summary (formulas)
ts = wb.create_sheet("Treatment_Summary")
ts.cell(row=1, column=1, value="Treatment means: healthy control plot vs mean of inoculated R1-R3 (formulas on Estimates)").font = FT
SC = ["RWC", "MSI", "Pn", "gs", "E", "Ci", "LAI", "TotChl", "Car", "Anth", "Prot"]
th = ["Treatment", "Growth stage", "Mean DI R1-R3 (%)"]
for c in SC:
    nm = next(e[1] for e in EST if e[0] == c)
    th += [f"{nm} healthy", f"{nm} R1-R3", "% change"]
header(ts, 2, th, [10, 12, 10] + [9] * (3 * len(SC)))
TRE = [f"D{a}V{b}" for a in range(1, 5) for b in range(1, 4)]
last = len(ORDER) + 1
for i, t in enumerate(TRE, start=3):
    hpl = next(p for p in ORDER if P[p]["t"] == t and P[p]["rep"] == "Control")
    put(ts, i, 1, t, FB)
    put(ts, i, 2, {"V": "Vegetative", "F": "Flowering", "P": "Pod formation"}[P[hpl]["stage"]])
    put(ts, i, 3, f'=AVERAGEIFS(Estimates!$G$2:$G${last},Estimates!$B$2:$B${last},A{i},Estimates!$E$2:$E${last},"<>Control")', fmt="0.0")
    for j, c in enumerate(SC):
        col = CL(EC[c])
        fm = next(e[2] for e in EST if e[0] == c)
        hc, dc, pc = 4 + 3 * j, 5 + 3 * j, 6 + 3 * j
        put(ts, i, hc, f"=Estimates!{col}{ESR[hpl]}", fmt=fm)
        put(ts, i, dc, f'=AVERAGEIFS(Estimates!${col}$2:${col}${last},Estimates!$B$2:$B${last},A{i},Estimates!$E$2:$E${last},"<>Control")', fmt=fm)
        put(ts, i, pc, f"={CL(dc)}{i}/{CL(hc)}{i}-1", fmt="+0.0%;-0.0%;0.0%")
ts.freeze_panes = "B3"

# Plot mean spectra
ms = wb.create_sheet("Plot_Mean_Spectra")
ms.cell(row=1, column=1, value="Plot-average ASD reflectance (raw, unsmoothed mean of the plot's spectra), 24/02/2026").font = FT
put(ms, 2, 1, "Plot", FW, HDR)
put(ms, 3, 1, "Treatment", FW, HDR)
put(ms, 4, 1, "Plot DI (%)", FW, HDR)
put(ms, 5, 1, "Wavelength (nm)", FW, HDR)
for j, p in enumerate(range(1, 49), start=2):
    put(ms, 2, j, p, FB)
    put(ms, 3, j, P[p]["t"])
    put(ms, 4, j, round(P[p]["di"], 1), fill=CF[P[p]["cls"]], fmt="0.0")
    put(ms, 5, j, f"P{p:02d}", FB)
    ms.column_dimensions[CL(j)].width = 8
for w_i, w in enumerate(WL):
    ms.cell(row=6 + w_i, column=1, value=int(w)).font = FB
    for j, p in enumerate(range(1, 49), start=2):
        ms.cell(row=6 + w_i, column=j, value=round(float(P[p]["spectrum"][w_i]), 5)).number_format = "0.0000"
ms.column_dimensions["A"].width = 15
ms.freeze_panes = "B6"

wb.move_sheet("Estimates", offset=1 - wb.sheetnames.index("Estimates"))
wb.save(OUT)
print("kappa", KAPPA, "saved", OUT)
